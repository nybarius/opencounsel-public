from __future__ import annotations

import copy
import hashlib
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import cast

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.paragraph import CT_P
from docx.shared import Pt
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph

from opencounsel.briefs.docx import BriefDocxError
from opencounsel.templates.build import (
    BLOCK_QUOTE_STYLE,
    BODY_NO_INDENT_STYLE,
    BODY_STYLE,
    CAPTION_STYLE,
    DOCUMENT_TITLE_STYLE,
    HEADING_STYLES,
    SECTION_HEADING_STYLE,
    SIGNATURE_LAST_STYLE,
    SIGNATURE_STYLE,
    TOA_TITLE_STYLE,
    TOC_TITLE_STYLE,
    _add_heading_numbering,
    _bind_heading_numbering,
    _clear_core_properties,
    _configure_section_geometry,
    _configure_settings,
    _configure_styles,
    _sanitize_package,
    configure_front_matter_section,
    configure_main_body_section,
)
from opencounsel.templates.profiles import FilingProfile, get_bundled_filing_profile
from opencounsel.templates.style_pack import resolve_word_style

_MAJOR_HEADINGS = {
    "ARGUMENT",
    "BACKGROUND",
    "CONCLUSION",
    "INTRODUCTION",
    "LEGAL STANDARD",
    "PRELIMINARY STATEMENT",
    "QUESTION PRESENTED",
    "QUESTIONS PRESENTED",
    "STATEMENT OF FACTS",
    "STATEMENT OF THE CASE",
    "STATEMENT OF THE QUESTIONS PRESENTED",
    "STANDARD OF REVIEW",
    "SUMMARY OF ARGUMENT",
}
_TOC_TITLE = "TABLE OF CONTENTS"
_TOA_TITLE = "TABLE OF AUTHORITIES"
_NUMBER_PREFIX = re.compile(
    r"^\s*(?P<label>(?:[IVXLCDM]+|[A-Z]|\d+|[a-z])\.)\s+(?P<body>\S.*)$"
)
_HEADING_NAME = re.compile(r"(?:heading|point|subpoint|argument)", re.IGNORECASE)
_TRACKED_TAGS = (
    b"<w:ins ",
    b"<w:ins>",
    b"<w:del ",
    b"<w:del>",
    b"<w:moveFrom ",
    b"<w:moveFrom>",
    b"<w:moveTo ",
    b"<w:moveTo>",
)


class FormattingProjectionError(BriefDocxError):
    pass


@dataclass(frozen=True, slots=True)
class FormattingDecision:
    paragraph_index: int
    original_style: str | None
    applied_style: str
    action: str


@dataclass(frozen=True, slots=True)
class FormattingProjectionResult:
    input_sha256: str
    output_sha256: str
    profile_id: str
    decisions: tuple[FormattingDecision, ...]
    inserted_toc_slot: bool
    inserted_toa_slot: bool
    inserted_merits_section: bool

    @property
    def applied_count(self) -> int:
        return len(self.decisions) + sum(
            (self.inserted_toc_slot, self.inserted_toa_slot, self.inserted_merits_section)
        )


def normalize_brief_formatting(
    source: Path,
    output: Path,
    profile: str | FilingProfile,
) -> FormattingProjectionResult:
    """Project a conventional brief onto the reviewed legal style system."""
    if not source.is_file() or source.is_symlink() or source.suffix.lower() != ".docx":
        raise FormattingProjectionError("brief must be a regular DOCX file")
    if source.resolve() == output.resolve():
        raise FormattingProjectionError("formatting projection must write a new DOCX")
    if output.exists() or output.is_symlink():
        raise FormattingProjectionError("formatting projection output already exists")
    _reject_unsafe_mutation(source)

    filing_profile = (
        get_bundled_filing_profile(profile) if isinstance(profile, str) else profile
    )
    resolved = resolve_word_style(filing_profile)
    document = Document(str(source))

    _configure_styles(document, filing_profile, resolved)
    numbering_id = _add_heading_numbering(
        document, filing_profile.structure.heading_numbering
    )
    _bind_heading_numbering(document, numbering_id)
    _configure_settings(document)
    _clear_core_properties(document)
    _configure_builtin_note_styles(document, resolved.font_family, resolved.footnote_font_size_pt)

    inserted_toc, inserted_toa = _ensure_front_matter_slots(document)
    decisions = _normalize_paragraphs(document)
    decisions.extend(_stabilize_standalone_page_breaks(document))
    inserted_section = _ensure_merits_section(document, filing_profile)
    for section in document.sections:
        _configure_section_geometry(section, filing_profile)

    output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    output.parent.chmod(0o700)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.stem}.", suffix=".docx", dir=output.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document.save(str(temporary))
        sanitized = _sanitize_package(temporary, resolved.font_family)
        os.replace(sanitized, output)
        output.chmod(0o600)
    finally:
        temporary.unlink(missing_ok=True)

    return FormattingProjectionResult(
        input_sha256=_hash_path(source),
        output_sha256=_hash_path(output),
        profile_id=filing_profile.profile_id,
        decisions=tuple(decisions),
        inserted_toc_slot=inserted_toc,
        inserted_toa_slot=inserted_toa,
        inserted_merits_section=inserted_section,
    )


def _reject_unsafe_mutation(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if any(name.startswith("_xmlsignatures/") for name in names):
                raise FormattingProjectionError("digitally signed DOCX cannot be reformatted")
            for name in names:
                if not name.startswith("word/") or not name.endswith(".xml"):
                    continue
                if any(tag in archive.read(name) for tag in _TRACKED_TAGS):
                    raise FormattingProjectionError(
                        "tracked changes must be accepted or rejected before formatting"
                    )
    except (OSError, zipfile.BadZipFile) as exc:
        raise FormattingProjectionError("DOCX is not a readable OOXML package") from exc


def _configure_builtin_note_styles(
    document: WordDocument, font_name: str, size_pt: float
) -> None:
    for style_name in ("Footnote Text", "Endnote Text"):
        try:
            style = cast(ParagraphStyle, document.styles[style_name])
        except KeyError:
            continue
        style.font.name = font_name
        style.font.size = Pt(size_pt)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)


def _ensure_front_matter_slots(document: WordDocument) -> tuple[bool, bool]:
    paragraphs = list(document.paragraphs)
    anchor = _first_paragraph(paragraphs, _MAJOR_HEADINGS) or (
        paragraphs[0] if paragraphs else document.add_paragraph()
    )
    toc_title = _first_paragraph(paragraphs, {_TOC_TITLE})
    toa_title = _first_paragraph(paragraphs, {_TOA_TITLE})
    inserted_toc = not _has_exact_slot(document, "[TOC]")
    inserted_toa = not _has_exact_slot(document, "[TOA]")

    if toc_title is None and toa_title is None:
        toc_title = _insert_before(anchor, _TOC_TITLE)
        if inserted_toc:
            _insert_before(anchor, "[TOC]")
        toa_title = _insert_before(anchor, _TOA_TITLE)
        if inserted_toa:
            _insert_before(anchor, "[TOA]")
    else:
        if toc_title is None:
            toc_title = _insert_before(toa_title or anchor, _TOC_TITLE)
        if inserted_toc:
            _insert_after(toc_title, "[TOC]")
        if toa_title is None:
            toa_title = _insert_before(anchor, _TOA_TITLE)
        if inserted_toa:
            _insert_after(toa_title, "[TOA]")

    toc_title.style = document.styles[TOC_TITLE_STYLE]
    toa_title.style = document.styles[TOA_TITLE_STYLE]
    return inserted_toc, inserted_toa


def _first_paragraph(
    paragraphs: list[Paragraph], normalized_values: set[str]
) -> Paragraph | None:
    for paragraph in paragraphs:
        if _normalized(paragraph.text) in normalized_values:
            return paragraph
    return None


def _has_exact_slot(document: WordDocument, text: str) -> bool:
    return sum(paragraph.text.strip() == text for paragraph in document.paragraphs) == 1


def _insert_before(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    inserted = Paragraph(cast(CT_P, element), paragraph._parent)
    inserted.add_run(text)
    return inserted


def _insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(cast(CT_P, element), paragraph._parent)
    inserted.add_run(text)
    return inserted


def _normalize_paragraphs(document: WordDocument) -> list[FormattingDecision]:
    decisions: list[FormattingDecision] = []
    in_argument = False
    for index, paragraph in enumerate(document.paragraphs):
        display_text = " ".join(paragraph.text.split())
        if not display_text:
            continue
        text = display_text.upper()
        original_style = paragraph.style.name if paragraph.style is not None else None
        major_heading = _major_heading_text(display_text)
        if major_heading is not None:
            target = SECTION_HEADING_STYLE
            replacement = major_heading if paragraph.text != major_heading else None
            in_argument = major_heading == "ARGUMENT"
        else:
            target, replacement = _target_style(
                paragraph,
                text,
                display_text,
                in_argument=in_argument,
            )
        if replacement is not None and replacement != paragraph.text:
            paragraph.text = replacement
        paragraph.style = document.styles[target]
        _clear_direct_paragraph_formatting(paragraph)
        _clear_direct_run_typography(paragraph)
        if original_style != target or replacement is not None:
            decisions.append(
                FormattingDecision(
                    paragraph_index=index,
                    original_style=original_style,
                    applied_style=target,
                    action="removed typed heading number and applied legal style"
                    if replacement is not None
                    else "applied legal style",
                )
            )
    return decisions


def _major_heading_text(display_text: str) -> str | None:
    normalized = display_text.upper()
    if normalized in _MAJOR_HEADINGS:
        return normalized
    prefix = _NUMBER_PREFIX.match(display_text)
    if prefix is None or _label_level(prefix.group("label")) != 0:
        return None
    body = " ".join(prefix.group("body").split()).upper()
    return body if body in _MAJOR_HEADINGS else None


def _stabilize_standalone_page_breaks(
    document: WordDocument,
) -> list[FormattingDecision]:
    """Move a bare page break onto its target so a full prior page cannot add a blank."""
    decisions: list[FormattingDecision] = []
    for index, paragraph in enumerate(list(document.paragraphs)):
        if paragraph.text.strip():
            continue
        breaks = paragraph._p.xpath(".//w:br[@w:type='page']")
        unsafe = paragraph._p.xpath(
            ".//w:sectPr | .//w:drawing | .//w:object | .//w:fldChar | .//w:instrText"
        )
        next_element = paragraph._p.getnext()
        if (
            len(breaks) != 1
            or unsafe
            or next_element is None
            or next_element.tag != qn("w:p")
        ):
            continue
        target = Paragraph(cast(CT_P, next_element), paragraph._parent)
        if not target.text.strip():
            continue
        target.paragraph_format.page_break_before = True
        parent = paragraph._p.getparent()
        if parent is None:
            raise FormattingProjectionError("page-break paragraph has no parent")
        parent.remove(paragraph._p)
        style_name = target.style.name if target.style is not None else BODY_STYLE
        decisions.append(
            FormattingDecision(
                paragraph_index=index,
                original_style=style_name,
                applied_style=style_name,
                action=(
                    "replaced standalone page-break paragraph with stable "
                    "page-break-before"
                ),
            )
        )
    return decisions


def _target_style(
    paragraph: Paragraph,
    text: str,
    display_text: str,
    *,
    in_argument: bool = False,
) -> tuple[str, str | None]:
    if text == _TOC_TITLE:
        return TOC_TITLE_STYLE, None
    if text == _TOA_TITLE:
        return TOA_TITLE_STYLE, None
    if text in _MAJOR_HEADINGS:
        return SECTION_HEADING_STYLE, None
    if text in {"[TOC]", "[TOA]"}:
        return BODY_NO_INDENT_STYLE, None

    style_name = paragraph.style.name if paragraph.style is not None else ""
    style_id = paragraph.style.style_id if paragraph.style is not None else ""
    generic_levels = {
        HEADING_STYLES[level][0]: level for level in range(len(HEADING_STYLES))
    }
    generic_levels.update(
        {HEADING_STYLES[level][1]: level for level in range(len(HEADING_STYLES))}
    )
    for candidate in (style_name, style_id):
        if candidate in generic_levels:
            return HEADING_STYLES[generic_levels[candidate]][0], None

    built_in_level = _built_in_heading_level(style_name, style_id)
    prefix = _NUMBER_PREFIX.match(paragraph.text)
    if prefix is not None and (
        built_in_level is not None
        or _looks_like_heading(paragraph, style_name, display_text)
    ):
        level = _label_level(prefix.group("label"))
        if in_argument and level > 0:
            level -= 1
        return HEADING_STYLES[level][0], prefix.group("body")
    if built_in_level is not None:
        return HEADING_STYLES[min(built_in_level, 3)][0], None

    if style_name in {SIGNATURE_STYLE, SIGNATURE_LAST_STYLE}:
        return style_name, None
    if "signature" in style_name.casefold() or text.casefold() == "respectfully submitted,":
        return SIGNATURE_STYLE, None
    if "caption" in style_name.casefold():
        return CAPTION_STYLE, None
    if "title" in style_name.casefold() and len(text) < 200:
        return DOCUMENT_TITLE_STYLE, None
    if "quote" in style_name.casefold() or _is_block_quote(paragraph):
        return BLOCK_QUOTE_STYLE, None
    if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER and len(text) < 200:
        return CAPTION_STYLE, None
    return BODY_STYLE, None


def _built_in_heading_level(style_name: str, style_id: str) -> int | None:
    for candidate in (style_name.replace(" ", ""), style_id.replace(" ", "")):
        match = re.fullmatch(r"Heading([1-4])", candidate, re.IGNORECASE)
        if match:
            return int(match.group(1)) - 1
    if _HEADING_NAME.search(style_name):
        digits = re.findall(r"[1-4]", style_name)
        return int(digits[-1]) - 1 if digits else 0
    return None


def _looks_like_heading(paragraph: Paragraph, style_name: str, text: str) -> bool:
    if _HEADING_NAME.search(style_name):
        return True
    letters = re.sub(r"[^A-Za-z]", "", text)
    if letters and letters == letters.upper():
        return True
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.bold is True for run in runs)


def _label_level(label: str) -> int:
    raw = label[:-1]
    if re.fullmatch(r"[IVXLCDM]+", raw):
        return 0
    if raw.isdigit():
        return 2
    if raw.islower():
        return 3
    if len(raw) == 1:
        return 1
    return 0


def _is_block_quote(paragraph: Paragraph) -> bool:
    left = paragraph.paragraph_format.left_indent
    first = paragraph.paragraph_format.first_line_indent
    return left is not None and left.inches >= 0.45 and (
        first is None or first.inches <= 0
    )


def _clear_direct_paragraph_formatting(paragraph: Paragraph) -> None:
    formatting = paragraph.paragraph_format
    formatting.alignment = None
    formatting.left_indent = None
    formatting.right_indent = None
    formatting.first_line_indent = None
    formatting.line_spacing = None
    formatting.space_before = None
    formatting.space_after = None
    formatting.keep_with_next = None
    formatting.keep_together = None
    formatting.page_break_before = None


def _clear_direct_run_typography(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.font.name = None
        run.font.size = None


def _ensure_merits_section(document: WordDocument, profile: FilingProfile) -> bool:
    paragraphs = list(document.paragraphs)
    merits = _first_paragraph(
        paragraphs,
        {
            "QUESTION PRESENTED",
            "QUESTIONS PRESENTED",
            "STATEMENT OF QUESTIONS PRESENTED",
            "STATEMENT OF THE QUESTIONS PRESENTED",
            "INTRODUCTION",
            "PRELIMINARY STATEMENT",
            "ARGUMENT",
        },
    )
    inserted = False
    if merits is not None and len(document.sections) == 1:
        index = paragraphs.index(merits)
        if index > 0:
            previous = paragraphs[index - 1]
            section_properties = copy.deepcopy(document.sections[0]._sectPr)
            section_type = section_properties.find(qn("w:type"))
            if section_type is None:
                section_type = OxmlElement("w:type")
                section_properties.insert(0, section_type)
            section_type.set(qn("w:val"), "nextPage")
            previous._p.get_or_add_pPr().append(section_properties)
            inserted = True

    sections = list(document.sections)
    resolved = resolve_word_style(profile)
    if sections:
        configure_front_matter_section(document, sections[0], profile, resolved)
        configure_main_body_section(document, sections[-1], profile, resolved)
    return inserted


def _normalized(text: str) -> str:
    return " ".join(text.split()).upper()


def _hash_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


__all__ = [
    "FormattingDecision",
    "FormattingProjectionError",
    "FormattingProjectionResult",
    "normalize_brief_formatting",
]
