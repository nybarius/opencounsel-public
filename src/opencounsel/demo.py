from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    RectangleObject,
)

from opencounsel.source.pdf_roa import package_searchable_pdf

_SOURCE_PDF = "https://cdn.openai.com/pdf/gov.uscourts.nysd.612697.52.0_1.pdf"
_AUTHORITY_LINKS = {
    "ashcroft-v-iqbal.pdf": (
        "https://www.supremecourt.gov/opinions/boundvolumes/556bv.pdf"
    ),
}

_BLOCKS = (
    ("heading", "I. INTRODUCTION", "Arial", 16),
    (
        "body",
        "The artificial intelligence tool known as ChatGPT is many things: a revolutionary "
        "technology with the potential to augment human capabilities; an accelerator for "
        "scientific and medical breakthroughs; a mechanism for making existing technologies "
        "accessible to more people; and a creative and computational tool.",
        "Times New Roman",
        11,
    ),
    (
        "body",
        "Contrary to the allegations in the Complaint, however, ChatGPT is not in any way a "
        "substitute for a subscription to The New York Times. In the ordinary course, one "
        "cannot use ChatGPT to serve up Times articles at will.",
        "Calibri",
        13,
    ),
    (
        "body",
        "There is a genuinely important issue at the heart of this lawsuit: whether it is fair "
        "use under copyright law to use publicly accessible content to train generative AI "
        "models to learn about language, grammar, syntax, and facts.",
        "Arial",
        10,
    ),
    (
        "body",
        "For more than a century, courts have recognized that knowledge and facts are free for "
        "common use. See Baker v. Selden, 101 U.S. 99 (1879). That principle frames the dispute "
        "but does not itself resolve the pleaded claims.",
        "Times New Roman",
        12,
    ),
    ("heading", "II. BACKGROUND", "Courier New", 13),
    ("heading", "A. OpenAI's Pioneering Research", "Arial", 11),
    (
        "body",
        "OpenAI began researching general-purpose language models years before this suit. The "
        "research culminated in systems capable of performing a broad range of language tasks.",
        "Calibri",
        10,
    ),
    (
        "heading",
        "B. Reliance on Longstanding Fair Use Principles",
        "Calibri",
        15,
    ),
    (
        "body",
        "Copyright is not a veto right over transformative technologies that use existing works "
        "internally without disseminating them. Authors Guild v. Google, Inc., 804 F.3d 202 "
        "(2d Cir. 2015), upheld book search as a transformative use. Sony Corp. of America v. "
        "Universal City Studios, Inc., 464 U.S. 417 (1984), likewise protects technologies "
        "capable of substantial noninfringing uses.",
        "Times New Roman",
        12,
    ),
    ("heading", "III. LEGAL STANDARD", "Times New Roman", 14),
    (
        "body",
        "A complaint must contain sufficient factual matter, accepted as true, to state a claim "
        "to relief that is plausible on its face. Ashcroft v. Iqbal, 556 U.S. 662, 678 (2009). "
        "Conclusory allegations do not suffice.",
        "Arial",
        10,
    ),
    ("heading", "IV. ARGUMENT", "Arial", 17),
    (
        "heading",
        "A. The Times Cannot Sue for Conduct Occurring More than Three Years Ago",
        "Courier New",
        11,
    ),
    (
        "body",
        "The direct-infringement claim rests in part on training activities alleged to have "
        "occurred outside the Copyright Act's three-year limitations period. See 17 U.S.C. "
        "section 507(b).",
        "Calibri",
        13,
    ),
    (
        "heading",
        "B. The Complaint Fails to State a Contributory Infringement Claim",
        "Times New Roman",
        15,
    ),
    (
        "body",
        "A contributory-infringement claim requires direct infringement, knowledge of the "
        "specific infringing activity, and material contribution. General awareness that a "
        "product might be misused is not enough.",
        "Arial",
        11,
    ),
    (
        "heading",
        "C. The DMCA Claim Fails for Multiple Independent Reasons",
        "Calibri",
        12,
    ),
    (
        "body",
        "The Complaint does not adequately identify the copyright-management information at "
        "issue, allege its removal from a particular work, or plead the required scienter and "
        "distribution elements.",
        "Courier New",
        10,
    ),
    ("heading", "V. CONCLUSION", "Times New Roman", 16),
    (
        "body",
        "The Court should dismiss the challenged claims and narrow the case to the issues that "
        "can be litigated under governing copyright law.",
        "Arial",
        12,
    ),
)

_AUTHORITIES = {
    "ashcroft-v-iqbal.pdf": (
        "ASHCROFT v. IQBAL, 556 U.S. 662 (2009)",
        "A complaint must contain sufficient factual matter, accepted as true, to state a claim "
        "to relief that is plausible on its face.",
    ),
    "authors-guild-v-google.pdf": (
        "AUTHORS GUILD v. GOOGLE, INC., 804 F.3d 202 (2d Cir. 2015)",
        "Google's making of a digital copy to provide a search function is a transformative use.",
    ),
    "baker-v-selden.pdf": (
        "BAKER v. SELDEN, 101 U.S. 99 (1879)",
        "The truths of a science or the methods of an art are the common property "
        "of the whole world.",
    ),
    "sony-v-universal.pdf": (
        "SONY CORP. OF AMERICA v. UNIVERSAL CITY STUDIOS, INC., 464 U.S. 417 (1984)",
        "The sale of copying equipment does not constitute contributory infringement if the "
        "product is widely used for legitimate, unobjectionable purposes.",
    ),
}


def build_synthetic_demo(inbox: Path) -> tuple[Path, Path]:
    """Generate a privacy-safe demo derived from a real public OpenAI filing."""
    inbox.mkdir(parents=True, exist_ok=True, mode=0o700)
    inbox.chmod(0o700)
    brief = inbox / "brief.docx"
    display_brief = inbox / "openai-mtd-deformatted.docx"
    record_pdf = inbox / "record.pdf"
    record_package = inbox / "record.zip"

    _write_deformatted_openai_brief(display_brief)
    shutil.copyfile(display_brief, brief)
    brief.chmod(0o600)
    _write_provenance(inbox / "demo-source.json")
    _write_authority_sources(inbox / "authority-sources")
    _write_text_pdf(
        record_pdf,
        (
            "SYNTHETIC PLACEHOLDER RECORD",
            "The contest demonstration uses the explicit no-record path.",
            "This PDF remains available for regression tests of raw-PDF intake.",
        ),
    )
    package_searchable_pdf(
        record_pdf,
        record_package,
        first_record_page=1,
        numbering_verified=True,
    )
    return brief, record_package


def _write_deformatted_openai_brief(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(1.15)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(1.18)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)

    _add_demo_caption(document)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _add_bad_paragraph(document, "TABLE OF CONTENTS", "Times New Roman", 14, bold=True)
    document.add_paragraph("[TOC]")
    _add_bad_paragraph(document, "TABLE OF AUTHORITIES", "Arial", 13, bold=True)
    document.add_paragraph("[TOA]")

    for kind, text, font, size in _BLOCKS:
        if kind == "heading" and text.startswith(("II.", "III.", "IV.", "V.")):
            document.add_page_break()
        if kind == "heading":
            _add_bad_paragraph(document, text, font, size, bold=True)
        else:
            _add_bad_paragraph(document, text, font, size)

    _add_bad_paragraph(
        document,
        "PUBLIC-FILING DEMONSTRATION - RECONSTRUCTED FROM DKT. 52 - NOT FOR FILING",
        "Courier New",
        9,
        bold=True,
    )
    document.save(str(path))
    path.chmod(0o600)


def _add_demo_caption(document: DocumentType) -> None:
    court = document.add_paragraph()
    court.alignment = WD_ALIGN_PARAGRAPH.CENTER
    court.paragraph_format.line_spacing = 1.0
    court.paragraph_format.space_before = Pt(0)
    court.paragraph_format.space_after = Pt(14)
    court.paragraph_format.keep_with_next = True
    _add_run(
        court,
        "UNITED STATES DISTRICT COURT\nSOUTHERN DISTRICT OF NEW YORK",
        "Times New Roman",
        12,
        bold=True,
    )

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = (Inches(3.75), Inches(2.75))
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width
    _fix_caption_table_geometry(table._tbl)

    left, right = table.rows[0].cells
    for cell, width in zip((left, right), widths, strict=True):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_width(cell, int(width.twips))
        _set_cell_margins(cell, top=60, right=100, bottom=60, left=100)
    _set_cell_borders(
        left,
        top="single",
        right="single",
        bottom="single",
        left="nil",
    )
    _set_cell_borders(
        right,
        top="nil",
        right="nil",
        bottom="nil",
        left="nil",
    )

    _replace_cell_paragraphs(
        left,
        (
            ("THE NEW YORK TIMES COMPANY,", WD_ALIGN_PARAGRAPH.LEFT, False, 2, 12),
            ("Plaintiff,", WD_ALIGN_PARAGRAPH.CENTER, False, 14, 12),
            ("v.", WD_ALIGN_PARAGRAPH.CENTER, False, 14, 12),
            (
                "MICROSOFT CORPORATION, OPENAI, INC.,\n"
                "OPENAI LP, OPENAI GP, LLC, OPENAI, LLC,\n"
                "OPENAI OPCO LLC, OPENAI GLOBAL LLC,\n"
                "OAI CORPORATION, LLC, OPENAI\n"
                "HOLDINGS, LLC,",
                WD_ALIGN_PARAGRAPH.LEFT,
                False,
                2,
                12,
            ),
            ("Defendants.", WD_ALIGN_PARAGRAPH.CENTER, False, 0, 12),
        ),
    )
    _replace_cell_paragraphs(
        right,
        (
            (
                "Case No. 1:23-cv-11195 (SHS) (OTW)",
                WD_ALIGN_PARAGRAPH.LEFT,
                False,
                24,
                10,
            ),
            (
                "MEMORANDUM OF LAW\n"
                "IN SUPPORT OF OPENAI\n"
                "DEFENDANTS\N{RIGHT SINGLE QUOTATION MARK} MOTION TO\n"
                "DISMISS",
                WD_ALIGN_PARAGRAPH.CENTER,
                True,
                24,
                12,
            ),
            (
                "ORAL ARGUMENT REQUESTED",
                WD_ALIGN_PARAGRAPH.CENTER,
                True,
                0,
                12,
            ),
        ),
    )


def _fix_caption_table_geometry(table: object) -> None:
    table_properties = table.tblPr  # type: ignore[attr-defined]
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    width = table_properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        table_properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), "9360")
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "nil")
    grid = table.tblGrid  # type: ignore[attr-defined]
    for column, value in zip(grid.gridCol_lst, ("5400", "3960"), strict=True):
        column.set(qn("w:w"), value)


def _set_cell_width(cell: _Cell, twips: int) -> None:
    width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(twips))


def _set_cell_margins(
    cell: _Cell,
    *,
    top: int,
    right: int,
    bottom: int,
    left: int,
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (
        ("top", top),
        ("right", right),
        ("bottom", bottom),
        ("left", left),
    ):
        margin = margins.find(qn(f"w:{edge}"))
        if margin is None:
            margin = OxmlElement(f"w:{edge}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def _set_cell_borders(cell: _Cell, **values: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "right", "bottom", "left"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), values[edge])
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")


def _replace_cell_paragraphs(
    cell: _Cell,
    values: tuple[tuple[str, int, bool, int, int], ...],
) -> None:
    for index, (text, alignment, bold, space_after, font_size) in enumerate(values):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = alignment
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(space_after)
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.left_indent = Inches(0)
        paragraph.paragraph_format.right_indent = Inches(0)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = index < len(values) - 1
        _add_run(
            paragraph,
            text,
            "Times New Roman",
            font_size,
            bold=bold,
        )


def _add_bad_paragraph(
    document: DocumentType,
    text: str,
    font: str,
    size: int,
    *,
    bold: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3 if size % 2 else 11)
    paragraph.paragraph_format.line_spacing = 1.0 if size < 12 else 1.35
    _add_run(paragraph, text, font, size, bold=bold)


def _add_run(
    paragraph: Paragraph,
    text: str,
    font: str,
    size: int,
    *,
    bold: bool = False,
) -> None:
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold


def _write_provenance(path: Path) -> None:
    payload = {
        "case": "The New York Times Company v. Microsoft Corporation, et al.",
        "court": "S.D.N.Y.",
        "docket": "1:23-cv-11195-SHS-OTW",
        "document": 52,
        "filed": "2024-02-26",
        "source_pdf": _SOURCE_PDF,
        "transformation": (
            "public filing text reconstructed as intentionally malformed DOCX; "
            "populated TOC and TOA omitted"
        ),
    }
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    path.chmod(0o600)


def _write_authority_sources(directory: Path) -> None:
    directory.mkdir(mode=0o700, exist_ok=True)
    for filename, lines in _AUTHORITIES.items():
        _write_text_pdf(
            directory / filename,
            lines,
            link_url=_AUTHORITY_LINKS.get(filename),
        )


def _write_text_pdf(
    path: Path,
    lines: tuple[str, ...],
    *,
    link_url: str | None = None,
) -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for page_index, line in enumerate(lines):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        stream = DecodedStreamObject()
        escaped = line.replace("\\", "\\\\")
        escaped = escaped.replace("(", "\\(").replace(")", "\\)")
        stream.set_data(
            f"BT\n/F1 11 Tf\n54 738 Td\n({escaped}) Tj\nET".encode("cp1252")
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
        if page_index == 0 and link_url is not None:
            writer.add_uri(
                page_index,
                link_url,
                RectangleObject((54, 716, 558, 742)),
            )
    with path.open("wb") as target:
        writer.write(target)
    path.chmod(0o600)


__all__ = ["build_synthetic_demo"]
