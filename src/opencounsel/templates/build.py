from __future__ import annotations

import hashlib
import os
import tempfile
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import cast

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Inches, Pt, RGBColor
from docx.styles.style import CharacterStyle, ParagraphStyle
from lxml import etree

from opencounsel.templates.profiles import (
    FilingProfile,
    HeadingNumbering,
    get_bundled_filing_profile,
)
from opencounsel.templates.style_pack import (
    CLEAN_SERIF_STYLE_PACK,
    ResolvedWordStyle,
    WordStylePack,
    resolve_word_style,
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"w": W_NS, "a": A_NS}

BODY_STYLE = "_LegalBody"
BODY_NO_INDENT_STYLE = "_LegalBodyNoIndent"
BLOCK_QUOTE_STYLE = "_LegalBlockQuote"
FOOTNOTE_STYLE = "_LegalFootnote"
SECTION_HEADING_STYLE = "_LegalHeadingBoldCtrUnderline"
ARGUMENT_OUTLINE_STYLE_NAME = "_LegalOutline"
ARGUMENT_OUTLINE_STYLE_ID = "LegalOutline"
DOCUMENT_TITLE_STYLE = "_LegalDocumentTitle"
CAPTION_STYLE = "_LegalCaption"
SIGNATURE_STYLE = "_LegalSignature"
SIGNATURE_LAST_STYLE = "_LegalSignatureLast"
CERTIFICATE_STYLE = "_LegalCertificate"
NUMBERED_ITEM_STYLE = "_LegalNumberedItem"
TOC_TITLE_STYLE = "_LegalTOCTitle"
TOA_TITLE_STYLE = "_LegalTOATitle"
TOA_CATEGORY_STYLE = "_LegalTOACategory"
FOOTER_STYLE = "_LegalFooter"
CASE_NAME_STYLE = "_LegalCaseName"
HYPERLINK_STYLE = "Hyperlink"

ARGUMENT_OUTLINE_NSID = "4C474C31"
ARGUMENT_OUTLINE_TEMPLATE = "4C474C32"

HEADING_STYLES: tuple[tuple[str, str, int], ...] = (
    ("_LegalHeadingNum1", "LegalHeadingNum1", 0),
    ("_LegalHeadingNum2", "LegalHeadingNum2", 1),
    ("_LegalHeadingNum3", "LegalHeadingNum3", 2),
    ("_LegalHeadingNum4", "LegalHeadingNum4", 3),
)
HEADING_TEXT_TWIPS: tuple[int, ...] = (720, 1440, 2160, 2880)
HEADING_HANGING_TWIPS = 720
TOC_LEVEL_COUNT = 5

STYLE_BINDINGS: tuple[tuple[str, str, int | None], ...] = (
    ("body", BODY_STYLE, None),
    ("body-no-indent", BODY_NO_INDENT_STYLE, None),
    ("block-quote", BLOCK_QUOTE_STYLE, None),
    ("footnote", FOOTNOTE_STYLE, None),
    ("section-heading", SECTION_HEADING_STYLE, 1),
    ("argument-outline", ARGUMENT_OUTLINE_STYLE_NAME, None),
    ("point-heading", HEADING_STYLES[0][0], 2),
    ("subpoint-heading", HEADING_STYLES[1][0], 3),
    ("subsubpoint-heading", HEADING_STYLES[2][0], 4),
    ("paragraph-heading", HEADING_STYLES[3][0], 5),
    ("toa-category-heading", TOA_CATEGORY_STYLE, None),
    ("document-title", DOCUMENT_TITLE_STYLE, None),
    ("caption", CAPTION_STYLE, None),
    ("signature", SIGNATURE_STYLE, None),
    ("signature-last", SIGNATURE_LAST_STYLE, None),
    ("certificate", CERTIFICATE_STYLE, None),
    ("numbered-item", NUMBERED_ITEM_STYLE, None),
    ("case-name", CASE_NAME_STYLE, None),
    ("hyperlink", HYPERLINK_STYLE, None),
)


class WordTemplateBuildError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class TemplateBuildResult:
    profile_id: str
    profile_version: str
    style_pack_id: str
    style_pack_version: str
    template_sha256: str
    output: Path
    manifest: Path | None


def build_blank_word_template(
    profile: str | FilingProfile,
    output: Path,
    *,
    style_pack: WordStylePack = CLEAN_SERIF_STYLE_PACK,
    manifest_path: Path | None = None,
) -> TemplateBuildResult:
    filing_profile = (
        get_bundled_filing_profile(profile) if isinstance(profile, str) else profile
    )
    resolved = resolve_word_style(filing_profile, style_pack)
    _validate_new_output(output, ".docx")
    if manifest_path is not None:
        _validate_new_output(manifest_path, ".toml")
    output.parent.mkdir(parents=True, exist_ok=True)
    if manifest_path is not None:
        manifest_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_section_geometry(document.sections[0], filing_profile)
    _configure_styles(document, filing_profile, resolved)
    numbering_id = _add_heading_numbering(
        document, filing_profile.structure.heading_numbering
    )
    _bind_heading_numbering(document, numbering_id)
    configure_front_matter_section(
        document, document.sections[0], filing_profile, resolved
    )
    _configure_settings(document)
    _clear_core_properties(document)
    document.add_paragraph(style=document.styles[BODY_STYLE])

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.stem}.", suffix=".docx", dir=output.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document.save(str(temporary))
        sanitized = _sanitize_package(temporary, resolved.font_family)
        os.replace(sanitized, output)
        output.chmod(0o600)
    finally:
        temporary.unlink(missing_ok=True)

    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    if manifest_path is not None:
        _write_manifest(manifest_path, filing_profile, style_pack, digest)
    return TemplateBuildResult(
        profile_id=filing_profile.profile_id,
        profile_version=filing_profile.version,
        style_pack_id=style_pack.style_pack_id,
        style_pack_version=style_pack.version,
        template_sha256=digest,
        output=output,
        manifest=manifest_path,
    )


def configure_front_matter_section(
    document: WordDocument,
    section: Section,
    profile: FilingProfile,
    resolved: ResolvedWordStyle | None = None,
) -> None:
    style = resolved or resolve_word_style(profile)
    _configure_section_geometry(section, profile)
    section.different_first_page_header_footer = False
    section.footer.is_linked_to_previous = False
    _set_page_numbering(section, fmt="lowerRoman", start=1)
    _set_page_number_footer(document, section.footer, style, placeholder="i")


def configure_main_body_section(
    document: WordDocument,
    section: Section,
    profile: FilingProfile,
    resolved: ResolvedWordStyle | None = None,
) -> None:
    style = resolved or resolve_word_style(profile)
    _configure_section_geometry(section, profile)
    section.different_first_page_header_footer = True
    section.footer.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False
    _set_page_numbering(section, fmt="decimal", start=1)
    _clear_footer(section.first_page_footer)
    _set_page_number_footer(document, section.footer, style, placeholder="2")


def _validate_new_output(path: Path, suffix: str) -> None:
    if path.suffix.lower() != suffix:
        raise WordTemplateBuildError(f"output must use the {suffix.upper()} extension")
    if path.exists() or path.is_symlink():
        raise WordTemplateBuildError("template outputs must be new files")


def _configure_section_geometry(section: Section, profile: FilingProfile) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    margin = Inches(profile.page.min_margin_in)
    section.top_margin = margin
    section.right_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.25)


def _configure_styles(
    document: WordDocument,
    profile: FilingProfile,
    resolved: ResolvedWordStyle,
) -> None:
    styles = document.styles
    body_alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
        if resolved.body_alignment == "left"
        else WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    text_width = 8.5 - (2 * profile.page.min_margin_in)

    normal = cast(ParagraphStyle, styles["Normal"])
    _font(normal, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        normal,
        alignment=body_alignment,
        line_spacing=2.0,
        first_line=0.5,
        keep_with_next=False,
        keep_together=False,
    )

    body = _paragraph_style(styles, BODY_STYLE, normal)
    _font(body, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        body,
        alignment=body_alignment,
        line_spacing=2.0,
        first_line=0.5,
        keep_with_next=False,
        keep_together=False,
    )

    body_no_indent = _paragraph_style(styles, BODY_NO_INDENT_STYLE, body)
    _font(body_no_indent, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        body_no_indent,
        alignment=body_alignment,
        line_spacing=2.0,
        first_line=0,
        keep_with_next=False,
        keep_together=False,
    )

    quote = _paragraph_style(styles, BLOCK_QUOTE_STYLE, body_no_indent)
    _font(quote, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        quote,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=1.0,
        left=0.5,
        right=0.5,
        keep_with_next=False,
        keep_together=True,
    )

    footnote = _paragraph_style(styles, FOOTNOTE_STYLE, body_no_indent)
    _font(footnote, resolved.font_family, resolved.footnote_font_size_pt)
    _paragraph(
        footnote,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=1.0,
        first_line=0,
        keep_with_next=False,
        keep_together=False,
    )

    section_heading = _paragraph_style(
        styles, SECTION_HEADING_STYLE, body_no_indent
    )
    _font(
        section_heading,
        resolved.font_family,
        resolved.body_font_size_pt,
        bold=True,
        underline=True,
    )
    _paragraph(
        section_heading,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
        space_before=12,
        space_after=12,
        keep_with_next=True,
        keep_together=True,
    )
    _outline_level(section_heading, 0)

    for (name, _, list_level), outline_level, space_before in zip(
        HEADING_STYLES,
        range(1, 5),
        (12, 6, 6, 6),
        strict=True,
    ):
        heading = _paragraph_style(styles, name, body_no_indent)
        heading.base_style = None
        _font(heading, resolved.font_family, resolved.body_font_size_pt, bold=True)
        _paragraph(
            heading,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            line_spacing=1.0,
            space_before=space_before,
            space_after=12,
            keep_with_next=True,
            keep_together=True,
        )
        _outline_level(heading, outline_level)
        _clear_heading_style_geometry(heading)
        assert list_level == outline_level - 1

    caption = _paragraph_style(styles, CAPTION_STYLE, body_no_indent)
    _font(caption, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        caption,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
        keep_with_next=False,
        keep_together=True,
    )

    document_title = _paragraph_style(styles, DOCUMENT_TITLE_STYLE, body_no_indent)
    _font(document_title, resolved.font_family, resolved.body_font_size_pt, bold=True)
    _paragraph(
        document_title,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
        space_after=12,
        keep_with_next=True,
        keep_together=True,
    )

    signature = _paragraph_style(styles, SIGNATURE_STYLE, body_no_indent)
    _font(signature, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        signature,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        left=3.25,
        keep_with_next=True,
        keep_together=True,
    )

    signature_last = _paragraph_style(styles, SIGNATURE_LAST_STYLE, signature)
    _font(signature_last, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        signature_last,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        left=3.25,
        keep_with_next=False,
        keep_together=True,
    )

    certificate = _paragraph_style(styles, CERTIFICATE_STYLE, body_no_indent)
    _font(certificate, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        certificate,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        space_before=12,
        keep_with_next=False,
        keep_together=True,
    )

    numbered = _paragraph_style(styles, NUMBERED_ITEM_STYLE, body_no_indent)
    _font(numbered, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        numbered,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=2.0,
        left=0.5,
        first_line=-0.5,
        keep_with_next=False,
        keep_together=True,
    )

    toc_title = _paragraph_style(styles, TOC_TITLE_STYLE, section_heading)
    _font(
        toc_title,
        resolved.font_family,
        resolved.body_font_size_pt,
        bold=True,
        underline=True,
    )
    _outline_level(toc_title, 9)

    for level in range(1, TOC_LEVEL_COUNT + 1):
        toc = _paragraph_style(styles, f"TOC {level}", body_no_indent)
        _font(toc, resolved.font_family, resolved.body_font_size_pt)
        left = 0.5 * (level - 1)
        first_line = -0.5 if level > 1 else 0.0
        _paragraph(
            toc,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            line_spacing=1.0,
            left=left,
            first_line=first_line,
            space_after=12,
            keep_with_next=False,
            keep_together=True,
        )
        _disable_contextual_spacing(toc)
        toc.paragraph_format.tab_stops.clear_all()
        if level > 1:
            toc.paragraph_format.tab_stops.add_tab_stop(
                Inches(left), WD_TAB_ALIGNMENT.LEFT
            )
        toc.paragraph_format.tab_stops.add_tab_stop(
            Inches(text_width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )

    toa_title = _paragraph_style(styles, TOA_TITLE_STYLE, section_heading)
    _font(
        toa_title,
        resolved.font_family,
        resolved.body_font_size_pt,
        bold=True,
        underline=True,
    )
    _paragraph(
        toa_title,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
        space_before=12,
        space_after=12,
        keep_with_next=True,
        keep_together=True,
        page_break_before=True,
    )
    _outline_level(toa_title, 9)

    native_toa_heading = _paragraph_style(styles, "TOA Heading", body_no_indent)
    _font(
        native_toa_heading,
        resolved.font_family,
        resolved.body_font_size_pt,
        small_caps=True,
    )
    _paragraph(
        native_toa_heading,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        space_before=12,
        space_after=6,
        keep_with_next=True,
        keep_together=True,
    )
    _outline_level(native_toa_heading, 9)

    toa_category = _paragraph_style(styles, TOA_CATEGORY_STYLE, native_toa_heading)
    _font(
        toa_category,
        resolved.font_family,
        resolved.body_font_size_pt,
        small_caps=True,
    )
    _paragraph(
        toa_category,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        space_before=12,
        space_after=6,
        keep_with_next=True,
        keep_together=True,
    )
    _outline_level(toa_category, 9)

    toa = _paragraph_style(styles, "Table of Authorities", body_no_indent)
    _font(toa, resolved.font_family, resolved.body_font_size_pt)
    _paragraph(
        toa,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        left=0.5,
        first_line=-0.5,
        space_after=12,
        keep_with_next=False,
        keep_together=True,
    )
    _disable_contextual_spacing(toa)
    toa.paragraph_format.tab_stops.clear_all()
    toa.paragraph_format.tab_stops.add_tab_stop(
        Inches(text_width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )

    footer = _paragraph_style(styles, FOOTER_STYLE, body_no_indent)
    _font(footer, resolved.font_family, min(resolved.body_font_size_pt, 12))
    _paragraph(
        footer,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
        keep_with_next=False,
        keep_together=False,
    )

    case_name = _character_style(styles, CASE_NAME_STYLE)
    _font(
        case_name,
        resolved.font_family,
        resolved.body_font_size_pt,
        italic=True,
    )

    hyperlink = _character_style(styles, HYPERLINK_STYLE)
    _font(
        hyperlink,
        resolved.font_family,
        resolved.body_font_size_pt,
        color=resolved.hyperlink_color,
        underline=resolved.hyperlink_underline,
    )


def _font(
    style: ParagraphStyle | CharacterStyle,
    family: str,
    size_pt: float,
    *,
    bold: bool = False,
    italic: bool = False,
    small_caps: bool = False,
    color: str | None = None,
    underline: bool = False,
) -> None:
    style.font.name = family
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.italic = italic
    style.font.small_caps = small_caps
    style.font.underline = underline
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), family)


def _paragraph(
    style: ParagraphStyle,
    *,
    alignment: WD_ALIGN_PARAGRAPH,
    line_spacing: float,
    first_line: float | None = None,
    left: float | None = None,
    right: float | None = None,
    space_before: float = 0,
    space_after: float = 0,
    keep_with_next: bool,
    keep_together: bool,
    page_break_before: bool = False,
) -> None:
    formatting = style.paragraph_format
    formatting.alignment = alignment
    formatting.line_spacing = line_spacing
    formatting.space_before = Pt(space_before)
    formatting.space_after = Pt(space_after)
    formatting.first_line_indent = (
        Inches(first_line) if first_line is not None else None
    )
    formatting.left_indent = Inches(left) if left is not None else None
    formatting.right_indent = Inches(right) if right is not None else None
    formatting.keep_with_next = keep_with_next
    formatting.keep_together = keep_together
    formatting.page_break_before = page_break_before
    formatting.widow_control = True


def _paragraph_style(styles: object, name: str, base: ParagraphStyle) -> ParagraphStyle:
    try:
        style = cast(ParagraphStyle, styles[name])  # type: ignore[index]
    except KeyError:
        matching_style: ParagraphStyle | None = next(
            (
                candidate
                for candidate in cast(Iterable[ParagraphStyle], styles)
                if candidate.type == WD_STYLE_TYPE.PARAGRAPH
                and candidate.name.casefold() == name.casefold()
            ),
            None,
        )
        if matching_style is None:
            matching_style = cast(
                ParagraphStyle,
                styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH),  # type: ignore[attr-defined]
            )
        style = matching_style
    style.base_style = base
    style.hidden = False
    style.quick_style = True
    return style


def _character_style(styles: object, name: str) -> CharacterStyle:
    try:
        style = cast(CharacterStyle, styles[name])  # type: ignore[index]
    except KeyError:
        style = cast(
            CharacterStyle,
            styles.add_style(name, WD_STYLE_TYPE.CHARACTER),  # type: ignore[attr-defined]
        )
    style.hidden = False
    style.quick_style = True
    return style


def _outline_level(style: ParagraphStyle, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    existing = p_pr.find(qn("w:outlineLvl"))
    if existing is not None:
        p_pr.remove(existing)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _disable_contextual_spacing(style: ParagraphStyle) -> None:
    p_pr = style.element.get_or_add_pPr()
    contextual = p_pr.find(qn("w:contextualSpacing"))
    if contextual is None:
        contextual = OxmlElement("w:contextualSpacing")
        p_pr.append(contextual)
    contextual.set(qn("w:val"), "0")


def _clear_heading_style_geometry(style: ParagraphStyle) -> None:
    p_pr = style.element.get_or_add_pPr()
    for child_name in ("w:ind", "w:tabs"):
        existing = p_pr.find(qn(child_name))
        if existing is not None:
            p_pr.remove(existing)


def _add_heading_numbering(
    document: WordDocument, numbering_grammar: HeadingNumbering
) -> int:
    if numbering_grammar not in {"point-roman", "roman-outline"}:
        raise WordTemplateBuildError("unsupported heading numbering grammar")
    numbering = document.part.numbering_part.element
    abstract_id = max(
        (
            int(element.get(qn("w:abstractNumId"), "0"))
            for element in numbering.findall(qn("w:abstractNum"))
        ),
        default=0,
    ) + 1
    num_id = max(
        (
            int(element.get(qn("w:numId"), "0"))
            for element in numbering.findall(qn("w:num"))
        ),
        default=0,
    ) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    for tag, value in (
        ("w:nsid", ARGUMENT_OUTLINE_NSID),
        ("w:multiLevelType", "hybridMultilevel"),
        ("w:tmpl", ARGUMENT_OUTLINE_TEMPLATE),
        ("w:name", ARGUMENT_OUTLINE_STYLE_NAME),
        ("w:styleLink", ARGUMENT_OUTLINE_STYLE_ID),
    ):
        child = OxmlElement(tag)
        child.set(qn("w:val"), value)
        abstract.append(child)

    for level, (_, style_id, _), number_format, level_text in zip(
        range(4),
        HEADING_STYLES,
        ("upperRoman", "upperLetter", "decimal", "lowerLetter"),
        ("%1.", "%2.", "%3.", "%4."),
        strict=True,
    ):
        abstract.append(
            _numbering_level(
                level=level,
                style_id=style_id,
                number_format=number_format,
                level_text=level_text,
                text_twips=HEADING_TEXT_TWIPS[level],
                hanging_twips=HEADING_HANGING_TWIPS,
            )
        )

    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _numbering_level(
    *,
    level: int,
    style_id: str,
    number_format: str,
    level_text: str,
    text_twips: int,
    hanging_twips: int,
) -> etree._Element:
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), str(level))

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), number_format)
    lvl.append(num_fmt)
    if level > 0:
        restart = OxmlElement("w:lvlRestart")
        restart.set(qn("w:val"), str(level))
        lvl.append(restart)
    p_style = OxmlElement("w:pStyle")
    p_style.set(qn("w:val"), style_id)
    lvl.append(p_style)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    lvl.append(suffix)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), level_text)
    lvl.append(text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    lvl.append(justification)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(text_twips))
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(text_twips))
    indent.set(qn("w:hanging"), str(hanging_twips))
    p_pr.append(indent)
    lvl.append(p_pr)
    return lvl


def _bind_heading_numbering(document: WordDocument, num_id: int) -> None:
    _bind_numbering_style(document, num_id)
    for name, style_id, level in HEADING_STYLES:
        style = cast(ParagraphStyle, document.styles[name])
        style.element.set(qn("w:styleId"), style_id)
        p_pr = style.element.get_or_add_pPr()
        existing = p_pr.find(qn("w:numPr"))
        if existing is not None:
            p_pr.remove(existing)
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        number = OxmlElement("w:numId")
        number.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(number)
        p_pr.append(num_pr)


def _bind_numbering_style(document: WordDocument, num_id: int) -> None:
    styles = document.styles
    try:
        style = styles[ARGUMENT_OUTLINE_STYLE_NAME]
    except KeyError:
        style = styles.add_style(ARGUMENT_OUTLINE_STYLE_NAME, WD_STYLE_TYPE.LIST)
    style.element.set(qn("w:styleId"), ARGUMENT_OUTLINE_STYLE_ID)
    style.hidden = False
    style.quick_style = True
    p_pr = style.element.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        style.element.append(p_pr)
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.append(number)
    p_pr.append(num_pr)


def _set_page_numbering(section: Section, *, fmt: str, start: int) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is None:
        existing = OxmlElement("w:pgNumType")
        sect_pr.append(existing)
    existing.set(qn("w:fmt"), fmt)
    existing.set(qn("w:start"), str(start))


def _clear_footer(footer: object) -> None:
    element = footer._element  # type: ignore[attr-defined]
    for child in list(element):
        element.remove(child)
    element.append(OxmlElement("w:p"))


def _set_page_number_footer(
    document: WordDocument,
    footer: object,
    resolved: ResolvedWordStyle,
    *,
    placeholder: str,
) -> None:
    _clear_footer(footer)
    width = Inches(6.5)
    table = footer.add_table(rows=1, cols=1, width=width)  # type: ignore[attr-defined]
    table.autofit = False
    row = table.rows[0]
    row.height = Inches(0.5)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    cell.width = width
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles[FOOTER_STYLE]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    _field(run._r, "PAGE", placeholder=placeholder)
    for page_run in paragraph.runs:
        page_run.font.name = resolved.font_family


def _field(
    run_element: etree._Element,
    instruction: str,
    *,
    placeholder: str = "1",
) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instruction_element.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction_element, separate, value, end):
        run_element.append(element)


def _configure_settings(document: WordDocument) -> None:
    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def _clear_core_properties(document: WordDocument) -> None:
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = ""
    properties.subject = ""
    properties.keywords = ""
    properties.comments = ""
    properties.category = ""


def _sanitize_package(path: Path, font_family: str) -> Path:
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.stem}.sanitized.", suffix=".docx", dir=path.parent
    )
    os.close(descriptor)
    output = Path(temporary_name)
    try:
        with zipfile.ZipFile(path) as incoming, zipfile.ZipFile(output, "w") as outgoing:
            for member in incoming.infolist():
                data = incoming.read(member)
                if member.filename == "docProps/core.xml":
                    data = _empty_core_properties(data)
                elif member.filename == "word/theme/theme1.xml":
                    data = _harmonize_theme(data, font_family)
                outgoing.writestr(member, data)
        return output
    except Exception:
        output.unlink(missing_ok=True)
        raise


def _empty_core_properties(data: bytes) -> bytes:
    root = etree.fromstring(data)
    for child in list(root):
        root.remove(child)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _harmonize_theme(data: bytes, font_family: str) -> bytes:
    root = etree.fromstring(data)
    elements = cast(
        list[etree._Element],
        root.xpath(
            ".//a:fontScheme/a:majorFont/a:latin | "
            ".//a:fontScheme/a:minorFont/a:latin",
            namespaces=NS,
        ),
    )
    for element in elements:
        element.set("typeface", font_family)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _write_manifest(
    path: Path,
    profile: FilingProfile,
    style_pack: WordStylePack,
    digest: str,
) -> None:
    lines = [
        "schema_version = 2",
        f'template_id = "{profile.profile_id}"',
        f'version = "{profile.version}"',
        f'template_sha256 = "{digest}"',
        f'filing_profile_id = "{profile.profile_id}"',
        f'style_pack_id = "{style_pack.style_pack_id}"',
        'slots = ["toc", "toa", "body", "signature", "certificate"]',
        "",
    ]
    for role, word_style_name, toc_level in STYLE_BINDINGS:
        lines.extend(
            [
                "[[styles]]",
                f'semantic_role = "{role}"',
                f'word_style_name = "{word_style_name}"',
            ]
        )
        if toc_level is not None:
            lines.append(f"toc_level = {toc_level}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    path.chmod(0o600)


__all__ = [
    "ARGUMENT_OUTLINE_STYLE_NAME",
    "BODY_NO_INDENT_STYLE",
    "BODY_STYLE",
    "CASE_NAME_STYLE",
    "DOCUMENT_TITLE_STYLE",
    "FOOTNOTE_STYLE",
    "HEADING_STYLES",
    "SECTION_HEADING_STYLE",
    "SIGNATURE_LAST_STYLE",
    "SIGNATURE_STYLE",
    "STYLE_BINDINGS",
    "TOA_CATEGORY_STYLE",
    "TOA_TITLE_STYLE",
    "TOC_TITLE_STYLE",
    "TemplateBuildResult",
    "WordTemplateBuildError",
    "build_blank_word_template",
    "configure_front_matter_section",
    "configure_main_body_section",
]
