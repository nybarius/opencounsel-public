from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import tempfile
from dataclasses import asdict
from importlib import resources
from pathlib import Path
from typing import cast

from docx import Document
from docx.enum.section import WD_SECTION
from opencounsel.templates.profiles import (
    FilingProfile,
    get_bundled_filing_profile,
    list_bundled_filing_profiles,
)

from opencounsel.templates.audit import audit_word_styles, write_style_audit
from opencounsel.templates.build import (
    BODY_NO_INDENT_STYLE,
    BODY_STYLE,
    CASE_NAME_STYLE,
    DOCUMENT_TITLE_STYLE,
    FOOTNOTE_STYLE,
    HEADING_STYLES,
    SECTION_HEADING_STYLE,
    SIGNATURE_LAST_STYLE,
    SIGNATURE_STYLE,
    TOA_CATEGORY_STYLE,
    TOA_TITLE_STYLE,
    TOC_TITLE_STYLE,
    build_blank_word_template,
    configure_main_body_section,
)

SAMPLE_PROFILES = (
    "ny-ad-appellant-brief",
    "fed-sdny-edny-motion-memorandum",
)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    build_review_packet(args.out)
    return 0


def build_review_packet(root: Path) -> None:
    if root.exists() or root.is_symlink():
        raise ValueError("review output must be a new directory")
    blank_dir = root / "blank-templates"
    sample_dir = root / "populated-samples"
    audit_dir = root / "style-audits"
    rule_dir = root / "filing-profiles"
    pdf_dir = root / "pdf-previews"
    for directory in (blank_dir, sample_dir, audit_dir, rule_dir, pdf_dir):
        directory.mkdir(parents=True, mode=0o700)

    profile_package = resources.files("opencounsel.templates.profile_data")
    for resource in profile_package.iterdir():
        if resource.name.endswith(".toml"):
            (rule_dir / resource.name).write_text(
                resource.read_text(encoding="utf-8"),
                encoding="utf-8",
            )

    repository_root = Path(__file__).resolve().parent.parent
    (root / "STYLE-GUIDE.md").write_text(
        (repository_root / "docs" / "WORD_STYLES.md").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (root / "README.md").write_text(_review_readme(), encoding="utf-8")

    index: dict[str, object] = {
        "schema_version": 1,
        "notice": "All content is synthetic and intended only for style review.",
        "profiles": [],
        "samples": [],
    }
    for profile in list_bundled_filing_profiles():
        template = blank_dir / f"{profile.profile_id}.docx"
        manifest = blank_dir / f"{profile.profile_id}.template.toml"
        result = build_blank_word_template(profile, template, manifest_path=manifest)
        audit_path = audit_dir / f"{profile.profile_id}.blank.audit.json"
        audit = audit_word_styles(template, profile)
        write_style_audit(audit, audit_path)
        if any(finding.severity == "error" for finding in audit.findings):
            raise RuntimeError(
                f"generated template failed style audit: {profile.profile_id}"
            )
        profiles = cast(list[object], index["profiles"])
        profiles.append(
            {
                "profile_id": profile.profile_id,
                "court": profile.court,
                "document_type": profile.document_type,
                "heading_numbering": profile.structure.heading_numbering,
                "body_pt": profile.typography.body_min_pt,
                "footnote_pt": profile.typography.footnote_min_pt,
                "template": str(template.relative_to(root)),
                "manifest": str(manifest.relative_to(root)),
                "audit": str(audit_path.relative_to(root)),
                "template_sha256": result.template_sha256,
            }
        )

    for profile_id in SAMPLE_PROFILES:
        profile = get_bundled_filing_profile(profile_id)
        base = sample_dir / f".{profile_id}.base.docx"
        build_blank_word_template(profile, base)
        sample = sample_dir / f"{profile_id}.style-sample.docx"
        _populate_sample(base, sample, profile)
        base.unlink()
        audit_path = audit_dir / f"{profile_id}.sample.audit.json"
        audit = audit_word_styles(sample, profile)
        write_style_audit(audit, audit_path)
        if any(finding.severity == "error" for finding in audit.findings):
            raise RuntimeError(
                f"populated sample failed style audit: {profile_id}"
            )
        pdf = _render_pdf(sample, pdf_dir)
        samples = cast(list[object], index["samples"])
        samples.append(
            {
                "profile_id": profile_id,
                "document": str(sample.relative_to(root)),
                "pdf": str(pdf.relative_to(root)) if pdf is not None else None,
                "audit": str(audit_path.relative_to(root)),
                "findings": [asdict(finding) for finding in audit.findings],
            }
        )

    (root / "review-index.json").write_text(
        json.dumps(index, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _review_readme() -> str:
    return """# OpenCounsel legal-style review packet

Everything in this packet is synthetic. It contains no client text, filename,
metadata, hash, or matter fact.

Open the populated DOCX samples in desktop Microsoft Word and update fields before
review. PDF previews are emitted when LibreOffice is available and validate the
publication layout, not every Microsoft Word client quirk. Review the generic legal
styles, justified four-level argument outline, five-level TOC, categorized TOA,
front-matter page break, merits section break, pagination, body text, and signature.
"""


def _populate_sample(
    source: Path, output: Path, profile: FilingProfile
) -> None:
    document = Document(str(source))
    first = document.paragraphs[0]
    first.style = document.styles[DOCUMENT_TITLE_STYLE]
    first.add_run("SYNTHETIC WORD-STYLE REVIEW - NOT FOR FILING")

    _paragraph(document, TOC_TITLE_STYLE, "TABLE OF CONTENTS")
    _toc_entry(document, "TOC 1", None, "PRELIMINARY STATEMENT", "1")
    _toc_entry(document, "TOC 1", None, "ARGUMENT", "2")
    _toc_entry(
        document,
        "TOC 2",
        "I.",
        "A Long Point Heading Wraps Under Its Text",
        "2",
    )
    _toc_entry(document, "TOC 3", "A.", "The Covenant Bars the Action", "3")
    _toc_entry(document, "TOC 4", "1.", "The Text Controls", "4")
    _toc_entry(document, "TOC 5", "a.", "The Record Confirms the Text", "4")
    _toc_entry(
        document,
        "TOC 3",
        "B.",
        "The Release Independently Applies",
        "4",
    )
    _toc_entry(
        document,
        "TOC 2",
        "II.",
        "The Remaining Objections Fail",
        "5",
    )
    _toc_entry(document, "TOC 1", None, "CONCLUSION", "5")

    _paragraph(document, TOA_TITLE_STYLE, "TABLE OF AUTHORITIES")
    _paragraph(document, TOA_CATEGORY_STYLE, "Cases")
    case = document.add_paragraph(style="Table of Authorities")
    case_name = case.add_run("Synthetic Corp. v. Example Co.")
    case_name.style = document.styles[CASE_NAME_STYLE]
    case.add_run(", 123 F.4th 456 (2d Cir. 2026)\t2, 4")
    _paragraph(document, TOA_CATEGORY_STYLE, "Statutes")
    _table_entry(document, "Table of Authorities", "28 U.S.C. § 1332", "3")
    _paragraph(document, TOA_CATEGORY_STYLE, "Rules")
    _table_entry(
        document,
        "Table of Authorities",
        "Fed. R. Civ. P. 12(b)(6)",
        "3",
    )
    _paragraph(document, TOA_CATEGORY_STYLE, "Regulations")
    _table_entry(
        document,
        "Table of Authorities",
        "17 C.F.R. § 240.10b-5",
        "4",
    )

    main = document.add_section(WD_SECTION.NEW_PAGE)
    configure_main_body_section(document, main, profile)
    _paragraph(document, SECTION_HEADING_STYLE, "PRELIMINARY STATEMENT")
    _body(document, profile, 4)
    _paragraph(
        document,
        FOOTNOTE_STYLE,
        "Footnote specimen: this text is justified, single-spaced, and sized "
        "by the filing profile.",
    )

    _paragraph(document, SECTION_HEADING_STYLE, "ARGUMENT")
    _paragraph(
        document,
        HEADING_STYLES[0][0],
        "THE ORDER SHOULD BE REVERSED BECAUSE THIS LONG POINT HEADING WRAPS "
        "WITH EVERY CONTINUATION LINE FLUSH AT THE TEXT POSITION",
    )
    _body(document, profile, 2)
    _paragraph(document, HEADING_STYLES[1][0], "The Covenant Bars the Action")
    _body(document, profile, 1)
    _paragraph(document, HEADING_STYLES[2][0], "The Agreement's Text Controls")
    _body(document, profile, 1)
    _paragraph(document, HEADING_STYLES[3][0], "The Record Confirms the Text")
    _body(document, profile, 1)
    _paragraph(
        document,
        HEADING_STYLES[1][0],
        "The Release Independently Applies",
    )
    _body(document, profile, 1)
    _paragraph(
        document,
        HEADING_STYLES[0][0],
        "THE REMAINING OBJECTIONS PROVIDE NO BASIS TO AFFIRM",
    )
    _body(document, profile, 1)

    _paragraph(document, SECTION_HEADING_STYLE, "CONCLUSION")
    _paragraph(
        document,
        BODY_NO_INDENT_STYLE,
        "The requested disposition would appear here.",
    )
    _paragraph(document, SIGNATURE_STYLE, "Respectfully submitted,")
    _paragraph(document, SIGNATURE_STYLE, "Example Counsel")
    _paragraph(document, SIGNATURE_LAST_STYLE, "Attorney for the Example Party")

    document.save(str(output))
    output.chmod(0o600)


def _body(document: object, profile: FilingProfile, count: int) -> None:
    text = (
        "This synthetic paragraph demonstrates the resolved body style for "
        f"{profile.court}. It uses the profile font size, double spacing, "
        "first-line indent, widow control, and justified alignment. No client "
        "fact or legal proposition appears here."
    )
    for _ in range(count):
        _paragraph(document, BODY_STYLE, text)


def _paragraph(document: object, style: str, text: str) -> object:
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def _toc_entry(
    document: object,
    style: str,
    prefix: str | None,
    text: str,
    page: str,
) -> None:
    paragraph = document.add_paragraph(style=style)
    if prefix is not None:
        paragraph.add_run(prefix)
        paragraph.add_run("\t")
    paragraph.add_run(text)
    paragraph.add_run("\t")
    paragraph.add_run(page)


def _table_entry(document: object, style: str, label: str, page: str) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(label)
    paragraph.add_run("\t")
    paragraph.add_run(page)


def _render_pdf(source: Path, output_dir: Path) -> Path | None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        return None
    with tempfile.TemporaryDirectory(prefix="opencounsel-style-review-lo-") as temporary:
        profile = Path(temporary) / "profile"
        subprocess.run(
            [
                executable,
                "--headless",
                "--nologo",
                "--nodefault",
                "--norestore",
                f"-env:UserInstallation={profile.resolve().as_uri()}",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                str(output_dir),
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=120,
        )
    output = output_dir / source.with_suffix(".pdf").name
    if not output.is_file():
        raise RuntimeError(
            f"LibreOffice did not create {output.name}"
        )
    output.chmod(0o600)
    return output


if __name__ == "__main__":
    raise SystemExit(main())
