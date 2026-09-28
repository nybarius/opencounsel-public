from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def main() -> int:
    if len(sys.argv) != 2:
        raise SystemExit("usage: make_seed.py CORPUS_DIRECTORY")
    output = Path(sys.argv[1])
    output.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("[TOC]")
    document.add_paragraph("[TOA]")
    document.add_heading("SYNTHETIC ARGUMENT", level=1)
    document.add_paragraph(
        "Synthetic citation , 550 U.S. 544 (2007), 42 U.S.C. § 1983, and R. 3."
    )
    document.save(output / "minimal.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
