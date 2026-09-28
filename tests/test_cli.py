from __future__ import annotations

import json
import shutil
from pathlib import Path

from sqlalchemy import create_engine
from sqlalchemy.exc import OperationalError

from opencounsel.cli import main
from opencounsel.config import Settings
from opencounsel.demo import build_synthetic_demo
from opencounsel.persistence.models import Base
from opencounsel.source.roa import inspect_roa_package


def _brief(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph("Twombly, 550 U.S. 544 (2007), confirms the rule (R 3).")
    document.save(path)
    return path


def test_init_and_inspect_cli(roa_package: Path, tmp_path: Path, capsys, monkeypatch) -> None:
    monkeypatch.setattr("opencounsel.config.DEFAULT_OBJECT_ROOT", tmp_path / "objects")
    config = tmp_path / "config.toml"
    assert main(["--config", str(config), "init"]) == 0
    assert json.loads(capsys.readouterr().out)["model_policy"] == "disabled"

    assert main(["inspect-roa", str(roa_package)]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["page_count"] == 3
    assert result["record_page_range"] == [3, 5]


def test_cli_reports_validation_error_without_traceback(tmp_path: Path, capsys) -> None:
    bad = tmp_path / "bad.zip"
    bad.write_bytes(b"not a zip")
    assert main(["inspect-roa", str(bad)]) == 2
    result = json.loads(capsys.readouterr().err)
    assert result == {"status": "error", "error": "package is not a valid ZIP archive"}


def test_package_roa_cli_exposes_machine_readable_converter(
    tmp_path: Path, capsys
) -> None:
    demo = tmp_path / "demo"
    build_synthetic_demo(demo)
    output = tmp_path / "machine-readable-roa.zip"

    assert (
        main(
            [
                "package-roa",
                str(demo / "record.pdf"),
                "--out",
                str(output),
                "--first-record-page",
                "1",
                "--numbering-verified",
            ]
        )
        == 0
    )
    result = json.loads(capsys.readouterr().out)
    assert result["status"] == "packaged"
    assert result["page_count"] == 3
    assert result["record_page_range"] == [1, 3]
    assert result["package_sha256"] == inspect_roa_package(output).package.sha256


def test_audit_brief_cli(roa_package: Path, tmp_path: Path, capsys) -> None:
    brief = _brief(tmp_path / "brief.docx")
    output = tmp_path / "audit"

    assert (
        main(
            [
                "audit-brief",
                str(brief),
                "--roa-package",
                str(roa_package),
                "--record-pdf",
                str(tmp_path / "record.pdf"),
                "--out",
                str(output),
            ]
        )
        == 0
    )
    result = json.loads(capsys.readouterr().out)
    assert result["status"] == "audited"
    assert result["record_citation_count"] == 1
    assert (output / "record-citations.csv").is_file()


def test_import_and_resolve_cli(roa_package: Path, tmp_path: Path, capsys, monkeypatch) -> None:
    database_url = f"sqlite+pysqlite:///{tmp_path / 'test.db'}"
    Base.metadata.create_all(create_engine(database_url))
    settings = Settings(database_url, tmp_path / "objects")
    monkeypatch.setattr("opencounsel.cli.load_settings", lambda _path: settings)

    assert (
        main(
            [
                "import-roa",
                str(roa_package),
                "--matter",
                "matter-001",
                "--document",
                "record",
                "--expect",
                "3:5",
            ]
        )
        == 0
    )
    assert json.loads(capsys.readouterr().out)["page_count"] == 3
    assert (
        main(
            [
                "resolve",
                "R 3",
                "--matter",
                "matter-001",
                "--document",
                "record",
            ]
        )
        == 0
    )
    assert json.loads(capsys.readouterr().out)["physical_page"] == 1


def test_cli_redacts_database_error(tmp_path: Path, capsys, monkeypatch) -> None:
    settings = Settings("postgresql+psycopg://secret:secret@invalid/db", tmp_path / "objects")
    monkeypatch.setattr("opencounsel.cli.load_settings", lambda _path: settings)
    monkeypatch.setattr(
        "opencounsel.cli.make_engine",
        lambda _url: (_ for _ in ()).throw(OperationalError("statement", {}, Exception("secret"))),
    )
    assert main(["resolve", "R 3", "--matter", "matter-001", "--document", "record"]) == 2
    assert json.loads(capsys.readouterr().err) == {
        "status": "error",
        "error": "database operation failed",
    }


def test_authority_acquisition_package_and_verification_cli(tmp_path: Path, capsys) -> None:
    brief = tmp_path / "authorities.docx"
    from docx import Document

    document = Document()
    document.add_paragraph(
        'Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234, held that "the agreement controls."'
    )
    document.save(brief)
    bundle = tmp_path / "bundle"

    assert main(["prepare-authorities", str(brief), "--out", str(bundle)]) == 0
    prepared = json.loads(capsys.readouterr().out)
    assert prepared["status"] == "prepared"
    assert prepared["detected_authority_count"] == 1
    assert prepared["authority_count"] == 1
    assert prepared["source_copy_required_count"] == 1
    assert prepared["case_source_copy_required_count"] == 1
    assert prepared["source_needed_count"] == 1

    package = tmp_path / "authorities.zip"
    assert main(["package-authorities", str(bundle), "--out", str(package)]) == 0
    packaged = json.loads(capsys.readouterr().out)
    assert packaged["status"] == "packaged"
    assert packaged["source_count"] == 0

    report = tmp_path / "verification.json"
    assert main(["verify-authorities", str(package), "--out", str(report)]) == 0
    verified = json.loads(capsys.readouterr().out)
    assert verified["status"] == "verified"
    assert verified["source_present_count"] == 0
    assert verified["citation_confirmed_count"] == 0


def test_process_brief_cli_emits_new_docx_and_separate_ledger(
    tmp_path: Path, capsys
) -> None:
    brief = _brief(tmp_path / "brief.docx")
    original = brief.read_bytes()
    work = tmp_path / "work"

    assert main(["process-brief", str(brief), "--work-root", str(work)]) == 0

    result = json.loads(capsys.readouterr().out)
    corrected = Path(result["corrected_document"])
    ledger = Path(result["correction_ledger"])
    front_matter = Path(result["front_matter_source"])
    assert result["status"] == "completed"
    assert result["correction_count"] == 2
    assert result["applied_correction_count"] == 0
    assert result["review_item_count"] == 2
    assert result["record_citation_count"] == 1
    assert result["authority_citation_count"] == 1
    assert result["citation_review_item_count"] == 2
    assert result["hyperlink_candidate_count"] == 0
    assert result["hyperlink_inserted_count"] == 0
    assert result["hyperlink_review_item_count"] == 0
    assert result["toc_entry_count"] == 1
    assert result["toa_authority_count"] == 1
    assert result["toa_occurrence_count"] == 1
    assert result["front_matter_field_candidate_count"] == 0
    assert result["lawyer_review_status"] == "pending"
    assert corrected.name == "corrected.docx"
    assert ledger.name == "corrections.json"
    assert front_matter.name == "front-matter.json"
    assert corrected.read_bytes() == original
    assert brief.read_bytes() == original
    assert json.loads(front_matter.read_text(encoding="utf-8"))["brief_sha256"] == (
        result["hyperlinked_sha256"]
    )
    assert [
        item["item_type"]
        for item in json.loads(ledger.read_text(encoding="utf-8"))["corrections"]
    ] == ["record-citation", "authority-citation"]


def test_clean_brief_cli_runs_profile_roa_links_and_front_matter_together(
    tmp_path: Path,
    roa_package: Path,
    capsys,
) -> None:
    brief = _brief(tmp_path / "brief.docx")
    work = tmp_path / "work"

    assert (
        main(
            [
                "clean-brief",
                str(brief),
                "--profile",
                "ny-ad-appellant-brief",
                "--roa-package",
                str(roa_package),
                "--work-root",
                str(work),
            ]
        )
        == 0
    )

    result = json.loads(capsys.readouterr().out)
    process = result["process"]
    assert result["status"] == "completed"
    assert process["resolved_record_citation_count"] == 1
    assert process["unresolved_record_citation_count"] == 0
    assert process["toc_field_inserted_count"] == 1
    assert process["toa_field_inserted_count"] == 1
    assert Path(result["corrected_document"]).is_file()
    assert Path(result["correction_ledger"]).is_file()
    assert Path(result["front_matter_source"]).is_file()

    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        return
    assert (
        main(
            [
                "publish-foss",
                process["process_id"],
                "--work-root",
                str(work),
                "--libreoffice",
                executable,
            ]
        )
        == 0
    )
    publication = json.loads(capsys.readouterr().out)
    assert Path(publication["published_document"]).is_file()
    assert Path(publication["published_pdf"]).is_file()
    assert Path(publication["publication_manifest"]).is_file()


def test_filing_profile_template_and_style_audit_cli(tmp_path: Path, capsys) -> None:
    assert main(["list-filing-profiles"]) == 0
    profiles = json.loads(capsys.readouterr().out)
    assert {item["profile_id"] for item in profiles["profiles"]} >= {
        "ny-ad-appellant-brief",
        "fed-sdny-edny-motion-memorandum",
    }

    output = tmp_path / "appellate.docx"
    manifest = tmp_path / "appellate.toml"
    assert (
        main(
            [
                "build-template",
                "ny-ad-appellant-brief",
                str(output),
                "--manifest",
                str(manifest),
            ]
        )
        == 0
    )
    built = json.loads(capsys.readouterr().out)
    assert built["status"] == "built"
    assert output.is_file()
    assert manifest.is_file()

    audit = tmp_path / "audit.json"
    assert (
        main(
            [
                "audit-word-styles",
                str(output),
                "--profile",
                "ny-ad-appellant-brief",
                "--out",
                str(audit),
            ]
        )
        == 0
    )
    result = json.loads(capsys.readouterr().out)
    assert result["status"] == "audited"
    assert result["error_count"] == 0
    assert audit.is_file()


def test_ui_cli_starts_private_local_server(tmp_path: Path, monkeypatch) -> None:
    called = {}

    def fake_run(**options) -> None:
        called.update(options)

    monkeypatch.setattr("opencounsel.web.server.run", fake_run)
    assert (
        main(
            [
                "ui",
                "--host",
                "127.0.0.1",
                "--port",
                "8877",
                "--work-root",
                str(tmp_path / "web"),
                "--libreoffice",
                "soffice",
            ]
        )
        == 0
    )
    assert called == {
        "host": "127.0.0.1",
        "port": 8877,
        "root": tmp_path / "web",
        "libreoffice": "soffice",
    }
