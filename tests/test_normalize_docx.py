from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from lxml import etree

from opencounsel.briefs.docx import inspect_brief_docx
from opencounsel.briefs.normalize_docx import (
    FormattingProjectionError,
    normalize_brief_formatting,
)
from opencounsel.templates.build import (
    BODY_STYLE,
    HEADING_STYLES,
    SECTION_HEADING_STYLE,
    TOA_TITLE_STYLE,
    TOC_TITLE_STYLE,
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def _legacy_brief(path: Path) -> Path:
    document = Document()
    title = document.add_paragraph("APPELLANT'S BRIEF")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("PRELIMINARY STATEMENT")
    document.add_paragraph("A synthetic introductory paragraph.")
    standard = document.add_paragraph("STANDARD OF REVIEW")
    standard.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading = document.add_paragraph("I. THE ORDER SHOULD BE REVERSED")
    heading.runs[0].bold = True
    document.add_paragraph("The claim arises under 42 U.S.C. § 1983.")
    document.add_paragraph("1. This numbered statement remains body text.")
    document.add_paragraph("CONCLUSION")
    document.add_paragraph("The requested relief should be granted.")
    document.save(path)
    return path


def test_normalizes_conventional_brief_and_adds_front_matter(tmp_path: Path) -> None:
    source = _legacy_brief(tmp_path / "legacy.docx")
    output = tmp_path / "clean.docx"

    result = normalize_brief_formatting(source, output, "ny-ad-appellant-brief")

    assert result.profile_id == "ny-ad-appellant-brief"
    assert result.inserted_toc_slot is True
    assert result.inserted_toa_slot is True
    assert result.inserted_merits_section is True
    assert result.input_sha256 != result.output_sha256

    document = Document(output)
    by_text = {paragraph.text: paragraph for paragraph in document.paragraphs}
    assert by_text["TABLE OF CONTENTS"].style.name == TOC_TITLE_STYLE
    assert by_text["TABLE OF AUTHORITIES"].style.name == TOA_TITLE_STYLE
    assert by_text["PRELIMINARY STATEMENT"].style.name == SECTION_HEADING_STYLE
    assert by_text["STANDARD OF REVIEW"].style.name == SECTION_HEADING_STYLE
    assert by_text["THE ORDER SHOULD BE REVERSED"].style.name == HEADING_STYLES[0][0]
    assert by_text["The claim arises under 42 U.S.C. § 1983."].style.name == BODY_STYLE
    assert by_text["1. This numbered statement remains body text."].style.name == BODY_STYLE
    assert "[TOC]" in by_text
    assert "[TOA]" in by_text
    assert len(document.sections) == 2
    assert [(heading.level, heading.text) for heading in inspect_brief_docx(output).headings] == [
        (1, "PRELIMINARY STATEMENT"),
        (1, "STANDARD OF REVIEW"),
        (2, "THE ORDER SHOULD BE REVERSED"),
        (1, "CONCLUSION"),
    ]

    with zipfile.ZipFile(output) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
        numbering = etree.fromstring(archive.read("word/numbering.xml"))
        document_xml = etree.fromstring(archive.read("word/document.xml"))
    assert styles.xpath(
        './/w:style[@w:styleId="LegalHeadingNum1"]', namespaces=NS
    )
    assert numbering.xpath(
        './/w:lvl[w:pStyle/@w:val="LegalHeadingNum1"]', namespaces=NS
    )
    assert len(document_xml.xpath(".//w:sectPr", namespaces=NS)) == 2


def test_starts_federal_memorandum_merits_at_introduction(tmp_path: Path) -> None:
    source = tmp_path / "memorandum.docx"
    document = Document()
    document.add_paragraph("TABLE OF CONTENTS")
    document.add_paragraph("[TOC]")
    document.add_paragraph("TABLE OF AUTHORITIES")
    document.add_paragraph("[TOA]")
    document.add_paragraph("I. INTRODUCTION")
    document.add_paragraph("Synthetic introductory text.")
    document.add_paragraph("II. BACKGROUND")
    document.add_paragraph("Synthetic background text.")
    document.add_paragraph("III. ARGUMENT")
    document.add_paragraph("Synthetic argument text.")
    document.save(source)

    output = tmp_path / "clean.docx"
    result = normalize_brief_formatting(
        source,
        output,
        "fed-sdny-edny-motion-memorandum",
    )

    assert result.inserted_merits_section is True
    with zipfile.ZipFile(output) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
    introduction_nodes = document_xml.xpath(
        './/w:p[.//w:t[contains(., "INTRODUCTION")]]',
        namespaces=NS,
    )
    assert len(introduction_nodes) == 1
    previous = introduction_nodes[0].getprevious()
    assert previous is not None
    assert previous.xpath('./w:pPr/w:sectPr', namespaces=NS)


def test_rejects_tracked_changes_before_mutation(tmp_path: Path) -> None:
    source = _legacy_brief(tmp_path / "legacy.docx")
    tracked = tmp_path / "tracked.docx"
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(tracked, "w") as outgoing:
        for member in incoming.infolist():
            data = incoming.read(member)
            if member.filename == "word/document.xml":
                data = data.replace(b"<w:body>", b"<w:body><w:ins w:id=\"1\"/>", 1)
            outgoing.writestr(member, data)

    with pytest.raises(FormattingProjectionError, match="tracked changes"):
        normalize_brief_formatting(
            tracked,
            tmp_path / "clean.docx",
            "ny-ad-appellant-brief",
        )


def test_places_front_matter_before_singular_question_presented(tmp_path: Path) -> None:
    source = tmp_path / "singular-question.docx"
    document = Document()
    document.add_paragraph("QUESTION PRESENTED")
    document.add_paragraph("1. Whether the order should be reversed.")
    document.add_paragraph("PRELIMINARY STATEMENT")
    document.add_paragraph("Synthetic body text.")
    document.save(source)

    output = tmp_path / "clean.docx"
    normalize_brief_formatting(source, output, "ny-ad-appellant-brief")

    cleaned = Document(output)
    texts = [paragraph.text for paragraph in cleaned.paragraphs]
    question_index = texts.index("QUESTION PRESENTED")
    assert texts.index("TABLE OF CONTENTS") < question_index
    assert texts.index("TABLE OF AUTHORITIES") < question_index
    assert cleaned.paragraphs[question_index].style.name == SECTION_HEADING_STYLE
    assert len(cleaned.sections) == 2

    with zipfile.ZipFile(output) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
    question_nodes = document_xml.xpath(
        './/w:p[.//w:t[contains(., "QUESTION PRESENTED")]]', namespaces=NS
    )
    assert len(question_nodes) == 1
    previous = question_nodes[0].getprevious()
    assert previous is not None
    assert previous.xpath('./w:pPr/w:sectPr', namespaces=NS)


def test_converts_standalone_page_break_to_stable_target_formatting(
    tmp_path: Path,
) -> None:
    source = tmp_path / "standalone-page-break.docx"
    document = Document()
    document.add_paragraph("PRELIMINARY STATEMENT")
    document.add_paragraph("Synthetic body text.")
    break_paragraph = document.add_paragraph()
    break_paragraph.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("PRINTING SPECIFICATIONS STATEMENT")
    document.save(source)

    output = tmp_path / "clean.docx"
    result = normalize_brief_formatting(source, output, "ny-ad-appellant-brief")

    assert any("standalone page-break" in item.action for item in result.decisions)
    with zipfile.ZipFile(output) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
    assert not document_xml.xpath('.//w:p[not(.//w:t)]//w:br[@w:type="page"]', namespaces=NS)
    target = document_xml.xpath(
        './/w:p[.//w:t[contains(., "PRINTING SPECIFICATIONS STATEMENT")]]',
        namespaces=NS,
    )
    assert len(target) == 1
    assert target[0].xpath('./w:pPr/w:pageBreakBefore', namespaces=NS)


def test_reuses_case_insensitive_native_toc_styles_without_duplicate_ids(
    tmp_path: Path,
) -> None:
    source = tmp_path / "native-toc-styles.docx"
    document = Document()
    for level in range(1, 4):
        style = document.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
        style.element.set(qn("w:styleId"), f"TOC{level}")
    document.add_paragraph("PRELIMINARY STATEMENT")
    document.add_paragraph("Synthetic body text.")
    document.save(source)

    output = tmp_path / "clean.docx"
    normalize_brief_formatting(source, output, "ny-ad-appellant-brief")

    with zipfile.ZipFile(output) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
    for level in range(1, 4):
        assert len(
            styles.xpath(
                f'.//w:style[@w:styleId="TOC{level}"]',
                namespaces=NS,
            )
        ) == 1
    toc3 = styles.xpath('.//w:style[@w:styleId="TOC3"]', namespaces=NS)[0]
    indent = toc3.xpath('./w:pPr/w:ind', namespaces=NS)[0]
    assert indent.get(qn("w:left")) == "1440"
    assert indent.get(qn("w:hanging")) == "720"
    tabs = {
        (tab.get(qn("w:val")), tab.get(qn("w:pos")))
        for tab in toc3.xpath('./w:pPr/w:tabs/w:tab', namespaces=NS)
    }
    assert ("left", "1440") in tabs
    assert ("right", "9360") in tabs


def test_centers_major_sections_and_rebases_argument_headings(tmp_path: Path) -> None:
    source = tmp_path / "federal-outline.docx"
    document = Document()
    for text in (
        "I. INTRODUCTION",
        "II. BACKGROUND",
        "III. LEGAL STANDARD",
        "IV. ARGUMENT",
        "A. Limitations",
        "1. Training-period conduct",
        "a. Earlier models",
        "V. CONCLUSION",
    ):
        paragraph = document.add_paragraph(text)
        paragraph.runs[0].bold = True
    document.save(source)

    output = tmp_path / "clean.docx"
    normalize_brief_formatting(
        source,
        output,
        "fed-sdny-edny-motion-memorandum",
    )

    cleaned = Document(output)
    by_text = {paragraph.text: paragraph for paragraph in cleaned.paragraphs}
    for text in ("INTRODUCTION", "BACKGROUND", "LEGAL STANDARD", "ARGUMENT", "CONCLUSION"):
        paragraph = by_text[text]
        assert paragraph.style.name == SECTION_HEADING_STYLE
        assert paragraph.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert paragraph.style.font.bold is True
        assert paragraph.style.font.underline is True

    assert by_text["Limitations"].style.name == HEADING_STYLES[0][0]
    assert by_text["Training-period conduct"].style.name == HEADING_STYLES[1][0]
    assert by_text["Earlier models"].style.name == HEADING_STYLES[2][0]

    with zipfile.ZipFile(output) as archive:
        numbering = etree.fromstring(archive.read("word/numbering.xml"))
    first_level = numbering.xpath(
        './/w:lvl[w:pStyle/@w:val="LegalHeadingNum1"]', namespaces=NS
    )[0]
    second_level = numbering.xpath(
        './/w:lvl[w:pStyle/@w:val="LegalHeadingNum2"]', namespaces=NS
    )[0]
    assert first_level.xpath('./w:pPr/w:tabs/w:tab/@w:pos', namespaces=NS) == ["720"]
    assert first_level.xpath('./w:pPr/w:ind/@w:left', namespaces=NS) == ["720"]
    assert first_level.xpath('./w:pPr/w:ind/@w:hanging', namespaces=NS) == ["720"]
    assert second_level.xpath('./w:pPr/w:tabs/w:tab/@w:pos', namespaces=NS) == ["1440"]
    assert second_level.xpath('./w:pPr/w:ind/@w:left', namespaces=NS) == ["1440"]
    assert second_level.xpath('./w:pPr/w:ind/@w:hanging', namespaces=NS) == ["720"]
