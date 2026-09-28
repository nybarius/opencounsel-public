from __future__ import annotations

import asyncio
import json
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from opencounsel.mcp.server import TOOL_NAMES, build_server, main

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _brief(path: Path) -> Path:
    document = Document()
    document.add_heading("QUESTION PRESENTED", level=1)
    document.add_paragraph("Synthetic question.")
    document.save(path)
    return path


def test_mcp_tools_have_closed_schemas_and_safety_annotations(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    server = build_server(tmp_path / "work", inbox)

    tools = asyncio.run(server.list_tools())

    assert tuple(tool.name for tool in tools) == TOOL_NAMES
    for tool in tools:
        assert tool.outputSchema is not None
        assert tool.annotations is not None
        assert tool.annotations.destructiveHint is False
        assert tool.annotations.idempotentHint is True
        assert tool.annotations.openWorldHint is False
    read_only = {tool.name: tool.annotations.readOnlyHint for tool in tools if tool.annotations}
    assert read_only == {
        "get_capabilities": True,
        "create_revision": False,
        "inspect_revision": True,
        "process_brief": False,
        "clean_brief": False,
    }


def test_mcp_create_and_inspect_return_metadata_only(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "private-name.docx")
    server = build_server(tmp_path / "work", inbox)

    capabilities = asyncio.run(server.call_tool("get_capabilities", {}))
    created = asyncio.run(server.call_tool("create_revision", {"source_path": str(source)}))
    assert isinstance(capabilities, tuple)
    assert isinstance(created, tuple)
    capabilities_data = capabilities[1]
    created_data = created[1]
    assert isinstance(capabilities_data, dict)
    assert isinstance(created_data, dict)
    revision_id = created_data["revision_id"]
    inspection = asyncio.run(
        server.call_tool("inspect_revision", {"revision_id": revision_id})
    )
    assert isinstance(inspection, tuple)
    inspection_data = inspection[1]
    assert isinstance(inspection_data, dict)

    assert capabilities_data["network_policy"] == "disabled"
    assert created_data["created"] is True
    assert inspection_data["paragraph_count"] == 2
    serialized = repr((capabilities_data, created_data, inspection_data))
    assert "private-name" not in serialized
    assert "Synthetic question" not in serialized


def test_mcp_process_brief_creates_delivery_without_returning_prose_or_paths(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "private-name.docx")
    work = tmp_path / "work"
    server = build_server(work, inbox)

    processed = asyncio.run(server.call_tool("process_brief", {"source_path": str(source)}))

    assert isinstance(processed, tuple)
    data = processed[1]
    assert isinstance(data, dict)
    assert data["status"] == "completed"
    assert data["correction_count"] == 0
    assert data["record_citation_count"] == 0
    assert data["authority_citation_count"] == 0
    assert data["citation_review_item_count"] == 0
    assert data["hyperlink_candidate_count"] == 0
    assert data["hyperlink_inserted_count"] == 0
    assert data["hyperlink_review_item_count"] == 0
    assert data["toc_entry_count"] == 1
    assert data["toa_authority_count"] == 0
    assert data["front_matter_field_candidate_count"] == 0
    assert data["lawyer_review_status"] == "not-required"
    assert (work / "deliveries" / data["process_id"] / "corrected.docx").is_file()
    assert (work / "deliveries" / data["process_id"] / "corrections.json").is_file()
    assert (work / "deliveries" / data["process_id"] / "front-matter.json").is_file()
    serialized = repr(data)
    assert "private-name" not in serialized
    assert "Synthetic question" not in serialized
    assert str(work) not in serialized


def test_mcp_process_brief_returns_proof_counts_and_pending_review(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = inbox / "proof.docx"
    document = Document()
    document.add_paragraph("Synthetic question , answered.")
    document.save(source)
    server = build_server(tmp_path / "work", inbox)

    processed = asyncio.run(server.call_tool("process_brief", {"source_path": str(source)}))

    assert isinstance(processed, tuple)
    data = processed[1]
    assert isinstance(data, dict)
    assert data["correction_count"] == 1
    assert data["applied_correction_count"] == 1
    assert data["review_item_count"] == 0
    assert data["citation_review_item_count"] == 0
    assert data["hyperlink_inserted_count"] == 0


def test_mcp_process_brief_returns_offline_link_counts_without_prose(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = inbox / "official.docx"
    document = Document()
    document.add_paragraph("Synthetic claim under 42 U.S.C. § 1983.")
    document.save(source)
    server = build_server(tmp_path / "work", inbox)

    processed = asyncio.run(server.call_tool("process_brief", {"source_path": str(source)}))

    assert isinstance(processed, tuple)
    data = processed[1]
    assert isinstance(data, dict)
    assert data["authority_citation_count"] == 1
    assert data["citation_review_item_count"] == 1
    assert data["hyperlink_candidate_count"] == 1
    assert data["hyperlink_inserted_count"] == 1
    assert data["hyperlink_review_item_count"] == 0
    assert data["toa_authority_count"] == 1
    assert data["toa_occurrence_count"] == 1
    serialized = repr(data)
    assert "42 U.S.C." not in serialized
    assert "Synthetic claim" not in serialized
    assert data["lawyer_review_status"] == "pending"


def test_mcp_rejects_unknown_arguments_and_remote_binding(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    server = build_server(tmp_path / "work", inbox)

    with pytest.raises(Exception, match="validation"):
        asyncio.run(server.call_tool("get_capabilities", {"unexpected": True}))
    with pytest.raises(ValueError, match="remote MCP binding"):
        build_server(tmp_path / "remote-work", inbox, host="0.0.0.0")
    with pytest.raises(ValueError, match="literal loopback"):
        build_server(tmp_path / "named-work", inbox, host="example.test")


def test_mcp_main_builds_selected_local_transport(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    calls: list[tuple[Path, Path, str, int, str]] = []

    class FakeServer:
        def run(self, *, transport: str) -> None:
            calls.append((tmp_path / "work", inbox, "localhost", 9010, transport))

    def fake_build(
        revision_root: Path, input_root: Path, *, host: str, port: int
    ) -> FakeServer:
        assert (revision_root, input_root, host, port) == (
            tmp_path / "work",
            inbox,
            "localhost",
            9010,
        )
        return FakeServer()

    monkeypatch.setattr("opencounsel.mcp.server.build_server", fake_build)

    assert main(
        [
            "--transport",
            "streamable-http",
            "--revision-root",
            str(tmp_path / "work"),
            "--input-root",
            str(inbox),
            "--host",
            "localhost",
            "--port",
            "9010",
        ]
    ) == 0
    assert calls[-1][-1] == "streamable-http"


def test_mcp_main_requires_explicit_storage_roots() -> None:
    with pytest.raises(SystemExit):
        main([])
    with pytest.raises(SystemExit):
        main(["--revision-root", "/tmp/work"])


def test_mcp_clean_brief_normalizes_formats_slots_and_verified_links(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = inbox / "legacy.docx"
    document = Document()
    document.add_heading("PRELIMINARY STATEMENT", level=1)
    document.add_paragraph("Synthetic introductory text.")
    document.add_heading("I. THE CLAIM IS PROPERLY PLEADED", level=2)
    document.add_paragraph(
        "Bell Atlantic Corp. v. Twombly, 550 U.S. 544, 570 (2007), and "
        "42 U.S.C. § 1983 control."
    )
    document.save(source)
    work = tmp_path / "work"
    server = build_server(work, inbox)

    cleaned = asyncio.run(
        server.call_tool(
            "clean_brief",
            {
                "source_path": str(source),
                "profile_id": "ny-ad-appellant-brief",
                "authority_links": {
                    "Bell Atlantic Corp. v. Twombly, 550 U.S. 544 (2007)":
                        "https://www.courtlistener.com/opinion/1444/bell-atlantic-corp-v-twombly/"
                },
            },
        )
    )

    assert isinstance(cleaned, tuple)
    data = cleaned[1]
    assert isinstance(data, dict)
    assert data["status"] == "completed"
    assert data["profile_id"] == "ny-ad-appellant-brief"
    assert data["formatting_applied_count"] > 0
    process = data["process"]
    corrected = work / "deliveries" / process["process_id"] / "corrected.docx"
    assert corrected.is_file()
    result = Document(corrected)
    texts = [paragraph.text for paragraph in result.paragraphs]
    assert texts[:4] == [
        "TABLE OF CONTENTS",
        "[TOC]",
        "TABLE OF AUTHORITIES",
        "[TOA]",
    ]
    assert "THE CLAIM IS PROPERLY PLEADED" in texts
    assert "I. THE CLAIM IS PROPERLY PLEADED" not in texts
    assert result.styles["_LegalBody"].font.name == "Times New Roman"
    assert process["toc_field_inserted_count"] == 1
    assert process["toa_field_inserted_count"] == 1
    assert data["authority_link_inserted_count"] == 2


def test_mcp_clean_brief_owns_complete_roa_style_citation_and_front_matter_pipeline(
    tmp_path: Path,
    roa_package: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    confined_roa = inbox / "synthetic-roa.zip"
    shutil.copyfile(roa_package, confined_roa)
    source = inbox / "legacy.docx"
    document = Document()
    document.add_heading("PRELIMINARY STATEMENT", level=1)
    document.add_paragraph("Synthetic introductory text.")
    document.add_heading("I. THE RECORD AND AUTHORITIES CONTROL", level=2)
    document.add_paragraph(
        "Bell Atlantic Corp. v. Twombly, 550 U.S. 544, 570 (2007), and "
        "42 U.S.C. § 1983 control; the record confirms the point (R. 3-4)."
    )
    document.add_paragraph("Twombly, 550 U.S. at 556, confirms the same rule.")
    document.save(source)
    original = source.read_bytes()
    work = tmp_path / "work"
    server = build_server(work, inbox)

    cleaned = asyncio.run(
        server.call_tool(
            "clean_brief",
            {
                "source_path": str(source),
                "profile_id": "ny-ad-appellant-brief",
                "roa_package_path": str(confined_roa),
                "authority_links": {
                    "Bell Atlantic Corp. v. Twombly, 550 U.S. 544 (2007)":
                        "https://www.courtlistener.com/opinion/1444/"
                        "bell-atlantic-corp-v-twombly/"
                },
            },
        )
    )

    assert isinstance(cleaned, tuple)
    data = cleaned[1]
    assert isinstance(data, dict)
    process = data["process"]
    assert source.read_bytes() == original
    assert process["record_citation_count"] == 1
    assert process["resolved_record_citation_count"] == 1
    assert process["unresolved_record_citation_count"] == 0
    assert process["record_source_sha256"] is not None
    assert process["authority_resolution_sha256"] is not None
    assert data["authority_link_inserted_count"] == 3

    delivery = work / "deliveries" / process["process_id"]
    corrected = delivery / "corrected.docx"
    ledger = json.loads((delivery / "corrections.json").read_text(encoding="utf-8"))
    record_items = [
        item for item in ledger["corrections"] if item["item_type"] == "record-citation"
    ]
    assert len(record_items) == 1
    assert "resolved against the supplied ROA" in record_items[0]["note"]
    assert ledger["record_source_sha256"] == process["record_source_sha256"]
    assert ledger["authority_resolution_sha256"] == process["authority_resolution_sha256"]

    result = Document(corrected)
    assert result.styles["_LegalBody"].font.name == "Times New Roman"
    assert process["toc_field_inserted_count"] == 1
    assert process["toa_field_inserted_count"] == 1
    with zipfile.ZipFile(corrected) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        relationships = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
    instructions = "".join(
        root.xpath(".//w:instrText/text()", namespaces={"w": W_NS})
    )
    assert " TOC " in instructions
    assert " TOA " in instructions
    external = {
        item.get("Target")
        for item in relationships
        if item.get("TargetMode") == "External"
    }
    assert "https://www.courtlistener.com/opinion/1444/bell-atlantic-corp-v-twombly/" in external
    assert any("uscode.house.gov" in target for target in external if target)
