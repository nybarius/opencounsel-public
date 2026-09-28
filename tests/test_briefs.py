from __future__ import annotations

import csv
import json
import zipfile
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from opencounsel.briefs.audit import audit_brief
from opencounsel.briefs.docx import BriefDocxError, inspect_brief_docx
from opencounsel.briefs.front_matter import build_front_matter_source
from opencounsel.briefs.links import AuthorityResolution, plan_authority_links
from opencounsel.briefs.report import write_brief_audit


def _brief(path: Path, *, linked: bool = False) -> Path:
    document = Document()
    document.styles.add_style("Argument Subheading", WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("ARGUMENT", level=1)
    paragraph = document.add_paragraph("The rule appears in Twombly, ")
    if linked:
        _add_hyperlink(paragraph, "550 U.S. 544", "https://www.courtlistener.com/opinion/1/")
    else:
        paragraph.add_run("550 U.S. 544")
    paragraph.add_run(" (2007), based on R 3-4.")
    document.save(path)
    return path


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def test_audits_brief_without_mutating_it(roa_package: Path, tmp_path: Path) -> None:
    brief = _brief(tmp_path / "brief.docx")
    before = brief.read_bytes()

    audit = audit_brief(brief, roa_package)

    assert brief.read_bytes() == before
    assert len(audit.headings) == 1
    assert audit.headings[0].text == "ARGUMENT"
    assert len(audit.record_citations) == 1
    assert audit.record_citations[0].record_pages == (3, 4)
    assert audit.record_citations[0].physical_pdf_pages == (1, 2)
    assert len(audit.authorities) == 1
    assert audit.authorities[0].normalized_text == "550 U.S. 544"
    assert audit.unlinked_authority_count == 1


def test_preserves_existing_authority_link(tmp_path: Path) -> None:
    brief = _brief(tmp_path / "linked.docx", linked=True)

    audit = audit_brief(brief)

    assert audit.authorities[0].hyperlink_target == ("https://www.courtlistener.com/opinion/1/")
    assert audit.unlinked_authority_count == 0


def test_builds_semantic_front_matter_source(tmp_path: Path) -> None:
    audit = audit_brief(_brief(tmp_path / "brief.docx"))

    source = build_front_matter_source(audit)

    assert source.brief_sha256 == audit.brief_sha256
    assert [(entry.level, entry.heading) for entry in source.toc] == [(1, "ARGUMENT")]
    assert len(source.toa) == 1
    assert source.toa[0].normalized_citation == "550 U.S. 544"
    assert len(source.toa[0].locations) == 1


def test_link_plan_requires_one_allowlisted_verified_resolution(tmp_path: Path) -> None:
    audit = audit_brief(_brief(tmp_path / "brief.docx"))
    resolution = AuthorityResolution(
        "550 U.S. 544",
        "verified",
        "https://www.courtlistener.com/opinion/1/",
        "courtlistener-citation-lookup-v4",
    )

    plan = plan_authority_links(audit, (resolution,))

    assert len(plan.insertions) == 1
    assert not plan.review
    assert plan.insertions[0].expected_text == "550 U.S. 544"

    rejected = plan_authority_links(
        audit,
        (
            AuthorityResolution(
                "550 U.S. 544",
                "verified",
                "https://example.com/opinion/1/",
                "untrusted",
            ),
        ),
    )
    assert not rejected.insertions
    assert rejected.review[0].reason == "resolver URL is not on the allowlist"


def test_writes_private_local_reports(roa_package: Path, tmp_path: Path) -> None:
    audit = audit_brief(_brief(tmp_path / "brief.docx"), roa_package)

    paths = write_brief_audit(audit, tmp_path / "audit", record_pdf_name="record.pdf")

    summary = json.loads(paths[0].read_text(encoding="utf-8"))
    assert summary["record_citation_count"] == 1
    with paths[1].open(encoding="utf-8", newline="") as stream:
        row = next(csv.DictReader(stream))
    assert row["pdf_targets"] == "record.pdf#page=1 record.pdf#page=2"
    assert paths[0].stat().st_mode & 0o777 == 0o600


def test_rejects_non_docx_and_symlink(tmp_path: Path) -> None:
    bad = tmp_path / "bad.docx"
    bad.write_text("not OOXML", encoding="utf-8")
    with pytest.raises(BriefDocxError, match="readable OOXML"):
        inspect_brief_docx(bad)

    link = tmp_path / "link.docx"
    link.symlink_to(bad)
    with pytest.raises(BriefDocxError, match="regular DOCX"):
        inspect_brief_docx(link)


def test_reads_footnotes_tabs_breaks_and_custom_heading(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.styles.add_style("_LegalCenter", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("CUSTOM SECTION", style="_LegalCenter")
    document.add_paragraph("Not a heading", style="_LegalCenter")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("left")
    run.add_tab()
    run.add_text("right")
    run.add_break(WD_BREAK.LINE)
    run.add_text("next")
    document.save(source)

    brief = tmp_path / "footnoted.docx"
    footnotes = b"""
    <w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:footnote w:id="-1"><w:p><w:r><w:t>separator</w:t></w:r></w:p></w:footnote>
      <w:footnote w:id="bad"><w:p><w:r><w:t>invalid</w:t></w:r></w:p></w:footnote>
      <w:footnote w:id="1"><w:p><w:r><w:t>Footnote R 5.</w:t></w:r></w:p></w:footnote>
    </w:footnotes>
    """
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(brief, "w") as outgoing:
        for member in incoming.infolist():
            outgoing.writestr(member, incoming.read(member))
        outgoing.writestr("word/footnotes.xml", footnotes)

    inspection = inspect_brief_docx(brief)

    assert [(heading.level, heading.text) for heading in inspection.headings] == [
        (1, "CUSTOM SECTION")
    ]
    assert any(paragraph.part == "footnote:1" for paragraph in inspection.paragraphs)
    assert any("left\tright\nnext" in paragraph.text for paragraph in inspection.paragraphs)


def test_ignores_non_web_hyperlink_target(tmp_path: Path) -> None:
    path = tmp_path / "brief.docx"
    document = Document()
    paragraph = document.add_paragraph()
    _add_hyperlink(paragraph, "550 U.S. 544", "javascript:alert(1)")
    document.save(path)

    inspection = inspect_brief_docx(path)

    assert not inspection.paragraphs[0].hyperlinks


def test_rejects_unsafe_or_incomplete_ooxml(tmp_path: Path, monkeypatch) -> None:
    empty = tmp_path / "empty.docx"
    with zipfile.ZipFile(empty, "w") as archive:
        archive.writestr("placeholder.txt", "x")
    with pytest.raises(BriefDocxError, match="no Word document body"):
        inspect_brief_docx(empty)

    unsafe = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(unsafe, "w") as archive:
        archive.writestr("../word/document.xml", "x")
    with pytest.raises(BriefDocxError, match="unsafe member path"):
        inspect_brief_docx(unsafe)

    invalid = tmp_path / "invalid.docx"
    with zipfile.ZipFile(invalid, "w") as archive:
        archive.writestr("word/document.xml", "<not-closed>")
    with pytest.raises(BriefDocxError, match="invalid XML"):
        inspect_brief_docx(invalid)

    valid = _brief(tmp_path / "large.docx")
    monkeypatch.setattr("opencounsel.briefs.docx.MAX_PACKAGE_BYTES", 1)
    with pytest.raises(BriefDocxError, match="50 MiB"):
        inspect_brief_docx(valid)


def test_rejects_corrupt_deflate_members(tmp_path: Path, monkeypatch) -> None:
    valid = _brief(tmp_path / "brief.docx")

    def corrupt_testzip(archive: zipfile.ZipFile) -> None:
        raise zlib.error("synthetic corrupt DEFLATE stream")

    monkeypatch.setattr(zipfile.ZipFile, "testzip", corrupt_testzip)
    with pytest.raises(BriefDocxError, match="unreadable members"):
        inspect_brief_docx(valid)

    def invalid_member_name(archive: zipfile.ZipFile) -> None:
        raise UnicodeDecodeError("utf-8", b"\xff", 0, 1, "synthetic invalid member name")

    monkeypatch.setattr(zipfile.ZipFile, "testzip", invalid_member_name)
    with pytest.raises(BriefDocxError, match="unreadable members"):
        inspect_brief_docx(valid)

    def overlapping_member(archive: zipfile.ZipFile) -> None:
        import warnings

        warnings.warn("synthetic overlapping member", UserWarning, stacklevel=1)

    monkeypatch.setattr(zipfile.ZipFile, "testzip", overlapping_member)
    with pytest.raises(BriefDocxError, match="unreadable members"):
        inspect_brief_docx(valid)


def test_rejects_invalid_zip_filename_encoding(tmp_path: Path, monkeypatch) -> None:
    valid = _brief(tmp_path / "brief.docx")

    def invalid_filename(path: Path) -> None:
        raise UnicodeDecodeError("utf-8", b"\xff", 0, 1, "synthetic invalid filename")

    monkeypatch.setattr(zipfile, "ZipFile", invalid_filename)
    with pytest.raises(BriefDocxError, match="readable OOXML"):
        inspect_brief_docx(valid)


def test_link_plan_flags_missing_and_ambiguous_resolutions(tmp_path: Path) -> None:
    audit = audit_brief(_brief(tmp_path / "brief.docx"))

    missing = plan_authority_links(audit, ())
    assert missing.review[0].reason == "no verified resolution"

    resolution = AuthorityResolution(
        "550 U.S. 544",
        "verified",
        "https://www.courtlistener.com/opinion/1/",
        "courtlistener",
    )
    ambiguous = plan_authority_links(audit, (resolution, resolution))
    assert ambiguous.review[0].reason == "multiple verified resolutions"
