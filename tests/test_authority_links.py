from __future__ import annotations

import json
import stat
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from urllib.error import HTTPError, URLError

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from opencounsel.adapters.citations import extract_citation_candidates
from opencounsel.briefs.audit import audit_brief
from opencounsel.briefs.authority_manifest import (
    build_authority_acquisition_manifest,
    write_authority_acquisition_manifest,
)
from opencounsel.briefs.docx import inspect_brief_docx
from opencounsel.briefs.link_docx import (
    LinkProjectionError,
    apply_authority_links,
    write_link_manifest,
)
from opencounsel.briefs.links import (
    AuthorityLinkPlan,
    AuthorityResolution,
    LinkInsertion,
    plan_authority_links,
)
from opencounsel.briefs.source_evidence import SourceEvidenceError, verify_resolution_urls
from opencounsel.briefs.sources import (
    SourceResolutionError,
    load_source_overrides,
    resolve_authority_sources,
)
from opencounsel.cli import main


def _statute_brief(path: Path) -> Path:
    document = Document()
    paragraph = document.add_paragraph("Jurisdiction follows 28 U.S.")
    paragraph.add_run("C. § 1291 and CPLR 3211(a)(7).")
    document.save(path)
    return path


def test_extracts_bounded_multi_section_and_vendor_citations() -> None:
    candidates = extract_citation_candidates(
        "See 28 U.S.C. §§ 1291, 1331; CPLR 3211(a)(7); and 2026 N.Y. Misc. LEXIS 1234."
    )

    assert [(value.normalized_text, value.category) for value in candidates] == [
        ("28 U.S.C. § 1291", "statutes"),
        ("28 U.S.C. § 1331", "statutes"),
        ("CPLR § 3211(a)(7)", "statutes"),
        ("2026 N.Y. Misc. LEXIS 1234", "cases"),
    ]
    assert [value.text for value in candidates[:2]] == ["28 U.S.C. §§ 1291", "1331"]


def test_short_case_citations_abstain_instead_of_becoming_authorities(tmp_path: Path) -> None:
    path = tmp_path / "short.docx"
    document = Document()
    document.add_paragraph("Twombly, 550 U.S. at 555, confirms the point. Id. at 556.")
    document.save(path)

    plan = plan_authority_links(audit_brief(path), ())

    assert not plan.insertions
    assert len(plan.review) == 2
    assert all("no uniquely approved antecedent" in item.reason for item in plan.review)


def test_resolves_unique_short_cite_and_id_across_body_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "shorts.docx"
    document = Document()
    document.add_paragraph("Bell Atlantic Corp. v. Twombly, 550 U.S. 544 (2007).")
    document.add_paragraph("Twombly, 550 U.S. at 555, controls.")
    document.add_paragraph("Id. at 556.")
    document.save(path)

    audit = audit_brief(path)
    resolution = AuthorityResolution(
        "550 U.S. 544",
        "verified",
        "https://www.courtlistener.com/opinion/1/",
        "test",
    )
    plan = plan_authority_links(audit, (resolution,))

    assert [item.resolved_citation for item in audit.authorities] == [
        "550 U.S. 544",
        "550 U.S. 544",
        "550 U.S. 544",
    ]
    assert [item.resolution_basis for item in audit.authorities] == [
        "full-citation",
        "unique-prior-full-citation",
        "immediate-antecedent",
    ]
    assert len(plan.insertions) == 3
    assert not plan.review


def test_does_not_resolve_footnote_id_from_body_authority(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Twombly, 550 U.S. 544 (2007).")
    document.save(source)
    target = tmp_path / "footnote.docx"
    footnotes = b"""
    <w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:footnote w:id="1"><w:p><w:r><w:t>Id. at 555.</w:t></w:r></w:p></w:footnote>
    </w:footnotes>
    """
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(target, "w") as outgoing:
        for member in incoming.infolist():
            outgoing.writestr(member, incoming.read(member))
        outgoing.writestr("word/footnotes.xml", footnotes)

    audit = audit_brief(target)

    assert audit.authorities[-1].part == "footnote:1"
    assert audit.authorities[-1].resolved_citation is None


def test_resolves_official_statute_sources_and_applies_across_runs(tmp_path: Path) -> None:
    source = _statute_brief(tmp_path / "source.docx")
    before = source.read_bytes()
    audit = audit_brief(source)
    resolutions = resolve_authority_sources(audit)
    plan = plan_authority_links(audit, resolutions)

    output = tmp_path / "linked.docx"
    result = apply_authority_links(source, output, plan)
    manifest = write_link_manifest(tmp_path / "linked.links.json", plan, result)

    assert source.read_bytes() == before
    assert result.inserted_count == 2
    assert not result.review
    assert stat.S_IMODE(output.stat().st_mode) == 0o600
    assert stat.S_IMODE(manifest.stat().st_mode) == 0o600
    inspection = inspect_brief_docx(output)
    targets = {
        hyperlink.target
        for paragraph in inspection.paragraphs
        for hyperlink in paragraph.hyperlinks
    }
    assert any(target.startswith("https://uscode.house.gov/view.xhtml") for target in targets)
    assert "https://www.nysenate.gov/legislation/laws/CVP/3211" in targets
    assert [paragraph.text for paragraph in inspection.paragraphs] == [
        paragraph.text for paragraph in inspect_brief_docx(source).paragraphs
    ]
    assert Document(output).paragraphs[0].text == Document(source).paragraphs[0].text
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    assert payload["input_sha256"] == result.input_sha256
    assert payload["output_sha256"] == result.output_sha256


def test_lexis_only_source_uses_explicit_local_override(tmp_path: Path) -> None:
    brief = tmp_path / "vendor.docx"
    document = Document()
    document.add_paragraph("The unpublished decision is 2026 N.Y. Misc. LEXIS 1234.")
    document.save(brief)
    overrides_path = tmp_path / "sources.toml"
    overrides_path.write_text(
        """schema_version = 1

[[authority]]
citation = "2026 N.Y. Misc. LEXIS 1234"
url = "https://plus.lexis.com/document/example"
note = "Firm subscription permalink"
""",
        encoding="utf-8",
    )

    overrides = load_source_overrides(overrides_path)
    resolutions = resolve_authority_sources(audit_brief(brief), overrides=overrides)
    plan = plan_authority_links(audit_brief(brief), resolutions)

    assert len(plan.insertions) == 1
    assert plan.insertions[0].source_type == "licensed-or-manual"
    assert plan.insertions[0].url == "https://plus.lexis.com/document/example"


def test_rejects_unsafe_override_and_hash_mismatch(tmp_path: Path) -> None:
    overrides = tmp_path / "bad.toml"
    overrides.write_text(
        'schema_version = 1\n[[authority]]\ncitation = "x"\nurl = "http://example.com"\n',
        encoding="utf-8",
    )
    with pytest.raises(SourceResolutionError, match="credential-free HTTPS"):
        load_source_overrides(overrides)

    source = _statute_brief(tmp_path / "source.docx")
    plan = plan_authority_links(audit_brief(source), resolve_authority_sources(audit_brief(source)))
    altered = tmp_path / "altered.docx"
    document = Document(source)
    document.add_paragraph("New text")
    document.save(altered)
    with pytest.raises(LinkProjectionError, match="does not match"):
        apply_authority_links(altered, tmp_path / "out.docx", plan)


def test_courtlistener_adapter_sends_only_citations_and_requires_unique_cluster(
    tmp_path: Path,
) -> None:
    path = tmp_path / "case.docx"
    document = Document()
    document.add_paragraph("See Twombly, 550 U.S. 544 (2007).")
    document.save(path)
    captured: dict[str, object] = {}

    class Response:
        def read(self, _limit: int) -> bytes:
            return json.dumps(
                [
                    {
                        "citation": "550 U.S. 544",
                        "status": 200,
                        "clusters": [{"absolute_url": "/opinion/1/example/"}],
                    }
                ]
            ).encode()

    def opener(request, *, timeout: int):
        captured["body"] = request.data
        captured["authorization"] = request.headers["Authorization"]
        captured["timeout"] = timeout
        return Response()

    resolutions = resolve_authority_sources(
        audit_brief(path),
        use_courtlistener=True,
        courtlistener_token="secret",
        opener=opener,
    )

    assert captured["body"] in {
        b"text=550+U.S.%C2%A0+544",
        b"text=550+U.S.+544",
    }
    assert captured["authorization"] == "Token secret"
    assert resolutions[0].url == "https://www.courtlistener.com/opinion/1/example/"


def test_link_brief_cli_writes_new_docx_and_manifest(tmp_path: Path, capsys) -> None:
    source = _statute_brief(tmp_path / "source.docx")
    output = tmp_path / "linked.docx"

    assert main(["link-brief", str(source), str(output)]) == 0

    response = json.loads(capsys.readouterr().out)
    assert response["status"] == "linked"
    assert response["inserted_count"] == 2
    assert output.is_file()
    assert output.with_suffix(".links.json").is_file()


def test_resolves_ecfr_and_rejects_missing_courtlistener_token(
    tmp_path: Path, monkeypatch
) -> None:
    path = tmp_path / "authorities.docx"
    document = Document()
    document.add_paragraph("See 28 C.F.R. § 0.1 and Twombly, 550 U.S. 544 (2007).")
    document.save(path)
    audit = audit_brief(path)

    resolutions = resolve_authority_sources(audit)

    assert resolutions[0].url == "https://www.ecfr.gov/current/title-28/section-0.1"
    assert any(authority.normalized_text == "550 U.S. 544" for authority in audit.authorities)
    monkeypatch.delenv("COURTLISTENER_API_TOKEN", raising=False)
    with pytest.raises(SourceResolutionError, match="is not set"):
        resolve_authority_sources(audit, use_courtlistener=True)


def test_resolves_new_york_slip_opinion_to_official_reporter(tmp_path: Path) -> None:
    path = tmp_path / "ny.docx"
    document = Document()
    document.add_paragraph(
        "See People v Doe, 2024 NY Slip Op 01234, ¶ 10; "
        "People v Doe, 2024 NY Slip Op 01234, ¶ 12."
    )
    document.save(path)

    resolutions = resolve_authority_sources(audit_brief(path))

    assert resolutions[0].url == (
        "https://www.nycourts.gov/reporter/3dseries/2024/2024_01234.htm"
    )


def test_recognizes_federal_and_new_york_rules_and_links_open_federal_source(
    tmp_path: Path,
) -> None:
    path = tmp_path / "rules.docx"
    document = Document()
    document.add_paragraph(
        "See Fed. R. Civ. P. 12(b)(6), Fed. R. Evid. 403, and 22 NYCRR § 1250.8(a)."
    )
    document.save(path)
    audit = audit_brief(path)
    resolutions = resolve_authority_sources(audit)
    plan = plan_authority_links(audit, resolutions)

    assert [item.normalized_text for item in audit.authorities] == [
        "Fed. R. Civ. P. 12(b)(6)",
        "Fed. R. Evid. 403",
        "22 N.Y.C.R.R. § 1250.8(a)",
    ]
    assert [item.url for item in resolutions] == [
        "https://www.law.cornell.edu/rules/frcp/rule_12",
        "https://www.law.cornell.edu/rules/fre/rule_403",
    ]
    assert len(plan.insertions) == 2
    assert plan.review[0].citation == "22 N.Y.C.R.R. § 1250.8(a)"


def test_prepares_private_missing_authority_bundle(tmp_path: Path) -> None:
    path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph(
        'Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234, held that "the agreement controls."'
    )
    document.save(path)
    manifest = build_authority_acquisition_manifest(audit_brief(path))

    paths = write_authority_acquisition_manifest(manifest, tmp_path / "bundle")

    payload = json.loads(paths[0].read_text(encoding="utf-8"))
    authority = payload["authorities"][0]
    assert authority["case_name"] == "Smith v. Jones"
    assert authority["source_status"] == "source-needed"
    assert len(manifest.authorities) == 1
    assert manifest.source_copy_required_count == 1
    assert manifest.case_source_copy_required_count == 1
    assert manifest.unresolved_occurrence_count == 0
    assert authority["verification_targets"] == [
        "citation-exists",
        "quotation-exact",
        "quotation-signals-complete",
        "characterization-supported",
    ]
    assert "the agreement controls" in authority["assertions"][0]["context"]
    assert paths[1].read_text(encoding="utf-8").startswith(
        "2026 N.Y. Misc. LEXIS 1234;Smith v. Jones;"
    )
    assert all(stat.S_IMODE(path.stat().st_mode) == 0o600 for path in paths)


def test_courtlistener_search_fallback_resolves_vendor_case(tmp_path: Path) -> None:
    path = tmp_path / "vendor.docx"
    document = Document()
    document.add_paragraph("Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234.")
    document.save(path)
    requests: list[object] = []

    class Response:
        def __init__(self, body: object):
            self.body = body

        def read(self, _limit: int) -> bytes:
            return json.dumps(self.body).encode()

    def opener(request, *, timeout: int):
        requests.append(request)
        if request.data is not None:
            return Response(
                [{"citation": "2026 N.Y. Misc. LEXIS 1234", "status": 404, "clusters": []}]
            )
        return Response(
            {
                "results": [
                    {
                        "caseName": "Smith v. Jones",
                        "dateFiled": "2026-02-03",
                        "absolute_url": "/opinion/42/smith-v-jones/",
                    }
                ]
            }
        )

    resolutions = resolve_authority_sources(
        audit_brief(path),
        use_courtlistener=True,
        courtlistener_token="token",
        opener=opener,
    )

    assert len(requests) == 2
    assert requests[1].data is None
    assert resolutions[0].resolver == "courtlistener-search-v4"
    assert resolutions[0].url == "https://www.courtlistener.com/opinion/42/smith-v-jones/"


def test_url_evidence_cache_reuses_fresh_success_and_fails_closed(tmp_path: Path) -> None:
    resolution = AuthorityResolution(
        "550 U.S. 544",
        "verified",
        "https://www.courtlistener.com/opinion/1/",
        "test",
    )
    calls = 0

    class Response:
        def __init__(self) -> None:
            self.status = 200
            self.url = "https://www.courtlistener.com/opinion/1/"
            self.headers = {"Content-Type": "text/html"}

    def opener(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        return Response()

    cache = tmp_path / "evidence.json"
    now = datetime(2026, 7, 15, tzinfo=UTC)
    first = verify_resolution_urls((resolution,), cache, opener=opener, now=now)
    second = verify_resolution_urls((resolution,), cache, opener=opener, now=now)

    assert first[0].status == second[0].status == "verified"
    assert calls == 1
    assert stat.S_IMODE(cache.stat().st_mode) == 0o600

    def broken(request, **_kwargs):
        raise HTTPError(request.full_url, 404, "missing", {}, None)

    failed = verify_resolution_urls(
        (resolution,), tmp_path / "broken.json", opener=broken, now=now
    )
    assert failed[0].status == "unverified"


def test_url_evidence_falls_back_to_bounded_get_and_handles_missing_url(
    tmp_path: Path,
) -> None:
    resolution = AuthorityResolution(
        "550 U.S. 544",
        "verified",
        "https://www.courtlistener.com/opinion/1/",
        "test",
    )
    methods: list[str] = []

    class Response:
        def __init__(self) -> None:
            self.status = 206
            self.url = "https://www.courtlistener.com/opinion/1/"
            self.headers: dict[str, str] = {}

    def opener(request, **_kwargs):
        methods.append(request.method)
        if request.method == "HEAD":
            raise HTTPError(request.full_url, 405, "method", {}, None)
        assert request.headers["Range"] == "bytes=0-0"
        return Response()

    no_url = AuthorityResolution("unknown", "verified", None, "test")
    results = verify_resolution_urls(
        (resolution, no_url), tmp_path / "cache.json", opener=opener
    )

    assert methods == ["HEAD", "GET"]
    assert [item.status for item in results] == ["verified", "unverified"]


def test_url_evidence_rechecks_expired_and_rejects_invalid_cache(tmp_path: Path) -> None:
    resolution = AuthorityResolution(
        "550 U.S. 544",
        "verified",
        "https://www.courtlistener.com/opinion/1/",
        "test",
    )
    cache = tmp_path / "cache.json"
    cache.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "evidence": [
                    {
                        "url": resolution.url,
                        "checked_at": "not-a-date",
                        "status_code": 200,
                        "final_url": resolution.url,
                        "content_type": None,
                        "valid": True,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    def unavailable(*_args, **_kwargs):
        raise URLError("offline")

    result = verify_resolution_urls((resolution,), cache, opener=unavailable)
    assert result[0].status == "unverified"

    invalid = tmp_path / "invalid.json"
    invalid.write_text("[]", encoding="utf-8")
    with pytest.raises(SourceEvidenceError, match="invalid"):
        verify_resolution_urls((resolution,), invalid)

    cache_link = tmp_path / "cache-link.json"
    cache_link.symlink_to(invalid)
    with pytest.raises(SourceEvidenceError, match="regular file"):
        verify_resolution_urls((resolution,), cache_link)

    naive = tmp_path / "naive.json"
    naive.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "evidence": [
                    {
                        "url": resolution.url,
                        "checked_at": "2026-07-15T00:00:00",
                        "status_code": 200,
                        "final_url": resolution.url,
                        "content_type": "text/html",
                        "valid": True,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    fresh = verify_resolution_urls(
        (resolution,), naive, opener=unavailable, now=datetime(2026, 7, 15, tzinfo=UTC)
    )
    assert fresh[0].status == "verified"


def test_url_evidence_get_fallback_failure_is_unverified(tmp_path: Path) -> None:
    resolution = AuthorityResolution(
        "550 U.S. 544",
        "verified",
        "https://www.courtlistener.com/opinion/1/",
        "test",
    )

    def opener(request, **_kwargs):
        code = 405 if request.method == "HEAD" else 503
        raise HTTPError(request.full_url, code, "failed", {}, None)

    result = verify_resolution_urls(
        (resolution,), tmp_path / "fallback-failed.json", opener=opener
    )
    assert result[0].status == "unverified"


@pytest.mark.parametrize(
    ("body", "message"),
    [
        (b"not-json", "invalid JSON"),
        (b'{"result": []}', "invalid result shape"),
        (b" " * (1024 * 1024 + 1), "exceeded 1 MiB"),
    ],
)
def test_rejects_invalid_courtlistener_responses(
    tmp_path: Path, body: bytes, message: str
) -> None:
    path = tmp_path / "case.docx"
    document = Document()
    document.add_paragraph("Twombly, 550 U.S. 544 (2007).")
    document.save(path)

    class Response:
        def read(self, _limit: int) -> bytes:
            return body

    with pytest.raises(SourceResolutionError, match=message):
        resolve_authority_sources(
            audit_brief(path),
            use_courtlistener=True,
            courtlistener_token="token",
            opener=lambda *_args, **_kwargs: Response(),
        )


def test_courtlistener_failure_and_ambiguous_cluster_abstain(tmp_path: Path) -> None:
    path = tmp_path / "case.docx"
    document = Document()
    document.add_paragraph("Twombly, 550 U.S. 544 (2007).")
    document.save(path)
    audit = audit_brief(path)

    def failed(*_args, **_kwargs):
        raise TimeoutError

    with pytest.raises(SourceResolutionError, match="lookup failed"):
        resolve_authority_sources(
            audit,
            use_courtlistener=True,
            courtlistener_token="token",
            opener=failed,
        )

    class AmbiguousResponse:
        def read(self, _limit: int) -> bytes:
            return json.dumps(
                [
                    "invalid",
                    {"citation": None, "status": 200, "clusters": []},
                    {
                        "citation": "550 U.S. 544",
                        "status": 200,
                        "clusters": [
                            {"absolute_url": "/opinion/1/"},
                            {"absolute_url": "/opinion/2/"},
                            {"absolute_url": "https://evil.example/opinion/3/"},
                        ],
                    },
                ]
            ).encode()

    assert not resolve_authority_sources(
        audit,
        use_courtlistener=True,
        courtlistener_token="token",
        opener=lambda *_args, **_kwargs: AmbiguousResponse(),
    )



def test_courtlistener_lookup_abstains_from_single_cross_host_result(tmp_path: Path) -> None:
    path = tmp_path / "case.docx"
    document = Document()
    document.add_paragraph("Twombly, 550 U.S. 544 (2007).")
    document.save(path)

    class CrossHostResponse:
        def read(self, _limit: int) -> bytes:
            return json.dumps(
                [
                    {
                        "citation": "550 U.S. 544",
                        "status": 200,
                        "clusters": [
                            {"absolute_url": "https://evil.example/opinion/3/"},
                        ],
                    }
                ]
            ).encode()

    assert not resolve_authority_sources(
        audit_brief(path),
        use_courtlistener=True,
        courtlistener_token="token",
        opener=lambda *_args, **_kwargs: CrossHostResponse(),
    )



@pytest.mark.parametrize(
    "content",
    [
        "schema_version = 2\n",
        'schema_version = 1\nauthority = "bad"\n',
        'schema_version = 1\nauthority = ["bad"]\n',
        'schema_version = 1\n[[authority]]\ncitation = ""\nurl = "https://example.com"\n',
        (
            'schema_version = 1\n[[authority]]\ncitation = "x"\n'
            'url = "https://example.com"\nnote = 3\n'
        ),
        (
            'schema_version = 1\n[[authority]]\ncitation = "x"\n'
            'url = "https://example.com/1"\n[[authority]]\ncitation = "X"\n'
            'url = "https://example.com/2"\n'
        ),
    ],
)
def test_rejects_malformed_override_manifests(tmp_path: Path, content: str) -> None:
    path = tmp_path / "bad.toml"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(SourceResolutionError):
        load_source_overrides(path)


def test_override_manifest_must_be_a_regular_file(tmp_path: Path) -> None:
    with pytest.raises(SourceResolutionError, match="regular TOML"):
        load_source_overrides(tmp_path / "missing.toml")


def test_projection_refuses_in_place_existing_signed_and_overlapping_outputs(
    tmp_path: Path,
) -> None:
    source = _statute_brief(tmp_path / "source.docx")
    audit = audit_brief(source)
    plan = plan_authority_links(audit, resolve_authority_sources(audit))
    with pytest.raises(LinkProjectionError, match="new DOCX"):
        apply_authority_links(source, source, plan)

    existing = tmp_path / "existing.docx"
    existing.touch()
    with pytest.raises(LinkProjectionError, match="already exists"):
        apply_authority_links(source, existing, plan)

    broken_link = tmp_path / "broken-output.docx"
    broken_link.symlink_to(tmp_path / "missing-target")
    with pytest.raises(LinkProjectionError, match="already exists"):
        apply_authority_links(source, broken_link, plan)

    first = plan.insertions[0]
    overlap = LinkInsertion(
        first.part,
        first.paragraph_index,
        first.start + 1,
        first.end,
        first.expected_text[1:],
        first.url,
        first.resolver,
    )
    bad_plan = AuthorityLinkPlan(plan.brief_sha256, (first, overlap), ())
    with pytest.raises(LinkProjectionError, match="overlapping"):
        apply_authority_links(source, tmp_path / "overlap.docx", bad_plan)

    signed = tmp_path / "signed.docx"
    signed.write_bytes(source.read_bytes())
    with zipfile.ZipFile(signed, "a") as archive:
        archive.writestr("_xmlsignatures/sig1.xml", "<signature/>")
    signed_audit = audit_brief(signed)
    signed_plan = plan_authority_links(signed_audit, resolve_authority_sources(signed_audit))
    with pytest.raises(LinkProjectionError, match="digitally signed"):
        apply_authority_links(signed, tmp_path / "signed-output.docx", signed_plan)


def test_projection_routes_existing_hyperlink_wrapper_to_review(tmp_path: Path) -> None:
    source = tmp_path / "linked.docx"
    document = Document()
    paragraph = document.add_paragraph("Twombly, ")
    relationship_id = paragraph.part.relate_to(
        "https://www.courtlistener.com/opinion/1/",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "550 U.S. 544"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    document.save(source)
    audit = audit_brief(source)
    authority = audit.authorities[0]
    insertion = LinkInsertion(
        authority.part,
        authority.paragraph_index,
        authority.start,
        authority.end,
        authority.text,
        "https://www.courtlistener.com/opinion/2/",
        "test",
    )
    plan = AuthorityLinkPlan(audit.brief_sha256, (insertion,), ())

    result = apply_authority_links(source, tmp_path / "reviewed.docx", plan)

    assert result.inserted_count == 0
    assert "crosses a field, hyperlink" in result.review[0].reason
