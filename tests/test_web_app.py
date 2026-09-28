from __future__ import annotations

import io
import shutil
import threading
import time
from pathlib import Path

from docx import Document
from starlette.testclient import TestClient

from opencounsel.demo import build_synthetic_demo
from opencounsel.web import app as web_app
from opencounsel.web import server as web_server
from opencounsel.web.app import create_app
from opencounsel.web.jobs import JobContext, JobManager, ProgressUpdate


def _docx_bytes() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph("Synthetic text (R. 3).")
    document.save(stream)
    return stream.getvalue()


def _fake_processor(
    context: JobContext,
    progress: ProgressUpdate,
) -> dict[str, object]:
    progress("preparing", "Resolving citations and applying filing styles")
    export = context.job_dir / "export"
    export.mkdir(mode=0o700)
    package = export / "opencounsel-filing-package.zip"
    document = export / "published.docx"
    pdf = export / "published.pdf"
    original_pdf = export / "original-brief.pdf"
    sources = export / "authority-source-report.json"
    package.write_bytes(b"PK synthetic")
    document.write_bytes(b"synthetic docx")
    pdf.write_bytes(b"%PDF synthetic")
    original_pdf.write_bytes(b"%PDF immutable original")
    sources.write_text("{}\n", encoding="utf-8")
    progress("packaging", "Packaging filing and review artifacts")
    return {
        "summary": {
            "record_citation_count": 3,
            "resolved_record_citation_count": 3,
            "unresolved_record_citation_count": 0,
            "toa_authority_count": 2,
            "toa_occurrence_count": 4,
            "toc_entry_count": 5,
            "hyperlink_inserted_count": 2,
            "review_item_count": 4,
            "page_count": 7,
            "body_font_pt": 14,
            "footnote_font_pt": 12,
        },
        "sources": [
            {
                "citation": "Example v. Case, 1 A.D.3d 2",
                "category": "cases",
                "status": "source-copy-required",
                "url": None,
                "occurrence_count": 2,
            }
        ],
        "audit": {
            "input_sha256": "1" * 64,
            "output_docx_sha256": "2" * 64,
            "output_pdf_sha256": "3" * 64,
        },
        "artifacts": {
            "package": {
                "path": "export/opencounsel-filing-package.zip",
                "filename": "opencounsel-filing-package.zip",
                "media_type": "application/zip",
            },
            "document": {
                "path": "export/published.docx",
                "filename": "published.docx",
                "media_type": (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
            },
            "pdf": {
                "path": "export/published.pdf",
                "filename": "published.pdf",
                "media_type": "application/pdf",
            },
            "original-pdf": {
                "path": "export/original-brief.pdf",
                "filename": "original-brief.pdf",
                "media_type": "application/pdf",
            },
            "sources": {
                "path": "export/authority-source-report.json",
                "filename": "authority-source-report.json",
                "media_type": "application/json",
            },
        },
    }


def _completed(client: TestClient, job_id: str) -> dict[str, object]:
    for _ in range(100):
        response = client.get(f"/api/jobs/{job_id}")
        payload = response.json()
        if payload["status"] in {"completed", "failed"}:
            return payload
        time.sleep(0.01)
    raise AssertionError("job did not finish")


def test_local_web_flow_uploads_processes_and_allowlists_artifacts(
    tmp_path: Path,
    roa_package: Path,
) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_fake_processor)
    app = create_app(manager=manager)
    with TestClient(app) as client:
        home = client.get("/")
        assert home.status_code == 200
        assert "Prepare an appellate brief" in home.text
        assert "Runs locally" in home.text
        assert home.headers["cache-control"] == "no-store"
        assert client.get("/assets/styles.css").status_code == 200
        assert client.get("/assets/app.js").status_code == 200

        profiles = client.get("/api/profiles").json()["profiles"]
        assert any(item["profile_id"] == "ny-ad-appellant-brief" for item in profiles)

        response = client.post(
            "/api/jobs",
            data={"profile": "ny-ad-appellant-brief"},
            files={
                "brief": (
                    "appeal.docx",
                    _docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
                "roa": ("record.zip", roa_package.read_bytes(), "application/zip"),
            },
        )
        assert response.status_code == 202
        job_id = response.json()["job_id"]
        result = _completed(client, job_id)

        assert result["status"] == "completed"
        assert result["summary"]["resolved_record_citation_count"] == 3
        assert result["summary"]["body_font_pt"] == 14
        package = client.get(f"/api/jobs/{job_id}/artifacts/package")
        assert package.status_code == 200
        assert package.content == b"PK synthetic"
        original = client.get(
            f"/api/jobs/{job_id}/artifacts/original-pdf?inline=1"
        )
        assert original.status_code == 200
        assert original.content == b"%PDF immutable original"
        assert original.headers["content-disposition"].startswith("inline;")
        assert client.get(f"/api/jobs/{job_id}/artifacts/../../state").status_code == 404
        assert client.delete(f"/api/jobs/{job_id}").status_code == 204
        assert client.get(f"/api/jobs/{job_id}").status_code == 404


def test_security_headers_allow_only_inline_pdf_to_be_same_origin_framed(
    tmp_path: Path,
    roa_package: Path,
) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_fake_processor)
    with TestClient(create_app(manager=manager)) as client:
        response = client.post(
            "/api/jobs",
            data={"profile": "ny-ad-appellant-brief"},
            files={
                "brief": ("appeal.docx", _docx_bytes()),
                "roa": ("record.zip", roa_package.read_bytes()),
            },
        )
        job_id = response.json()["job_id"]
        assert _completed(client, job_id)["status"] == "completed"

        inline_pdf = client.get(f"/api/jobs/{job_id}/artifacts/pdf?inline=1")
        assert inline_pdf.headers["content-type"] == "application/pdf"
        assert inline_pdf.headers["content-disposition"].startswith("inline;")
        assert inline_pdf.headers["x-frame-options"] == "SAMEORIGIN"
        assert "frame-ancestors 'self'" in inline_pdf.headers[
            "content-security-policy"
        ]
        assert "frame-ancestors 'none'" not in inline_pdf.headers[
            "content-security-policy"
        ]

        protected_responses = (
            client.get("/"),
            client.get("/api/profiles"),
            client.get("/assets/app.js"),
            client.get(f"/api/jobs/{job_id}/artifacts/pdf"),
            client.get(f"/api/jobs/{job_id}/artifacts/document?inline=1"),
        )
        for protected in protected_responses:
            assert protected.headers["x-frame-options"] == "DENY"
            assert "frame-ancestors 'none'" in protected.headers[
                "content-security-policy"
            ]
            assert "frame-ancestors 'self'" not in protected.headers[
                "content-security-policy"
            ]


def test_web_rejects_unapproved_input_types(tmp_path: Path, roa_package: Path) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_fake_processor)
    with TestClient(create_app(manager=manager)) as client:
        response = client.post(
            "/api/jobs",
            data={"profile": "ny-ad-appellant-brief"},
            files={
                "brief": ("appeal.txt", b"not a docx", "text/plain"),
                "roa": ("record.zip", roa_package.read_bytes(), "application/zip"),
            },
        )
        assert response.status_code == 400
        assert response.json()["error"] == "The brief must be a .docx file."
        assert client.get("/api/jobs/not-a-job").status_code == 404


def test_web_accepts_attested_searchable_pdf_and_reports_input_errors(
    tmp_path: Path,
) -> None:
    demo_dir = tmp_path / "demo-input"
    build_synthetic_demo(demo_dir)
    pdf = (demo_dir / "record.pdf").read_bytes()
    manager = JobManager(tmp_path / "jobs", processor=_fake_processor)
    with TestClient(create_app(manager=manager)) as client:
        accepted = client.post(
            "/api/jobs",
            data={
                "profile": "ny-ad-appellant-brief",
                "first_record_page": "3",
                "numbering_verified": "on",
            },
            files={
                "brief": ("appeal.docx", _docx_bytes()),
                "roa": ("record.pdf", pdf, "application/pdf"),
            },
        )
        assert accepted.status_code == 202
        result = _completed(client, accepted.json()["job_id"])
        assert result["status"] == "completed"

        invalid_page = client.post(
            "/api/jobs",
            data={
                "profile": "ny-ad-appellant-brief",
                "first_record_page": "first",
                "numbering_verified": "on",
            },
            files={
                "brief": ("appeal.docx", _docx_bytes()),
                "roa": ("record.pdf", pdf, "application/pdf"),
            },
        )
        assert invalid_page.status_code == 400
        assert "positive integer" in invalid_page.json()["error"]

        missing_brief = client.post(
            "/api/jobs",
            data={"profile": "ny-ad-appellant-brief"},
            files={"roa": ("record.pdf", pdf, "application/pdf")},
        )
        assert missing_brief.status_code == 400
        assert missing_brief.json()["error"] == "The brief file is required."


def test_failed_job_and_running_delete_are_safe(
    tmp_path: Path,
    roa_package: Path,
) -> None:
    def failing(context: JobContext, progress: ProgressUpdate) -> dict[str, object]:
        progress("validating", "Validating brief and record")
        raise RuntimeError("")

    failed_manager = JobManager(tmp_path / "failed", processor=failing)
    with TestClient(create_app(manager=failed_manager)) as client:
        response = client.post(
            "/api/jobs",
            data={"profile": "ny-ad-appellant-brief"},
            files={
                "brief": ("appeal.docx", _docx_bytes()),
                "roa": ("record.zip", roa_package.read_bytes()),
            },
        )
        result = _completed(client, response.json()["job_id"])
        assert result["status"] == "failed"
        assert result["error"].startswith("OpenCounsel could not prepare")
        assert any(stage["status"] == "failed" for stage in result["stages"])
        assert client.get(
            f"/api/jobs/{response.json()['job_id']}/artifacts/missing"
        ).status_code == 404

    started = threading.Event()
    release = threading.Event()

    def blocking(context: JobContext, progress: ProgressUpdate) -> dict[str, object]:
        started.set()
        release.wait(timeout=2)
        return _fake_processor(context, progress)

    running_manager = JobManager(tmp_path / "running", processor=blocking)
    with TestClient(create_app(manager=running_manager)) as client:
        response = client.post(
            "/api/jobs",
            data={"profile": "ny-ad-appellant-brief"},
            files={
                "brief": ("appeal.docx", _docx_bytes()),
                "roa": ("record.zip", roa_package.read_bytes()),
            },
        )
        job_id = response.json()["job_id"]
        assert started.wait(timeout=1)
        assert client.delete(f"/api/jobs/{job_id}").status_code == 409
        release.set()
        assert _completed(client, job_id)["status"] == "completed"


def test_demo_failure_and_default_app_root_are_handled(
    tmp_path: Path,
    monkeypatch,
) -> None:
    manager = JobManager(tmp_path / "failed-demo", processor=_fake_processor)
    monkeypatch.setattr(
        web_app,
        "build_synthetic_demo",
        lambda _inbox: (_ for _ in ()).throw(RuntimeError("synthetic failure")),
    )
    with TestClient(create_app(manager=manager)) as client:
        assert client.post("/api/demo").status_code == 500

    monkeypatch.setenv("OPENCOUNSEL_WEB_ROOT", str(tmp_path / "default-root"))
    with TestClient(create_app()) as client:
        assert client.get("/healthz").json()["status"] == "ready"


def test_web_server_uses_private_runtime_defaults(tmp_path: Path, monkeypatch) -> None:
    captured = {}

    def fake_run(app, **options) -> None:
        captured["app"] = app
        captured.update(options)

    monkeypatch.setattr(web_server.uvicorn, "run", fake_run)
    web_server.run(host="127.0.0.1", port=8765, root=tmp_path, libreoffice="soffice")

    assert captured["host"] == "127.0.0.1"
    assert captured["port"] == 8765
    assert captured["access_log"] is False


def test_synthetic_demo_uses_the_same_job_path(tmp_path: Path) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_fake_processor)
    with TestClient(create_app(manager=manager)) as client:
        response = client.post("/api/demo")
        assert response.status_code == 202
        result = _completed(client, response.json()["job_id"])
        assert result["status"] == "completed"


def test_real_synthetic_demo_reaches_verified_publication(tmp_path: Path) -> None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        return
    manager = JobManager(tmp_path / "jobs", libreoffice=executable)
    reservation = manager.reserve(
        profile_id="ny-ad-appellant-brief",
        brief_name="synthetic-appellant-brief.docx",
        roa_name="synthetic-record.zip",
        demo=True,
    )
    build_synthetic_demo(reservation.context.inbox_dir)
    manager.submit(reservation.context)
    try:
        for _ in range(300):
            result = manager.snapshot(reservation.context.job_id)
            if result["status"] in {"completed", "failed"}:
                break
            time.sleep(0.05)
        else:
            raise AssertionError("real demonstration did not finish")
        assert result["status"] == "completed", result.get("error")
        assert result["summary"]["record_citation_count"] == 0
        assert result["summary"]["resolved_record_citation_count"] == 0
        assert result["summary"]["unresolved_record_citation_count"] == 0
        assert result["summary"]["toc_entry_count"] >= 7
        assert result["summary"]["source_copy_required_count"] == 4
        roa_artifact = manager.artifact(reservation.context.job_id, "roa-package")
        assert roa_artifact.filename == "machine-readable-roa.zip"
        assert roa_artifact.path == reservation.context.roa_path
        assert {item["status"] for item in result["sources"]} == {
            "official-or-open-source",
            "source-copy-required",
        }
        assert manager.artifact(reservation.context.job_id, "package").path.is_file()
        original_pdf = manager.artifact(
            reservation.context.job_id, "original-pdf"
        )
        assert original_pdf.filename == "original-brief.pdf"
        assert original_pdf.path.read_bytes().startswith(b"%PDF")
        published_docx = manager.artifact(
            reservation.context.job_id, "document"
        )
        published_document = Document(published_docx.path)
        published_text = [
            paragraph.text
            for paragraph in published_document.paragraphs
            if paragraph.text.strip()
        ]
        assert published_text[0] == (
            "UNITED STATES DISTRICT COURT\nSOUTHERN DISTRICT OF NEW YORK"
        )
        assert len(published_document.tables) == 1
        caption_text = "\n".join(
            cell.text for cell in published_document.tables[0].rows[0].cells
        )
        assert "THE NEW YORK TIMES COMPANY," in caption_text
        assert "DEFENDANTS\N{RIGHT SINGLE QUOTATION MARK} MOTION TO" in caption_text
        assert "TABLE OF CONTENTS" in published_text
        assert any(
            text.startswith("17 U.S.C. § 507(b)\t") for text in published_text
        )
    finally:
        manager.close()
