from __future__ import annotations

import csv
import io
import json
import time
import zipfile
from pathlib import Path
from typing import cast

from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, RectangleObject
from starlette.testclient import TestClient

from opencounsel.web.app import create_app
from opencounsel.web.jobs import JobContext, JobManager, ProgressUpdate

WORD_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
AUTHORITY_ID = "auth-1234567890abcdef"
CITATION = "Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234"


def _docx_bytes() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph(f'{CITATION}, held that "the agreement is unambiguous."')
    document.save(stream)
    return stream.getvalue()


def _source_pdf(url: str) -> bytes:
    stream = io.BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    text = f"Smith v. Jones. {CITATION}. The agreement is unambiguous."
    content = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.add_uri(0, url, RectangleObject((72, 690, 300, 710)))
    writer.write(stream)
    return stream.getvalue()


def _processor(context: JobContext, progress: ProgressUpdate) -> dict[str, object]:
    progress("reporting", "Building source and review reports")
    export = context.job_dir / "export"
    bundle = export / "reports" / "authority-sources"
    bundle.mkdir(parents=True, mode=0o700)
    manifest = {
        "schema_version": 1,
        "brief_sha256": "1" * 64,
        "unresolved_occurrence_count": 0,
        "authorities": [
            {
                "authority_id": AUTHORITY_ID,
                "canonical_citation": CITATION,
                "case_name": "Smith v. Jones",
                "category": "cases",
                "court": "Synthetic Court",
                "year": "2026",
                "docket": None,
                "source_status": "source-needed",
                "source_url": None,
                "source_resolver": None,
                "expected_filename": f"{AUTHORITY_ID}.pdf",
                "assertions": [
                    {
                        "part": "body",
                        "paragraph_index": 1,
                        "citation_text": "2026 N.Y. Misc. LEXIS 1234",
                        "citation_type": "case",
                        "context": 'Smith held that "the agreement is unambiguous."',
                        "resolution_basis": "full-citation",
                    }
                ],
                "verification_targets": [
                    "citation-exists",
                    "quotation-exact",
                    "quotation-signals-complete",
                    "characterization-supported",
                ],
            }
        ],
    }
    (bundle / "authority-manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    with (bundle / "authority-intake.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=(
                "authority_id",
                "canonical_citation",
                "case_name",
                "expected_filename",
                "source_file",
                "source_url",
            ),
        )
        writer.writeheader()
        writer.writerow(
            {
                "authority_id": AUTHORITY_ID,
                "canonical_citation": CITATION,
                "case_name": "Smith v. Jones",
                "expected_filename": f"{AUTHORITY_ID}.pdf",
                "source_file": "",
                "source_url": "",
            }
        )
    (export / "published.docx").write_bytes(_docx_bytes())
    (export / "published.pdf").write_bytes(_source_pdf("https://example.test/first.pdf"))
    (export / "opencounsel-filing-package.zip").write_bytes(b"PK first pass")
    return {
        "summary": {"toa_authority_count": 1, "page_count": 1},
        "sources": [
            {
                "authority_id": AUTHORITY_ID,
                "citation": CITATION,
                "category": "cases",
                "status": "source-copy-required",
                "url": None,
                "occurrence_count": 1,
            }
        ],
        "audit": {"input_sha256": "1" * 64},
        "artifacts": {
            "document": {
                "path": "export/published.docx",
                "filename": "published.docx",
                "media_type": WORD_MEDIA_TYPE,
            },
            "pdf": {
                "path": "export/published.pdf",
                "filename": "published.pdf",
                "media_type": "application/pdf",
            },
        },
    }


def _prepared(client: TestClient) -> str:
    response = client.post(
        "/api/jobs",
        data={
            "profile": "fed-sdny-edny-motion-memorandum",
            "record_mode": "no-record",
        },
        files={"brief": ("memorandum.docx", _docx_bytes(), WORD_MEDIA_TYPE)},
    )
    assert response.status_code == 202
    job_id = cast(str, response.json()["job_id"])
    for _ in range(100):
        if client.get(f"/api/jobs/{job_id}").json()["status"] == "completed":
            return job_id
        time.sleep(0.01)
    raise AssertionError("job did not complete")


def test_verified_pdf_link_requires_approval_and_creates_separate_artifacts(
    tmp_path: Path,
) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_processor)
    with TestClient(create_app(manager=manager)) as client:
        job_id = _prepared(client)
        original_docx = client.get(f"/api/jobs/{job_id}/artifacts/document").content
        original_pdf = client.get(f"/api/jobs/{job_id}/artifacts/pdf").content
        durable_url = "https://law.example/opinions/smith-v-jones-2026.pdf"
        uploaded = client.post(
            f"/api/jobs/{job_id}/sources",
            data={"authority_id": AUTHORITY_ID},
            files={"source": ("opinion.pdf", _source_pdf(durable_url), "application/pdf")},
        )
        assert uploaded.status_code == 200, uploaded.text
        pending = uploaded.json()["link_disposition"]
        assert pending["candidate_count"] == 1
        assert pending["approved_link_count"] == 0
        link = pending["links"][0]
        assert link["original_target_url"] == durable_url
        assert link["approval_status"] == "pending"
        assert "link-disposition" in uploaded.json()["artifacts"]

        approved = client.post(
            f"/api/jobs/{job_id}/links",
            json={"link_ids": [link["link_id"]]},
        )
        assert approved.status_code == 200, approved.text
        result = approved.json()
        assert result["link_disposition"]["approved_link_count"] == 1
        assert result["link_disposition"]["links"][0]["final_inserted_url"] == durable_url
        assert {"linked-document", "linked-pdf"} <= set(result["artifacts"])
        assert client.get(f"/api/jobs/{job_id}/artifacts/document").content == original_docx
        assert client.get(f"/api/jobs/{job_id}/artifacts/pdf").content == original_pdf

        linked_docx = client.get(f"/api/jobs/{job_id}/artifacts/linked-document").content
        with zipfile.ZipFile(io.BytesIO(linked_docx)) as archive:
            relationships = archive.read("word/_rels/document.xml.rels")
        assert durable_url.encode() in relationships
        linked_pdf = client.get(f"/api/jobs/{job_id}/artifacts/linked-pdf")
        assert linked_pdf.content.startswith(b"%PDF")
        assert len(PdfReader(io.BytesIO(linked_pdf.content)).pages) >= 1

        manifest = client.get(f"/api/jobs/{job_id}/artifacts/link-disposition").json()
        assert manifest["links"][0]["annotation_rectangle"] == [72.0, 690.0, 300.0, 710.0]
        assert manifest["links"][0]["original_target_url"] == durable_url
        with zipfile.ZipFile(
            io.BytesIO(client.get(f"/api/jobs/{job_id}/artifacts/final-package").content)
        ) as archive:
            names = set(archive.namelist())
        assert "review/private-link-disposition.json" in names
        assert "final/linked-document.docx" in names
        assert "final/linked-document.pdf" in names


def test_ineligible_or_undeclared_link_approval_is_rejected(tmp_path: Path) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_processor)
    with TestClient(create_app(manager=manager)) as client:
        job_id = _prepared(client)
        uploaded = client.post(
            f"/api/jobs/{job_id}/sources",
            data={"authority_id": AUTHORITY_ID},
            files={
                "source": (
                    "opinion.pdf",
                    _source_pdf("https://law.example/search?session=temporary"),
                    "application/pdf",
                )
            },
        )
        link = uploaded.json()["link_disposition"]["links"][0]
        rejected = client.post(f"/api/jobs/{job_id}/links", json={"link_ids": [link["link_id"]]})
        assert rejected.status_code == 400
        assert "eligible" in rejected.json()["error"].lower()
        undeclared = client.post(
            f"/api/jobs/{job_id}/links", json={"link_ids": ["link-not-declared"]}
        )
        assert undeclared.status_code == 400
        assert "declared" in undeclared.json()["error"].lower()
        state = client.get(f"/api/jobs/{job_id}").json()
        assert "linked-document" not in state["artifacts"]
        assert "linked-pdf" not in state["artifacts"]
