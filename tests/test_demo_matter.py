from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

from opencounsel.demo import build_synthetic_demo
from opencounsel.source.authority_package import extract_pdf_link_annotations


def test_demo_reconstructs_public_openai_filing_as_broken_word_input(
    tmp_path: Path,
) -> None:
    brief_path, record_path = build_synthetic_demo(tmp_path)
    display_brief = tmp_path / "openai-mtd-deformatted.docx"

    assert brief_path.name == "brief.docx"
    assert display_brief.is_file()
    assert brief_path.read_bytes() == display_brief.read_bytes()
    assert record_path.name == "record.zip"

    document = Document(display_brief)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    text = "\n".join(paragraphs)

    assert paragraphs[0] == (
        "UNITED STATES DISTRICT COURT\nSOUTHERN DISTRICT OF NEW YORK"
    )
    assert len(document.tables) == 1
    caption_text = "\n".join(cell.text for cell in document.tables[0].rows[0].cells)
    assert "THE NEW YORK TIMES COMPANY," in caption_text
    assert "Case No. 1:23-cv-11195 (SHS) (OTW)" in caption_text
    assert "MEMORANDUM OF LAW\nIN SUPPORT OF OPENAI" in caption_text
    assert "ORAL ARGUMENT REQUESTED" in caption_text
    assert "ECF" not in caption_text
    assert "I. INTRODUCTION" in text
    assert "III. LEGAL STANDARD" in text
    assert "IV. ARGUMENT" in text
    assert "Ashcroft v. Iqbal, 556 U.S. 662, 678 (2009)" in text
    assert "Authors Guild v. Google, Inc., 804 F.3d 202 (2d Cir. 2015)" in text
    assert (
        "Sony Corp. of America v. Universal City Studios, Inc., "
        "464 U.S. 417 (1984)"
        in text
    )

    assert "[TOC]" in paragraphs
    assert "[TOA]" in paragraphs
    assert not any("........" in paragraph for paragraph in paragraphs)

    body_runs = [
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    ]
    font_names = {run.font.name for run in body_runs if run.font.name}
    font_sizes = {
        round(run.font.size.pt)
        for run in body_runs
        if run.font.size is not None
    }
    assert len(font_names) >= 2
    assert len(font_sizes) >= 3
    introduction = document.paragraphs[paragraphs.index("I. INTRODUCTION")]
    assert introduction.style.name == "Normal"

    source_path = tmp_path / "demo-source.json"
    provenance = json.loads(source_path.read_text(encoding="utf-8"))
    assert provenance == {
        "case": "The New York Times Company v. Microsoft Corporation, et al.",
        "court": "S.D.N.Y.",
        "docket": "1:23-cv-11195-SHS-OTW",
        "document": 52,
        "filed": "2024-02-26",
        "source_pdf": (
            "https://cdn.openai.com/pdf/"
            "gov.uscourts.nysd.612697.52.0_1.pdf"
        ),
        "transformation": (
            "public filing text reconstructed as intentionally malformed DOCX; "
            "populated TOC and TOA omitted"
        ),
    }


def test_demo_caption_uses_a_fixed_bordered_word_table(tmp_path: Path) -> None:
    build_synthetic_demo(tmp_path)
    source = tmp_path / "openai-mtd-deformatted.docx"
    document = Document(source)
    table = document.tables[0]
    left, right = table.rows[0].cells

    assert table.autofit is False
    assert round(table.columns[0].width.inches, 2) == 3.75
    assert round(table.columns[1].width.inches, 2) == 2.75
    assert left.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.TOP
    assert right.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.TOP
    assert [paragraph.text for paragraph in left.paragraphs] == [
        "THE NEW YORK TIMES COMPANY,",
        "Plaintiff,",
        "v.",
        (
            "MICROSOFT CORPORATION, OPENAI, INC.,\n"
            "OPENAI LP, OPENAI GP, LLC, OPENAI, LLC,\n"
            "OPENAI OPCO LLC, OPENAI GLOBAL LLC,\n"
            "OAI CORPORATION, LLC, OPENAI\n"
            "HOLDINGS, LLC,"
        ),
        "Defendants.",
    ]
    assert [paragraph.alignment for paragraph in left.paragraphs] == [
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.CENTER,
    ]
    assert [paragraph.text for paragraph in right.paragraphs] == [
        "Case No. 1:23-cv-11195 (SHS) (OTW)",
        (
            "MEMORANDUM OF LAW\n"
            "IN SUPPORT OF OPENAI\n"
            "DEFENDANTS\N{RIGHT SINGLE QUOTATION MARK} MOTION TO\n"
            "DISMISS"
        ),
        "ORAL ARGUMENT REQUESTED",
    ]
    assert right.paragraphs[0].runs[0].font.size is not None
    assert right.paragraphs[0].runs[0].font.size.pt == 10
    assert all(
        paragraph.paragraph_format.space_after is not None
        for paragraph in (*left.paragraphs, *right.paragraphs)
    )
    assert all(
        paragraph.paragraph_format.first_line_indent is not None
        and paragraph.paragraph_format.first_line_indent.inches == 0
        and paragraph.paragraph_format.left_indent is not None
        and paragraph.paragraph_format.left_indent.inches == 0
        and paragraph.paragraph_format.right_indent is not None
        and paragraph.paragraph_format.right_indent.inches == 0
        for paragraph in (*left.paragraphs, *right.paragraphs)
    )

    with zipfile.ZipFile(source) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    }
    tables = root.xpath(".//w:body/w:tbl", namespaces=namespaces)
    assert len(tables) == 1
    table_xml = tables[0]
    assert table_xml.xpath(
        "./w:tblPr/w:tblLayout[@w:type='fixed']", namespaces=namespaces
    )
    assert table_xml.xpath(
        "./w:tblGrid/w:gridCol/@w:w", namespaces=namespaces
    ) == ["5400", "3960"]
    cells = table_xml.xpath("./w:tr/w:tc", namespaces=namespaces)
    assert len(cells) == 2
    assert [
        cells[0].xpath(
            f"./w:tcPr/w:tcBorders/w:{edge}/@w:val", namespaces=namespaces
        )[0]
        for edge in ("top", "right", "bottom", "left")
    ] == ["single", "single", "single", "nil"]
    assert [
        cells[1].xpath(
            f"./w:tcPr/w:tcBorders/w:{edge}/@w:val", namespaces=namespaces
        )[0]
        for edge in ("top", "right", "bottom", "left")
    ] == ["nil", "nil", "nil", "nil"]
    assert not table_xml.xpath(
        ".//w:tab | .//w:txbxContent | .//w:drawing | .//w:pict",
        namespaces=namespaces,
    )


def test_demo_includes_real_public_authority_source_copies(tmp_path: Path) -> None:
    build_synthetic_demo(tmp_path)

    authority_dir = tmp_path / "authority-sources"
    expected = {
        "ashcroft-v-iqbal.pdf",
        "authors-guild-v-google.pdf",
        "baker-v-selden.pdf",
        "sony-v-universal.pdf",
    }
    sources = list(authority_dir.glob("*.pdf"))
    assert {path.name for path in sources} == expected
    assert all(path.read_bytes().startswith(b"%PDF") for path in sources)

    iqbal = extract_pdf_link_annotations(
        authority_dir / "ashcroft-v-iqbal.pdf",
        authority_id="auth-public-iqbal",
    )
    links = iqbal["links"]
    assert isinstance(links, list)
    assert len(links) == 1
    assert links[0]["classification"] == "durable-candidate"
    assert links[0]["approval_status"] == "pending"
    assert links[0]["original_target_url"] == (
        "https://www.supremecourt.gov/opinions/boundvolumes/556bv.pdf"
    )
