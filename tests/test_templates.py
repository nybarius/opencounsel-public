from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document

from opencounsel.templates.registry import (
    TemplateManifestError,
    load_template_manifest,
    validate_word_template,
)


def _template_and_manifest(tmp_path: Path) -> tuple[Path, Path]:
    template = tmp_path / "appellate.docx"
    document = Document()
    document.add_heading("Blank heading", level=1)
    document.save(template)
    digest = hashlib.sha256(template.read_bytes()).hexdigest()
    manifest = tmp_path / "appellate.toml"
    manifest.write_text(
        f'''schema_version = 1
template_id = "ny-appellate-brief"
version = "2026.1"
template_sha256 = "{digest}"
slots = ["front-matter", "body"]

[[styles]]
semantic_role = "section-heading"
word_style_name = "Heading 1"
toc_level = 1
''',
        encoding="utf-8",
    )
    return template, manifest


def test_loads_and_validates_separate_template_manifest(tmp_path: Path) -> None:
    template, manifest_path = _template_and_manifest(tmp_path)

    manifest = load_template_manifest(manifest_path)
    result = validate_word_template(template, manifest)

    assert result.template_id == "ny-appellate-brief"
    assert result.bound_roles == ("section-heading",)
    assert result.template_sha256 == manifest.template_sha256


def test_rejects_hash_or_style_drift(tmp_path: Path) -> None:
    template, manifest_path = _template_and_manifest(tmp_path)
    manifest = load_template_manifest(manifest_path)

    changed = Document(template)
    changed.add_paragraph("Matter content must not enter a registered blank template.")
    changed.save(template)
    with pytest.raises(TemplateManifestError, match="hash"):
        validate_word_template(template, manifest)

    template, manifest_path = _template_and_manifest(tmp_path)
    text = manifest_path.read_text(encoding="utf-8").replace("Heading 1", "Missing Style")
    manifest_path.write_text(text, encoding="utf-8")
    missing_style_manifest = load_template_manifest(manifest_path)
    with pytest.raises(TemplateManifestError, match="missing required Word styles"):
        validate_word_template(template, missing_style_manifest)


@pytest.mark.parametrize(
    ("text", "message"),
    [
        ("not = [valid", "malformed"),
        (
            """schema_version = 3
template_id = "valid-id"
version = "1"
template_sha256 = "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa"
[[styles]]
semantic_role = "heading"
word_style_name = "Heading 1"
""",
            "unsupported",
        ),
    ],
)
def test_rejects_invalid_manifest(tmp_path: Path, text: str, message: str) -> None:
    path = tmp_path / "manifest.toml"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(TemplateManifestError, match=message):
        load_template_manifest(path)


def test_manifest_v2_pins_filing_profile_and_style_pack(tmp_path: Path) -> None:
    template, _ = _template_and_manifest(tmp_path)
    digest = hashlib.sha256(template.read_bytes()).hexdigest()
    manifest_path = tmp_path / "manifest-v2.toml"
    manifest_path.write_text(
        f'''schema_version = 2
template_id = "ny-appellate-brief"
version = "2026.1"
template_sha256 = "{digest}"
filing_profile_id = "ny-ad-appellant-brief"
style_pack_id = "opencounsel-clean-serif"
slots = ["body"]

[[styles]]
semantic_role = "section-heading"
word_style_name = "Heading 1"
toc_level = 1
''',
        encoding="utf-8",
    )

    manifest = load_template_manifest(manifest_path)

    assert manifest.filing_profile_id == "ny-ad-appellant-brief"
    assert manifest.style_pack_id == "opencounsel-clean-serif"


def test_manifest_rejects_unknown_fields(tmp_path: Path) -> None:
    _, manifest_path = _template_and_manifest(tmp_path)
    manifest_path.write_text(
        manifest_path.read_text(encoding="utf-8").replace(
            "\n[[styles]]",
            "\nunknown = true\n\n[[styles]]",
        ),
        encoding="utf-8",
    )

    with pytest.raises(TemplateManifestError, match="unknown fields"):
        load_template_manifest(manifest_path)


def test_manifest_boundary_rejects_invalid_shapes_and_values(tmp_path: Path) -> None:
    base = '''schema_version = 1
template_id = "valid-id"
version = "1"
template_sha256 = "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa"
slots = ["body"]

[[styles]]
semantic_role = "heading"
word_style_name = "Heading 1"
toc_level = 1
'''
    style_block = '''[[styles]]
semantic_role = "heading"
word_style_name = "Heading 1"
toc_level = 1
'''
    cases = [
        (base.replace('template_id = "valid-id"', 'template_id = "Invalid"'), "template_id"),
        (base.replace('version = "1"', 'version = " "'), "version"),
        (base.replace("a" * 64, "A" * 64), "SHA-256"),
        (base.replace('slots = ["body"]', 'slots = "body"'), "slots must be a list"),
        (base.replace(style_block, "styles = []\n"), "bind at least one style"),
        (base.replace(style_block, 'styles = ["Heading 1"]\n'), "style binding is malformed"),
        (base.replace("toc_level = 1", "toc_level = 1\nunknown = true"), "unknown fields"),
        (base.replace('word_style_name = "Heading 1"\n', ""), "style binding is malformed"),
        (base.replace('semantic_role = "heading"', 'semantic_role = ""'), "nonempty and unique"),
        (base + style_block, "nonempty and unique"),
        (
            base.replace('word_style_name = "Heading 1"', 'word_style_name = ""'),
            "must not be empty",
        ),
        (base.replace("toc_level = 1", "toc_level = 10"), "between 1 and 9"),
        (base.replace('slots = ["body"]', 'slots = ["body", "body"]'), "nonempty and unique"),
        (
            base.replace(
                "schema_version = 1",
                'schema_version = 2\nstyle_pack_id = "valid-pack"',
            ),
            "malformed",
        ),
        (
            base.replace(
                "schema_version = 1",
                'schema_version = 2\nfiling_profile_id = "Invalid"\nstyle_pack_id = "valid-pack"',
            ),
            "lowercase stable identifiers",
        ),
    ]

    for index, (text, message) in enumerate(cases):
        path = tmp_path / f"invalid-{index}.toml"
        path.write_text(text, encoding="utf-8")
        with pytest.raises(TemplateManifestError, match=message):
            load_template_manifest(path)


def test_manifest_boundary_rejects_non_toml_and_oversized_files(tmp_path: Path) -> None:
    with pytest.raises(TemplateManifestError, match="regular TOML"):
        load_template_manifest(tmp_path / "missing.toml")

    wrong_suffix = tmp_path / "manifest.txt"
    wrong_suffix.write_text("schema_version = 1", encoding="utf-8")
    with pytest.raises(TemplateManifestError, match="regular TOML"):
        load_template_manifest(wrong_suffix)

    oversized = tmp_path / "manifest.toml"
    oversized.write_text(" " * (64 * 1024 + 1), encoding="utf-8")
    with pytest.raises(TemplateManifestError, match="64 KiB"):
        load_template_manifest(oversized)
