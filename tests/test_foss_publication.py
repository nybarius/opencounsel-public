from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from pypdf import PdfReader, PdfWriter

from opencounsel.briefs.clean import clean_brief
from opencounsel.briefs.sources import SourceOverride
from opencounsel.contracts.models import (
    FrontMatterSource,
    ToaSourceEntry,
    TocSourceEntry,
)
from opencounsel.publication import foss
from opencounsel.publication.foss import publish_foss
from opencounsel.revisions import RevisionStore

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _synthetic_authority() -> ToaSourceEntry:
    citation = "Example v. Case, 1 A.D.3d 2"
    return ToaSourceEntry.model_validate(
        {
            "authority_id": f"fm-toa-{'1' * 64}",
            "category": "cases",
            "word_category": 1,
            "category_heading": "Cases",
            "display_name": citation,
            "short_name": "Case",
            "normalized_citation": citation,
            "locations": [
                {
                    "part": "footnote:1",
                    "paragraph_index": 0,
                    "start_offset": 0,
                    "end_offset": len(citation),
                    "original_text": citation,
                    "citation_type": "FullCaseCitation",
                }
            ],
        }
    )


def test_reference_pdf_render_preserves_the_original_docx(tmp_path: Path) -> None:
    source = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("Immutable original brief.")
    document.save(source)
    original = source.read_bytes()
    output = tmp_path / "original-brief.pdf"

    def fake_run(
        command: list[str],
        *,
        check: bool,
        capture_output: bool,
        text: bool,
        timeout: int,
    ) -> subprocess.CompletedProcess[str]:
        assert check is False
        assert capture_output is True
        assert text is True
        assert timeout == 120
        outdir = Path(command[command.index("--outdir") + 1])
        input_path = Path(command[-1])
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with (outdir / f"{input_path.stem}.pdf").open("wb") as stream:
            writer.write(stream)
        return subprocess.CompletedProcess(command, 0, "converted\n", "")

    page_count = foss.render_docx_reference_pdf(
        source,
        output,
        executable="libreoffice",
        runner=fake_run,
    )

    assert page_count == 1
    assert source.read_bytes() == original
    assert len(PdfReader(output).pages) == 1
    assert output.stat().st_mode & 0o777 == 0o600


def test_reference_pdf_render_requires_a_separate_new_output(tmp_path: Path) -> None:
    source = tmp_path / "brief.docx"
    Document().save(source)

    with pytest.raises(foss.FossPublicationError, match="separate PDF"):
        foss.render_docx_reference_pdf(source, source)

    existing = tmp_path / "existing.pdf"
    existing.write_bytes(b"existing")
    with pytest.raises(foss.FossPublicationError, match="already exists"):
        foss.render_docx_reference_pdf(source, existing)

    with pytest.raises(foss.FossPublicationError, match="directory is missing"):
        foss.render_docx_reference_pdf(source, tmp_path / "missing" / "output.pdf")


def test_reference_pdf_render_rejects_source_mutation(tmp_path: Path) -> None:
    source = tmp_path / "brief.docx"
    Document().save(source)
    output = tmp_path / "original-brief.pdf"

    def mutating_run(
        command: list[str],
        *,
        check: bool,
        capture_output: bool,
        text: bool,
        timeout: int,
    ) -> subprocess.CompletedProcess[str]:
        outdir = Path(command[command.index("--outdir") + 1])
        input_path = Path(command[-1])
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with (outdir / f"{input_path.stem}.pdf").open("wb") as stream:
            writer.write(stream)
        with input_path.open("ab") as stream:
            stream.write(b"mutated")
        return subprocess.CompletedProcess(command, 0, "converted\n", "")

    with pytest.raises(foss.FossPublicationError, match="changed its source"):
        foss.render_docx_reference_pdf(
            source,
            output,
            executable="libreoffice",
            runner=mutating_run,
        )

    assert not output.exists()


def test_rejects_pageref_in_terminal_front_matter(tmp_path: Path) -> None:
    path = tmp_path / "pageref.docx"
    document = Document()
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGEREF bookmark_1 \\h "
    document.add_paragraph().add_run()._r.append(instruction)
    document.save(path)

    with pytest.raises(foss.FossPublicationError, match="still contains PAGEREF"):
        foss._validate_static_front_matter(
            path,
            FrontMatterSource(brief_sha256="0" * 64),
        )


def test_rejects_incomplete_terminal_front_matter(tmp_path: Path) -> None:
    path = tmp_path / "incomplete.docx"
    Document().save(path)
    front_matter = FrontMatterSource(
        brief_sha256="0" * 64,
        toc=(
            TocSourceEntry(
                entry_id=f"fm-toc-{'1' * 64}",
                level=1,
                heading="ARGUMENT",
                paragraph_index=0,
                basis="paragraph-style",
            ),
        ),
    )

    with pytest.raises(foss.FossPublicationError, match="entry counts are incomplete"):
        foss._validate_static_front_matter(path, front_matter)


def test_reports_unreadable_first_pass_page_sources(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def unreadable(_path: Path) -> None:
        raise OSError("synthetic unreadable PDF")

    monkeypatch.setattr(foss, "PdfReader", unreadable)
    authority = _synthetic_authority()

    with pytest.raises(foss.FossPublicationError, match="TOA page references"):
        foss._resolve_toa_occurrence_pages(Path("unused.pdf"), [authority])
    with pytest.raises(foss.FossPublicationError, match="first-pass source pages"):
        foss._resolve_source_occurrence_pages(Path("unused.pdf"), [authority])


def test_toa_page_parser_does_not_treat_body_word_as_roman_page(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    citation = "42 U.S.C. § 1983"
    authority = ToaSourceEntry.model_validate(
        {
            "authority_id": f"fm-toa-{'2' * 64}",
            "category": "statutes",
            "word_category": 2,
            "category_heading": "Statutes",
            "display_name": citation,
            "short_name": citation,
            "normalized_citation": citation,
            "locations": [
                {
                    "part": "document",
                    "paragraph_index": 1,
                    "start_offset": 0,
                    "end_offset": len(citation),
                    "original_text": citation,
                    "citation_type": "FullLawCitation",
                }
            ],
        }
    )
    pages = [
        f"ii TABLE OF AUTHORITIES STATUTES {citation}........1",
        f"1 The remedy under {citation} confirms the rule.",
    ]
    monkeypatch.setattr(
        foss,
        "PdfReader",
        lambda _path: SimpleNamespace(
            pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in pages]
        ),
    )

    assert foss._resolve_toa_occurrence_pages(Path("unused.pdf"), [authority]) == {
        authority.authority_id: ("1",)
    }


def test_rejects_empty_or_missing_first_pass_source_occurrences(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        foss,
        "PdfReader",
        lambda _path: SimpleNamespace(pages=[]),
    )
    authority = _synthetic_authority()
    with pytest.raises(foss.FossPublicationError, match="has no pages"):
        foss._resolve_source_occurrence_pages(Path("unused.pdf"), [authority])

    monkeypatch.setattr(
        foss,
        "PdfReader",
        lambda _path: SimpleNamespace(
            pages=[SimpleNamespace(extract_text=lambda: "1 Body without the citation")]
        ),
    )
    location = authority.locations[0].model_copy(update={"original_text": ""})
    authority = authority.model_copy(update={"locations": (location,)})
    with pytest.raises(foss.FossPublicationError, match="occurrence is empty"):
        foss._resolve_source_occurrence_pages(Path("unused.pdf"), [authority])

    location = location.model_copy(update={"original_text": "missing citation"})
    authority = authority.model_copy(update={"locations": (location,)})
    with pytest.raises(foss.FossPublicationError, match="source occurrence"):
        foss._resolve_source_occurrence_pages(Path("unused.pdf"), [authority])


def test_resolves_footnote_toa_pages_from_first_pass_body_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pages = [
        "i TABLE OF CONTENTS",
        "ii TABLE OF AUTHORITIES Example v. Case, 1 A.D.3d 2 Error: Reference source not found",
        "QUESTION PRESENTED",
        "2 Body text.",
        "3 Note text cites Example v. Case, 1 A.D.3d 2.",
    ]
    monkeypatch.setattr(
        foss,
        "PdfReader",
        lambda _path: SimpleNamespace(
            pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in pages]
        ),
    )
    authority = ToaSourceEntry.model_validate(
        {
            "authority_id": f"fm-toa-{'1' * 64}",
            "category": "cases",
            "word_category": 1,
            "category_heading": "Cases",
            "display_name": "Example v. Case, 1 A.D.3d 2",
            "short_name": "Case",
            "normalized_citation": "Example v. Case, 1 A.D.3d 2",
            "italic_spans": [],
            "existing_targets": [],
            "locations": [
                {
                    "part": "footnote:1",
                    "paragraph_index": 0,
                    "start_offset": 0,
                    "end_offset": len("1 A.D.3d 2"),
                    "original_text": "1 A.D.3d 2",
                    "citation_type": "FullCaseCitation",
                }
            ],
        }
    )

    assert foss._resolve_source_occurrence_pages(Path("unused.pdf"), [authority]) == {
        f"fm-toa-{'1' * 64}": ("3",)
    }


def test_resolves_toc_pages_from_body_headings_not_front_matter_echo(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pages = [
        "i TABLE OF CONTENTS\nARGUMENT 3\nI. THE AUTHORITIES CONTROL 3",
        "1\nPRELIMINARY STATEMENT\nSynthetic text.",
        "2\nSUMMARY OF ARGUMENT\nSynthetic summary.",
        "3\nARGUMENT\nI. THE AUTHORITIES CONTROL\nSynthetic argument.",
    ]
    monkeypatch.setattr(
        foss,
        "PdfReader",
        lambda _path: SimpleNamespace(
            pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in pages]
        ),
    )
    entries = [
        TocSourceEntry(
            entry_id=f"fm-toc-{'1' * 64}",
            level=1,
            heading="SUMMARY OF ARGUMENT",
            paragraph_index=1,
            basis="paragraph-style",
        ),
        TocSourceEntry(
            entry_id=f"fm-toc-{'2' * 64}",
            level=1,
            heading="ARGUMENT",
            paragraph_index=2,
            basis="paragraph-style",
        ),
        TocSourceEntry(
            entry_id=f"fm-toc-{'3' * 64}",
            level=2,
            heading="THE AUTHORITIES CONTROL",
            number_label="I.",
            paragraph_index=3,
            basis="paragraph-style",
        ),
    ]

    assert foss._resolve_toc_pages(Path("unused.pdf"), entries) == {
        f"fm-toc-{'1' * 64}": "2",
        f"fm-toc-{'2' * 64}": "3",
        f"fm-toc-{'3' * 64}": "3",
    }


@pytest.mark.parametrize(
    "pages",
    [
        [
            "i\nTABLE OF CONTENTS",
            "ii\nTABLE OF AUTHORITIES",
            "I. INTRODUCTION\nFirst body page without a visible footer.",
            "2\nII. BACKGROUND",
        ],
        [
            "TABLE OF CONTENTS\ni",
            "TABLE OF AUTHORITIES\nii",
            "I. INTRODUCTION\nFirst body page without a visible footer.",
            "II. BACKGROUND\n2",
        ],
    ],
    ids=("page-label-first", "page-label-last"),
)
def test_body_start_does_not_treat_roman_heading_as_page_label(
    pages: list[str],
) -> None:
    assert foss._body_start_page(pages) == 2


def test_rejects_toc_heading_repeated_on_distinct_body_pages(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pages = [
        "i\nTABLE OF CONTENTS",
        "1\nARGUMENT\nFirst occurrence.",
        "2\nARGUMENT\nSecond occurrence.",
    ]
    monkeypatch.setattr(
        foss,
        "PdfReader",
        lambda _path: SimpleNamespace(
            pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in pages]
        ),
    )
    entry = TocSourceEntry(
        entry_id=f"fm-toc-{'4' * 64}",
        level=1,
        heading="ARGUMENT",
        paragraph_index=1,
        basis="paragraph-style",
    )

    with pytest.raises(foss.FossPublicationError, match="unambiguous TOC heading"):
        foss._resolve_toc_pages(Path("unused.pdf"), [entry])


def test_long_case_toa_entry_soft_breaks_before_reporter() -> None:
    case_name = "Uygur v. Superior Walls of Hudson Valley, Inc."
    display = f"{case_name}, 35 A.D.3d 447 (2d Dep't 2006)"

    paragraph = foss._toa_static_entry_paragraph(
        display,
        ((0, len(case_name)),),
        [("32", None)],
        {},
        relationship_id=None,
    )

    assert len(paragraph.xpath(".//w:br", namespaces={"w": W_NS})) == 1
    assert "".join(paragraph.xpath(".//w:t/text()", namespaces={"w": W_NS})) == (
        f"{display}32"
    )


def test_foss_publication_compiles_visible_linked_front_matter_and_pdf(
    tmp_path: Path,
    roa_package: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    confined_roa = inbox / "synthetic-roa.zip"
    shutil.copyfile(roa_package, confined_roa)
    source = inbox / "brief.docx"
    document = Document()
    first_heading = document.add_heading(level=1)
    first_heading.add_run().add_break()
    first_heading.add_run("PRELIMINARY STATEMENT")
    document.add_paragraph("Synthetic opening.")
    document.add_heading("I. THE AUTHORITIES CONTROL", level=2)
    document.add_heading(
        "A. A LONG SUBPOINT HEADING THAT WRAPS WITH ITS TEXT ALIGNED "
        "UNDER THE FIRST-LINE TEXT",
        level=3,
    )
    document.add_paragraph(
        "Bell Atlantic Corp. v. Twombly, 550 U.S. 544, 570 (2007), and "
        "42 U.S.C. § 1983 control (R. 3-4)."
    )
    document.save(source)
    work = tmp_path / "work"
    store = RevisionStore(work, inbox)
    cleaned = clean_brief(
        store,
        source,
        "ny-ad-appellant-brief",
        roa_package_path=confined_roa,
        authority_overrides=(
            SourceOverride(
                "Bell Atlantic Corp. v. Twombly, 550 U.S. 544 (2007)",
                "https://www.courtlistener.com/opinion/1444/"
                "bell-atlantic-corp-v-twombly/",
            ),
        ),
    )
    monkeypatch.setattr(
        foss,
        "_resolve_toc_pages",
        lambda _pdf, entries: {entry.entry_id: "1" for entry in entries},
    )
    monkeypatch.setattr(
        foss,
        "_resolve_toa_occurrence_pages",
        lambda _pdf, authorities: {
            authority.authority_id: tuple("1" for _ in authority.locations)
            for authority in authorities
        },
    )

    def fake_run(
        command: list[str],
        *,
        check: bool,
        capture_output: bool,
        text: bool,
        timeout: int,
    ) -> subprocess.CompletedProcess[str]:
        assert check is False
        assert capture_output is True
        assert text is True
        assert timeout > 0
        if "--version" in command:
            return subprocess.CompletedProcess(command, 0, "LibreOffice 26.2.0\n", "")
        outdir = Path(command[command.index("--outdir") + 1])
        input_path = Path(command[-1])
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with (outdir / f"{input_path.stem}.pdf").open("wb") as stream:
            writer.write(stream)
        return subprocess.CompletedProcess(command, 0, "converted\n", "")

    result = publish_foss(
        store,
        cleaned.process.process_id,
        executable="libreoffice",
        runner=fake_run,
    )

    delivery = work / "deliveries" / cleaned.process.process_id
    published = delivery / "published.docx"
    pdf = delivery / "published.pdf"
    manifest_path = delivery / "publication.json"
    assert published.is_file()
    assert pdf.is_file()
    assert manifest_path.is_file()
    assert result.engine == "libreoffice"
    assert result.page_count == 1
    assert result.toc_entry_count >= 3
    assert result.toa_authority_count == 2

    rendered = Document(published)
    texts = [paragraph.text for paragraph in rendered.paragraphs]
    assert "[TOC]" not in texts
    assert "[TOA]" not in texts
    assert "Cases" in texts
    assert "Statutes" in texts
    assert any("THE AUTHORITIES CONTROL" in text for text in texts)
    assert any(text == "I.\tTHE AUTHORITIES CONTROL\t1" for text in texts)
    assert any(
        text
        == (
            "A.\tA LONG SUBPOINT HEADING THAT WRAPS WITH ITS TEXT ALIGNED "
            "UNDER THE FIRST-LINE TEXT\t1"
        )
        for text in texts
    )
    assert any("Bell Atlantic Corp. v. Twombly" in text for text in texts)

    with zipfile.ZipFile(published) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    with zipfile.ZipFile(delivery / "corrected.docx") as archive:
        corrected_root = etree.fromstring(archive.read("word/document.xml"))
    instructions = "".join(
        root.xpath(".//w:instrText/text()", namespaces={"w": W_NS})
    )
    assert " PAGEREF " not in instructions
    assert root.xpath(".//w:bookmarkStart", namespaces={"w": W_NS})
    assert root.xpath(".//w:hyperlink[@w:anchor]", namespaces={"w": W_NS})
    assert len(root.xpath(".//w:sectPr", namespaces={"w": W_NS})) == len(
        corrected_root.xpath(".//w:sectPr", namespaces={"w": W_NS})
    )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["input_docx_sha256"] == cleaned.process.corrected_sha256
    assert manifest["output_docx_sha256"] == result.output_docx_sha256
    assert manifest["output_pdf_sha256"] == result.output_pdf_sha256
    assert publish_foss(
        store,
        cleaned.process.process_id,
        executable="unused-on-idempotent-load",
        runner=fake_run,
    ) == result


@pytest.mark.skipif(
    shutil.which("libreoffice") is None and shutil.which("soffice") is None,
    reason="LibreOffice is not installed",
)
def test_real_libreoffice_updates_pageref_front_matter(
    tmp_path: Path,
    roa_package: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    confined_roa = inbox / "synthetic-roa.zip"
    shutil.copyfile(roa_package, confined_roa)
    source = inbox / "brief.docx"
    document = Document()
    document.add_paragraph("PRELIMINARY STATEMENT")
    document.add_paragraph("Synthetic opening.")
    heading = document.add_paragraph("I. THE AUTHORITIES CONTROL")
    heading.runs[0].bold = True
    document.add_paragraph("42 U.S.C. § 1983 controls (R. 3).")
    document.add_paragraph("Section 1983 remains controlling under 42 U.S.C. § 1983.")
    document.save(source)
    store = RevisionStore(tmp_path / "work", inbox)
    cleaned = clean_brief(
        store,
        source,
        "ny-ad-appellant-brief",
        roa_package_path=confined_roa,
    )

    executable = shutil.which("libreoffice") or shutil.which("soffice")
    assert executable is not None
    result = publish_foss(
        store,
        cleaned.process.process_id,
        executable=executable,
    )

    pdf = store.deliveries / cleaned.process.process_id / "published.pdf"
    published = store.deliveries / cleaned.process.process_id / "published.docx"
    text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    assert result.page_count >= 3
    assert "[TOC]" not in text
    assert "[TOA]" not in text
    assert "THE AUTHORITIES CONTROL" in text
    assert "1, 1" not in text
    first_page = PdfReader(pdf).pages[0].extract_text() or ""
    assert "THE AUTHORITIES CONTROL" in first_page
    assert " 0" not in first_page
    with zipfile.ZipFile(published) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    pageref_paragraphs = root.xpath(
        ".//w:p[.//w:instrText[contains(., 'PAGEREF')]]",
        namespaces={"w": W_NS},
    )
    assert not pageref_paragraphs
