from __future__ import annotations

import io
import time
from pathlib import Path

from docx import Document
from starlette.testclient import TestClient

from opencounsel.briefs.audit import audit_brief
from opencounsel.web.app import create_app
from opencounsel.web.jobs import JobContext, JobManager, ProgressUpdate


def _brief_bytes(text: str) -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph(text)
    document.save(stream)
    return stream.getvalue()


def _completed(client: TestClient, job_id: str) -> dict[str, object]:
    for _ in range(100):
        payload = client.get(f"/api/jobs/{job_id}").json()
        if payload["status"] in {"completed", "failed"}:
            return payload
        time.sleep(0.01)
    raise AssertionError("job did not finish")


def test_no_record_mode_accepts_brief_without_creating_record_artifacts(
    tmp_path: Path,
) -> None:
    contexts: list[JobContext] = []

    def processor(context: JobContext, progress: ProgressUpdate) -> dict[str, object]:
        contexts.append(context)
        progress("validating", "Validating brief")
        audit = audit_brief(context.brief_path, context.roa_path)
        export = context.job_dir / "export"
        export.mkdir(mode=0o700)
        document = export / "published.docx"
        document.write_bytes(b"immutable prepared docx")
        return {
            "summary": {
                "record_citation_count": len(audit.record_citations),
                "resolved_record_citation_count": 0,
                "unresolved_record_citation_count": audit.unresolved_record_citation_count,
            },
            "sources": [],
            "audit": {},
            "artifacts": {
                "document": {
                    "path": "export/published.docx",
                    "filename": "published.docx",
                    "media_type": (
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                }
            },
        }

    manager = JobManager(tmp_path / "jobs", processor=processor)
    with TestClient(create_app(manager=manager)) as client:
        response = client.post(
            "/api/jobs",
            data={
                "profile": "fed-sdny-edny-motion-memorandum",
                "record_mode": "no-record",
            },
            files={"brief": ("memorandum.docx", _brief_bytes("No record is cited."))},
        )
        assert response.status_code == 202, response.text
        result = _completed(client, response.json()["job_id"])

    assert result["status"] == "completed"
    assert result["record_mode"] == "no-record"
    assert result["roa_name"] is None
    assert result["summary"]["record_citation_count"] == 0
    assert "record" not in result["artifacts"]
    assert "roa-package" not in result["artifacts"]
    assert contexts[0].roa_path is None
    assert not (contexts[0].inbox_dir / "record.zip").exists()


def test_record_mode_still_requires_and_preserves_the_roa_path(
    tmp_path: Path,
    roa_package: Path,
) -> None:
    contexts: list[JobContext] = []

    def processor(context: JobContext, progress: ProgressUpdate) -> dict[str, object]:
        contexts.append(context)
        return {"summary": {}, "sources": [], "audit": {}, "artifacts": {}}

    manager = JobManager(tmp_path / "jobs", processor=processor)
    with TestClient(create_app(manager=manager)) as client:
        missing = client.post(
            "/api/jobs",
            data={
                "profile": "fed-sdny-edny-motion-memorandum",
                "record_mode": "record",
            },
            files={"brief": ("memorandum.docx", _brief_bytes("Record-backed filing."))},
        )
        assert missing.status_code == 400
        assert "record" in missing.json()["error"].lower()

        accepted = client.post(
            "/api/jobs",
            data={
                "profile": "fed-sdny-edny-motion-memorandum",
                "record_mode": "record",
            },
            files={
                "brief": ("memorandum.docx", _brief_bytes("Record-backed filing.")),
                "roa": ("record.zip", roa_package.read_bytes(), "application/zip"),
            },
        )
        assert accepted.status_code == 202
        result = _completed(client, accepted.json()["job_id"])

    assert result["status"] == "completed"
    assert result["record_mode"] == "record"
    assert contexts[0].roa_path is not None
    assert contexts[0].roa_path.is_file()
