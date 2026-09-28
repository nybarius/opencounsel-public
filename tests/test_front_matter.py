from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from opencounsel.briefs.audit import audit_brief
from opencounsel.briefs.docx import NS, W_NS, inspect_brief_docx
from opencounsel.briefs.front_matter import build_front_matter_source
from opencounsel.briefs.front_matter_docx import (
    FrontMatterProjectionError,
    project_front_matter_fields,
)
from opencounsel.briefs.ooxml import XML_SPACE
from opencounsel.briefs.word_fields import (
    ta_field_runs,
    tc_field_runs,
    toa_field_runs,
    toc_field_runs,
)
from opencounsel.contracts.schema import validate_contract


def _brief(path: Path) -> Path:
    document = Document()
    document.add_paragraph("[TOC]")
    document.add_paragraph("[TOA]")
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph(
        "Bell Atlantic Corp. v. Twombly, 550 U.S. 544, 570 (2007), controls. "
        "Twombly, 550 U.S. at 555, confirms the rule. "
        "See 28 U.S.C. § 1332; CPLR 3211(a)(7); and 17 C.F.R. § 240.10b-5."
    )
    document.save(path)
    return path


def _instructions(runs: tuple[etree._Element, ...]) -> str:
    return "".join(
        str(value)
        for run in runs
        for value in run.xpath("./w:instrText/text()", namespaces=NS)
    )


def _document_xml(path: Path) -> etree._Element:
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def _rewrite_document_xml(source: Path, output: Path, body: bytes) -> Path:
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(output, "w") as outgoing:
        for member in incoming.infolist():
            outgoing.writestr(
                member,
                body if member.filename == "word/document.xml" else incoming.read(member),
            )
    return output


def test_compiles_stable_page_free_toc_and_canonical_toa(tmp_path: Path) -> None:
    audit = audit_brief(_brief(tmp_path / "brief.docx"))

    first = build_front_matter_source(audit)
    second = build_front_matter_source(audit)

    assert first == second
    assert first.brief_sha256 == audit.brief_sha256
    assert len(first.toc) == 1
    assert re.fullmatch(r"fm-toc-[0-9a-f]{64}", first.toc[0].entry_id)
    assert first.toc[0].heading == "ARGUMENT"
    assert first.toc[0].level == 1
    assert not hasattr(first.toc[0], "page")

    by_category = {entry.category: entry for entry in first.toa}
    assert set(by_category) == {"cases", "statutes", "regulations"}
    case = by_category["cases"]
    assert re.fullmatch(r"fm-toa-[0-9a-f]{64}", case.authority_id)
    assert case.word_category == 1
    assert case.display_name == (
        "Bell Atlantic Corp. v. Twombly, 550 U.S. 544 (2007)"
    )
    assert case.short_name == "Twombly"
    assert case.italic_spans == ((0, len("Bell Atlantic Corp. v. Twombly")),)
    assert len(case.locations) == 2
    assert {location.citation_type for location in case.locations} == {
        "FullCaseCitation",
        "ShortCaseCitation",
    }
    assert not hasattr(case.locations[0], "page")

    statutes = by_category["statutes"]
    assert statutes.word_category == 2
    assert {entry.normalized_citation for entry in first.toa if entry.category == "statutes"} == {
        "28 U.S.C. § 1332",
        "CPLR § 3211(a)(7)",
    }
    regulation = by_category["regulations"]
    assert regulation.word_category == 6
    assert regulation.normalized_citation == "17 C.F.R. § 240.10b-5"
    validate_contract(
        "front-matter-source.schema.json", first.model_dump(mode="json")
    )


def test_toa_canonicalizes_spelled_out_federal_code_section(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(
        "The claim is untimely under 17 U.S.C. section 507(b)."
    )
    source = tmp_path / "brief.docx"
    document.save(source)

    front_matter = build_front_matter_source(audit_brief(source))

    assert len(front_matter.toa) == 1
    authority = front_matter.toa[0]
    assert authority.normalized_citation == "17 U.S.C. § 507(b)"
    assert authority.display_name == "17 U.S.C. § 507(b)"
    assert authority.locations[0].original_text == "17 U.S.C. section 507(b)"


def test_front_matter_locations_preserve_full_eyecite_spans(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(
        "Polar International Brokerage Corp. v. Richman, "
        "32 A.D.3d 717, 719 (2d Dep't 2006)."
    )
    document.add_paragraph(
        "Polar International, 32 A.D.3d at 719-20. Id. at 720."
    )
    source = tmp_path / "ranged-pincites.docx"
    document.save(source)

    audit = audit_brief(source)
    semantic = build_front_matter_source(audit)
    inspected = inspect_brief_docx(source)
    locations = semantic.toa[0].locations

    assert [location.original_text for location in locations] == [
        "32 A.D.3d 717",
        "32 A.D.3d at 719-20",
        "Id. at 720",
    ]
    assert all(
        inspected.paragraphs[location.paragraph_index].text[
            location.start_offset : location.end_offset
        ]
        == location.original_text
        for location in locations
    )


def test_builds_native_word_fields_with_case_italics_and_safe_quoting() -> None:
    long_form = 'Example "Quoted" Corp. v. Defendant, 123 F.3d 456 (2d Cir. 2020)'
    case_name = 'Example "Quoted" Corp. v. Defendant'
    ta = ta_field_runs(
        long_form,
        "Example",
        1,
        italic_spans=((0, len(case_name)),),
        case_soft_break_threshold=None,
    )

    instruction = _instructions(ta)
    assert instruction == (
        ' TA \\l "Example ""Quoted"" Corp. v. Defendant, 123 F.3d 456 '
        '(2d Cir. 2020)" \\s "Example" \\c 1 '
    )
    assert any(
        run.xpath("./w:rPr/w:i", namespaces=NS)
        and case_name.replace('"', '""')
        in "".join(run.xpath("./w:instrText/text()", namespaces=NS))
        for run in ta
    )
    assert not any(run.xpath(".//w:vanish", namespaces=NS) for run in ta)

    assert _instructions(tc_field_runs("ARGUMENT", 1)) == (
        ' TC "ARGUMENT" \\f O \\l "1" '
    )
    assert _instructions(toc_field_runs("[TOC]")) == ' TOC \\f O \\h \\z '
    assert _instructions(toa_field_runs("[TOA]")) == ' TOA \\h \\p '


def test_word_field_contract_rejects_invalid_values_and_handles_long_cases() -> None:
    with pytest.raises(ValueError, match="must not be empty"):
        ta_field_runs("", "Example", 1)
    with pytest.raises(ValueError, match="must not be empty"):
        ta_field_runs("Example", "", 1)
    with pytest.raises(ValueError, match="between 1 and 16"):
        ta_field_runs("Example", "Example", 17)
    with pytest.raises(ValueError, match="must not be empty"):
        tc_field_runs("", 1)
    with pytest.raises(ValueError, match="between 1 and 9"):
        tc_field_runs("ARGUMENT", 0)
    for identifier in ("", "AB", "1", "é"):
        with pytest.raises(ValueError, match="one ASCII letter"):
            tc_field_runs("ARGUMENT", 1, identifier=identifier)

    case_name = (
        "The Very Long Synthetic Corporate Appellant With Several Words "
        "v. Synthetic Corporate Appellee"
    )
    long_form = f"{case_name}, 123 F.4th 456 (2d Cir. 2026)"
    runs = ta_field_runs(
        long_form,
        "Synthetic Corporate Appellee",
        1,
        italic_spans=((-4, 5), (3, len(case_name)), (999, 1_000)),
        case_soft_break_threshold=20,
    )
    assert "\n" in _instructions(runs)
    assert any(
        run.xpath("./w:rPr/w:i", namespaces=NS)
        and case_name in "".join(run.xpath("./w:instrText/text()", namespaces=NS))
        for run in runs
    )

    no_comma = "A" * 100
    assert "\n" not in _instructions(
        ta_field_runs(
            no_comma,
            "A",
            1,
            italic_spans=((0, 10),),
            case_soft_break_threshold=20,
        )
    )
    trailing_comma = f"{'B' * 100},"
    assert "\n" not in _instructions(
        ta_field_runs(
            trailing_comma,
            "B",
            1,
            italic_spans=((0, 10),),
            case_soft_break_threshold=20,
        )
    )
    newline_after_comma = f"{'C' * 100},\n123 F.4th 456"
    assert "\n123" in _instructions(
        ta_field_runs(
            newline_after_comma,
            "C",
            1,
            italic_spans=((0, 10),),
            case_soft_break_threshold=20,
        )
    )

    result_runs = toc_field_runs(" [TOC] ")
    result_node = result_runs[3].find(f"{{{W_NS}}}t")
    assert result_node is not None
    assert result_node.get(XML_SPACE) == "preserve"


def test_maps_constitutions_and_treatises_to_native_word_categories(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("[TOA]")
    document.add_paragraph(
        "U.S. Const. amend. XIV; Restatement (Second) of Contracts § 90."
    )
    source = tmp_path / "categories.docx"
    document.save(source)

    semantic = build_front_matter_source(audit_brief(source))
    assert [(item.category, item.word_category) for item in semantic.toa] == [
        ("treatises", 5),
        ("constitutional-provisions", 7),
    ]

    output = tmp_path / "projected.docx"
    project_front_matter_fields(source, output, semantic)
    instructions = "".join(
        _document_xml(output).xpath(".//w:instrText/text()", namespaces=NS)
    )
    assert "\\c 5" in instructions
    assert "\\c 7" in instructions


def test_projects_front_matter_fields_without_changing_visible_text(tmp_path: Path) -> None:
    source = _brief(tmp_path / "brief.docx")
    original = source.read_bytes()
    semantic = build_front_matter_source(audit_brief(source))
    output = tmp_path / "projected.docx"

    result = project_front_matter_fields(source, output, semantic)

    assert source.read_bytes() == original
    assert result.input_sha256 == semantic.brief_sha256
    assert result.toc_marker_candidate_count == 1
    assert result.toc_marker_inserted_count == 1
    assert result.toa_marker_candidate_count == 5
    assert result.toa_marker_inserted_count == 5
    assert result.toc_field_inserted_count == 1
    assert result.toa_field_inserted_count == 1
    assert result.review_count == 0
    assert [paragraph.text for paragraph in inspect_brief_docx(source).paragraphs] == [
        paragraph.text for paragraph in inspect_brief_docx(output).paragraphs
    ]

    root = _document_xml(output)
    instructions = "".join(root.xpath(".//w:instrText/text()", namespaces=NS))
    assert instructions.count(" TC ") == 1
    assert instructions.count(" TA ") == 5
    assert " TOC " in instructions
    assert " TOA " in instructions
    assert not root.xpath(".//w:vanish", namespaces=NS)
    assert root.xpath(
        './/w:r[w:rPr/w:i]/w:instrText[contains(., "Bell Atlantic Corp. v. Twombly")]',
        namespaces=NS,
    )


def test_projection_is_idempotent_for_existing_fields(tmp_path: Path) -> None:
    source = _brief(tmp_path / "brief.docx")
    first_plan = build_front_matter_source(audit_brief(source))
    first_output = tmp_path / "first.docx"
    project_front_matter_fields(source, first_output, first_plan)
    second_plan = build_front_matter_source(audit_brief(first_output))
    second_output = tmp_path / "second.docx"

    result = project_front_matter_fields(first_output, second_output, second_plan)

    assert result.toc_marker_inserted_count == 0
    assert result.toa_marker_inserted_count == 0
    assert result.existing_marker_count == 8
    root = _document_xml(second_output)
    instructions = "".join(root.xpath(".//w:instrText/text()", namespaces=NS))
    assert instructions.count(" TC ") == 1
    assert instructions.count(" TA ") == 5
    assert instructions.count(" TOC ") == 1
    assert instructions.count(" TOA ") == 1


def test_projection_rejects_mismatched_or_overwriting_requests(tmp_path: Path) -> None:
    source = _brief(tmp_path / "brief.docx")
    semantic = build_front_matter_source(audit_brief(source))

    mismatched = semantic.model_copy(update={"brief_sha256": "0" * 64})
    with pytest.raises(FrontMatterProjectionError, match="does not match"):
        project_front_matter_fields(source, tmp_path / "wrong-hash.docx", mismatched)
    with pytest.raises(FrontMatterProjectionError, match="write a new DOCX"):
        project_front_matter_fields(source, source, semantic)

    existing = tmp_path / "existing.docx"
    existing.write_bytes(b"occupied")
    with pytest.raises(FrontMatterProjectionError, match="already exists"):
        project_front_matter_fields(source, existing, semantic)


def test_projection_rejects_invalid_semantic_locations(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("[TOA]")
    document.add_paragraph("See 28 U.S.C. § 1332.")
    source = tmp_path / "brief.docx"
    document.save(source)
    semantic = build_front_matter_source(audit_brief(source))
    entry = semantic.toa[0]
    location = entry.locations[0]

    invalid_plans = (
        (
            "duplicate correction ids",
            semantic.model_copy(update={"toa": (entry, entry)}),
        ),
        (
            "missing paragraph",
            semantic.model_copy(
                update={
                    "toa": (
                        entry.model_copy(
                            update={
                                "locations": (
                                    location.model_copy(update={"paragraph_index": 999}),
                                )
                            }
                        ),
                    )
                }
            ),
        ),
        (
            "outside the paragraph",
            semantic.model_copy(
                update={
                    "toa": (
                        entry.model_copy(
                            update={
                                "locations": (
                                    location.model_copy(update={"end_offset": 10_000}),
                                )
                            }
                        ),
                    )
                }
            ),
        ),
        (
            "no longer matches",
            semantic.model_copy(
                update={
                    "toa": (
                        entry.model_copy(
                            update={
                                "locations": (
                                    location.model_copy(update={"original_text": "wrong"}),
                                )
                            }
                        ),
                    )
                }
            ),
        ),
    )
    for index, (message, plan) in enumerate(invalid_plans):
        with pytest.raises(FrontMatterProjectionError, match=message):
            project_front_matter_fields(source, tmp_path / f"invalid-{index}.docx", plan)


@pytest.mark.parametrize(
    ("mode", "expected_note"),
    [
        ("tracked-slot", "tracked change or hyperlink"),
        ("slot-hyperlink", "tracked change or hyperlink"),
        ("authority-field", "existing Word field"),
        ("authority-smart-tag", "unsupported Word wrapper"),
        ("authority-complex-run", "complex Word run"),
        ("authority-partial-hyperlink", "inside a hyperlink"),
    ],
)
def test_projection_abstains_on_unsafe_word_structures(
    tmp_path: Path,
    mode: str,
    expected_note: str,
) -> None:
    document = Document()
    document.add_paragraph("[TOA]")
    document.add_paragraph("See 28 U.S.C. § 1332 extra text.")
    plain = tmp_path / f"{mode}-plain.docx"
    document.save(plain)
    root = _document_xml(plain)
    paragraphs = root.xpath(".//w:body/w:p", namespaces=NS)
    paragraph = paragraphs[0] if mode in {"tracked-slot", "slot-hyperlink"} else paragraphs[1]
    run = paragraph.find(f"{{{W_NS}}}r")
    assert run is not None

    if mode == "authority-complex-run":
        etree.SubElement(run, f"{{{W_NS}}}tab")
    else:
        paragraph.remove(run)
        wrapper_name = {
            "tracked-slot": "ins",
            "slot-hyperlink": "hyperlink",
            "authority-field": "fldSimple",
            "authority-smart-tag": "smartTag",
            "authority-partial-hyperlink": "hyperlink",
        }[mode]
        wrapper = etree.SubElement(paragraph, f"{{{W_NS}}}{wrapper_name}")
        wrapper.append(run)

    unsafe = _rewrite_document_xml(
        plain,
        tmp_path / f"{mode}.docx",
        etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    )
    semantic = build_front_matter_source(audit_brief(unsafe))
    output = tmp_path / f"{mode}-projected.docx"

    result = project_front_matter_fields(unsafe, output, semantic)

    assert any(
        item.application_status == "review-only" and expected_note in item.note
        for item in result.corrections
    )
    assert [paragraph.text for paragraph in inspect_brief_docx(unsafe).paragraphs] == [
        paragraph.text for paragraph in inspect_brief_docx(output).paragraphs
    ]


def test_projection_splits_a_styled_run_without_losing_formatting(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("[TOA]")
    paragraph = document.add_paragraph()
    paragraph.add_run("28 U.S.C. § 1332 followed by synthetic text.").bold = True
    source = tmp_path / "styled.docx"
    document.save(source)
    semantic = build_front_matter_source(audit_brief(source))
    output = tmp_path / "projected.docx"

    result = project_front_matter_fields(source, output, semantic)

    assert result.toa_marker_inserted_count == 1
    root = _document_xml(output)
    paragraph_xml = root.xpath(".//w:body/w:p", namespaces=NS)[1]
    suffixes = paragraph_xml.xpath(
        './w:r[w:rPr/w:b]/w:t[contains(., "followed by synthetic text")]',
        namespaces=NS,
    )
    assert suffixes


def test_signed_package_abstains_and_stays_byte_identical(tmp_path: Path) -> None:
    unsigned = _brief(tmp_path / "unsigned.docx")
    signed = tmp_path / "signed.docx"
    with zipfile.ZipFile(unsigned) as incoming, zipfile.ZipFile(signed, "w") as outgoing:
        for member in incoming.infolist():
            outgoing.writestr(member, incoming.read(member))
        outgoing.writestr("_xmlsignatures/sig1.xml", "<Signature/>")
    semantic = build_front_matter_source(audit_brief(signed))
    output = tmp_path / "output.docx"

    result = project_front_matter_fields(signed, output, semantic)

    assert output.read_bytes() == signed.read_bytes()
    assert result.applied_count == 0
    assert result.review_count == 8
    assert all(item.application_status == "review-only" for item in result.corrections)
    assert all("digitally signed" in item.note for item in result.corrections)


def test_multiple_exact_slots_abstain_instead_of_guessing_placement(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("[TOA]")
    document.add_paragraph("[TOA]")
    document.add_paragraph("See 28 U.S.C. § 1332.")
    source = tmp_path / "ambiguous.docx"
    document.save(source)
    semantic = build_front_matter_source(audit_brief(source))
    output = tmp_path / "output.docx"

    result = project_front_matter_fields(source, output, semantic)

    assert result.toa_marker_candidate_count == 0
    assert result.toa_field_candidate_count == 2
    assert result.applied_count == 0
    assert result.review_count == 2
    assert output.read_bytes() == source.read_bytes()
    assert all("placement ambiguous" in item.note for item in result.corrections)


def test_tracked_authority_abstains_without_losing_visible_text(tmp_path: Path) -> None:
    plain = Document()
    plain.add_paragraph("[TOA]")
    plain.add_paragraph("See 28 U.S.C. § 1332.")
    source = tmp_path / "plain.docx"
    plain.save(source)
    root = _document_xml(source)
    paragraph = root.xpath(".//w:body/w:p", namespaces=NS)[1]
    run = paragraph.find(f"{{{W_NS}}}r")
    assert run is not None
    paragraph.remove(run)
    insertion = etree.SubElement(paragraph, f"{{{W_NS}}}ins")
    insertion.append(run)
    tracked = _rewrite_document_xml(
        source,
        tmp_path / "tracked.docx",
        etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    )
    semantic = build_front_matter_source(audit_brief(tracked))
    output = tmp_path / "projected.docx"

    result = project_front_matter_fields(tracked, output, semantic)

    assert result.toa_marker_candidate_count == 1
    assert result.toa_marker_inserted_count == 0
    assert result.review_count == 1
    tracked_item = next(item for item in result.corrections if item.item_type == "toa-authority")
    assert "tracked change" in tracked_item.note
    assert [paragraph.text for paragraph in inspect_brief_docx(tracked).paragraphs] == [
        paragraph.text for paragraph in inspect_brief_docx(output).paragraphs
    ]
