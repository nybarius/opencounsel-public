from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

from opencounsel.templates.audit import audit_word_styles, write_style_audit
from opencounsel.templates.build import (
    ARGUMENT_OUTLINE_STYLE_NAME,
    BODY_STYLE,
    CASE_NAME_STYLE,
    FOOTNOTE_STYLE,
    HEADING_STYLES,
    SECTION_HEADING_STYLE,
    TOA_CATEGORY_STYLE,
    TOA_TITLE_STYLE,
    build_blank_word_template,
    configure_main_body_section,
)
from opencounsel.templates.profiles import get_bundled_filing_profile
from opencounsel.templates.registry import load_template_manifest, validate_word_template

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HEADING_STYLE_IDS = tuple(style_id for _, style_id, _ in HEADING_STYLES)


@pytest.mark.parametrize(
    ("profile_id", "body_pt", "footnote_pt"),
    [
        ("ny-ad-appellant-brief", 14, 12),
        ("ny-coa-principal-brief", 14, 12),
        ("ny-commercial-motion-memorandum", 12, 12),
        ("fed-ca2-principal-brief", 14, 14),
        ("fed-sdny-edny-motion-memorandum", 12, 10),
    ],
)
def test_blank_template_resolves_profile_typography_and_geometry(
    tmp_path: Path, profile_id: str, body_pt: int, footnote_pt: int
) -> None:
    output = tmp_path / f"{profile_id}.docx"
    manifest_path = tmp_path / f"{profile_id}.toml"
    result = build_blank_word_template(profile_id, output, manifest_path=manifest_path)
    document = Document(output)
    section = document.sections[0]

    assert result.profile_id == profile_id
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.top_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.bottom_margin == Inches(1)
    assert section.left_margin == Inches(1)
    assert document.styles[BODY_STYLE].font.name == "Times New Roman"
    assert document.styles[BODY_STYLE].font.size == Pt(body_pt)
    assert document.styles[BODY_STYLE].paragraph_format.line_spacing == 2.0
    assert document.styles[BODY_STYLE].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    footnote = document.styles[FOOTNOTE_STYLE]
    assert footnote.font.size == Pt(footnote_pt)
    assert footnote.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert all(not paragraph.text for paragraph in document.paragraphs)

    manifest = load_template_manifest(manifest_path)
    assert manifest.schema_version == 2
    assert manifest.filing_profile_id == profile_id
    assert manifest.style_pack_id == "opencounsel-clean-serif"
    assert validate_word_template(output, manifest).template_sha256 == result.template_sha256


def test_major_and_numbered_heading_hierarchy(tmp_path: Path) -> None:
    output = tmp_path / "appellate.docx"
    build_blank_word_template("ny-ad-appellant-brief", output)
    document = Document(output)

    major = document.styles[SECTION_HEADING_STYLE]
    assert major.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert major.font.bold is True
    assert major.font.underline is True
    assert _outline(major) == 0
    assert document.styles[ARGUMENT_OUTLINE_STYLE_NAME].type == WD_STYLE_TYPE.LIST

    for index, (name, _, _) in enumerate(HEADING_STYLES, start=1):
        style = document.styles[name]
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert style.paragraph_format.keep_with_next is True
        assert style.paragraph_format.keep_together is True
        assert style.paragraph_format.space_after == Pt(12)
        assert style.base_style is None
        assert style.paragraph_format.left_indent is None
        assert style.paragraph_format.first_line_indent is None
        assert not list(style.paragraph_format.tab_stops)
        assert _outline(style) == index

    long_heading = (
        "THE ORDER SHOULD BE REVERSED BECAUSE THE RECORD DOES NOT SUPPORT "
        "THE ONLY GROUND ON WHICH THE COURT RELIED"
    )
    paragraph = document.paragraphs[0]
    paragraph.style = document.styles[HEADING_STYLES[0][0]]
    paragraph.add_run(long_heading)
    document.save(output)

    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
        styles_xml = archive.read("word/styles.xml")
        numbering_xml = archive.read("word/numbering.xml")
    assert b"<w:br" not in document_xml
    assert b"<w:tab" not in document_xml
    assert document_xml.count(long_heading.encode()) == 1

    num_ids = {_style_num_id(styles_xml, style_id.encode()) for style_id in HEADING_STYLE_IDS}
    assert len(num_ids) == 1
    for level, style_id in enumerate(HEADING_STYLE_IDS):
        assert _style_list_level(styles_xml, style_id.encode()) == str(level).encode()
        assert not _style_has_list_indent(styles_xml, style_id.encode())

    abstract = _heading_abstract(numbering_xml)
    assert _value(abstract, "w:multiLevelType") == "hybridMultilevel"
    assert _value(abstract, "w:name") == ARGUMENT_OUTLINE_STYLE_NAME
    assert _value(abstract, "w:styleLink") == "LegalOutline"
    levels = _levels_by_index(abstract)
    assert [_value(levels[level], "w:lvlText") for level in range(4)] == ["%1.", "%2.", "%3.", "%4."]
    assert [_value(levels[level], "w:numFmt") for level in range(4)] == [
        "upperRoman", "upperLetter", "decimal", "lowerLetter"
    ]
    assert [_value(levels[level], "w:pStyle") for level in range(4)] == list(HEADING_STYLE_IDS)
    assert [_indent(levels[level]) for level in range(4)] == [
        (720, 720), (1440, 720), (2160, 720), (2880, 720)
    ]
    assert [_tab_position(levels[level]) for level in range(4)] == [720, 1440, 2160, 2880]
    assert all(_value(levels[level], "w:suff") == "tab" for level in range(4))


def test_motion_uses_same_i_a_1_a_grammar(tmp_path: Path) -> None:
    output = tmp_path / "motion.docx"
    build_blank_word_template("fed-sdny-edny-motion-memorandum", output)
    with zipfile.ZipFile(output) as archive:
        levels = _levels_by_index(_heading_abstract(archive.read("word/numbering.xml")))
    assert [_value(levels[level], "w:lvlText") for level in range(4)] == ["%1.", "%2.", "%3.", "%4."]


def test_front_matter_styles_have_five_levels_and_real_spacing(tmp_path: Path) -> None:
    output = tmp_path / "front-matter.docx"
    build_blank_word_template("ny-ad-appellant-brief", output)
    document = Document(output)

    expected = ((0.0, 0.0), (0.5, -0.5), (1.0, -0.5), (1.5, -0.5), (2.0, -0.5))
    for level, (left, first_line) in enumerate(expected, start=1):
        style = document.styles[f"TOC {level}"]
        formatting = style.paragraph_format
        assert formatting.line_spacing == 1.0
        assert formatting.left_indent == Inches(left)
        assert formatting.first_line_indent == Inches(first_line)
        assert formatting.space_after == Pt(12)
        expected_tabs = [9360] if level == 1 else [int(left * 1440), 9360]
        assert _style_tab_positions(style) == expected_tabs
        assert _contextual_spacing_value(style) == "0"

    assert document.styles[TOA_TITLE_STYLE].paragraph_format.page_break_before is True
    toa = document.styles["Table of Authorities"].paragraph_format
    assert toa.left_indent == Inches(0.5)
    assert toa.first_line_indent == Inches(-0.5)
    assert toa.space_after == Pt(12)
    for name in ("TOA Heading", TOA_CATEGORY_STYLE):
        category = document.styles[name]
        assert category.font.bold is False
        assert category.font.underline is False
        assert category.font.small_caps is True
        assert category.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
        assert category.paragraph_format.space_after == Pt(6)
    assert document.styles[CASE_NAME_STYLE].font.italic is True
    assert all(not style.name.startswith("OC ") for style in document.styles)


def test_front_and_main_sections_use_required_page_numbering(tmp_path: Path) -> None:
    output = tmp_path / "sections.docx"
    profile = get_bundled_filing_profile("ny-ad-appellant-brief")
    build_blank_word_template(profile, output)
    document = Document(output)
    main = document.add_section(WD_SECTION.NEW_PAGE)
    configure_main_body_section(document, main, profile)
    document.save(output)

    with zipfile.ZipFile(output) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
        footer_xml = [
            etree.fromstring(archive.read(name))
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        ]
    sections = document_xml.findall(".//" + qn("w:sectPr"))
    assert len(sections) == 2
    front_num = sections[0].find(qn("w:pgNumType"))
    main_num = sections[1].find(qn("w:pgNumType"))
    assert front_num is not None and main_num is not None
    assert front_num.get(qn("w:fmt")) == "lowerRoman"
    assert front_num.get(qn("w:start")) == "1"
    assert main_num.get(qn("w:fmt")) == "decimal"
    assert main_num.get(qn("w:start")) == "1"
    assert sections[1].find(qn("w:titlePg")) is not None
    assert any(root.xpath(".//w:tbl", namespaces={"w": W_NS}) for root in footer_xml)


def test_generated_template_has_clean_metadata_and_no_external_relationships(tmp_path: Path) -> None:
    output = tmp_path / "clean.docx"
    build_blank_word_template("ny-ad-appellant-brief", output)
    with zipfile.ZipFile(output) as archive:
        core = archive.read("docProps/core.xml")
        settings = archive.read("word/settings.xml")
        external = sum(
            archive.read(name).count(b'TargetMode="External"')
            for name in archive.namelist()
            if name.endswith(".rels")
        )
    assert b"dc:creator" not in core
    assert b"cp:lastModifiedBy" not in core
    assert b"dcterms:created" not in core
    assert b"dcterms:modified" not in core
    assert b"w:updateFields" in settings
    assert external == 0


def test_privacy_safe_audit_reports_structure_without_document_text(tmp_path: Path) -> None:
    path = tmp_path / "legacy.docx"
    document = Document()
    document.core_properties.author = "Private Author"
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.0
    heading = document.add_heading("A confidential proposition", level=1)
    heading.add_run().add_break(WD_BREAK.LINE)
    body = document.add_paragraph("Private body text")
    body.paragraph_format.left_indent = Inches(0.25)
    body.add_run(" private continuation").bold = True
    body.add_run().add_tab()
    document.add_paragraph()
    document.save(path)

    report = audit_word_styles(path, get_bundled_filing_profile("ny-ad-appellant-brief"))
    codes = {finding.code for finding in report.findings}
    assert {
        "body-font-too-small", "body-font-kind-not-allowed", "body-not-double-spaced",
        "blank-spacer-paragraph", "direct-paragraph-formatting", "direct-run-formatting",
        "manual-line-break-in-heading", "manual-tab", "metadata-present",
    } <= codes
    assert "full-justification" not in codes

    output = tmp_path / "style-audit.json"
    write_style_audit(report, output)
    serialized = output.read_text(encoding="utf-8")
    assert "A confidential proposition" not in serialized
    assert "Private body text" not in serialized
    assert "Private Author" not in serialized
    assert json.loads(serialized)["schema_version"] == 1


def test_generated_template_passes_high_severity_style_audit(tmp_path: Path) -> None:
    output = tmp_path / "clean.docx"
    build_blank_word_template("ny-ad-appellant-brief", output)
    report = audit_word_styles(output, get_bundled_filing_profile("ny-ad-appellant-brief"))
    assert not [finding for finding in report.findings if finding.severity == "error"]


def _outline(style: object) -> int:
    outline = style._element.find(qn("w:pPr") + "/" + qn("w:outlineLvl"))
    assert outline is not None
    return int(outline.get(qn("w:val")))


def _style_tab_positions(style: object) -> list[int]:
    return [
        int(tab.get(qn("w:pos"), "-1"))
        for tab in style._element.findall(f"{qn('w:pPr')}/{qn('w:tabs')}/{qn('w:tab')}")
    ]


def _contextual_spacing_value(style: object) -> str | None:
    element = style._element.find(f"{qn('w:pPr')}/{qn('w:contextualSpacing')}")
    return None if element is None else element.get(qn("w:val"))


def _style_xml(styles_xml: bytes, style_id: bytes) -> bytes:
    marker = b'w:styleId="' + style_id + b'"'
    start = styles_xml.rindex(b"<w:style", 0, styles_xml.index(marker) + 1)
    end = styles_xml.index(b"</w:style>", start) + len(b"</w:style>")
    return styles_xml[start:end]


def _style_num_id(styles_xml: bytes, style_id: bytes) -> bytes:
    style_xml = _style_xml(styles_xml, style_id)
    marker = b'<w:numId w:val="'
    start = style_xml.index(marker) + len(marker)
    return style_xml[start : style_xml.index(b'"', start)]


def _style_list_level(styles_xml: bytes, style_id: bytes) -> bytes:
    style_xml = _style_xml(styles_xml, style_id)
    marker = b'<w:ilvl w:val="'
    start = style_xml.index(marker) + len(marker)
    return style_xml[start : style_xml.index(b'"', start)]


def _style_has_list_indent(styles_xml: bytes, style_id: bytes) -> bool:
    style_xml = _style_xml(styles_xml, style_id)
    return b"w:left=" in style_xml or b"w:hanging=" in style_xml


def _heading_abstract(numbering_xml: bytes) -> etree._Element:
    root = etree.fromstring(numbering_xml)
    for abstract in root.findall(qn("w:abstractNum")):
        style_ids = {
            element.get(qn("w:val"))
            for element in abstract.findall(".//" + qn("w:pStyle"))
        }
        if HEADING_STYLE_IDS[0] in style_ids:
            return abstract
    raise AssertionError("heading numbering definition not found")


def _levels_by_index(abstract: etree._Element) -> dict[int, etree._Element]:
    return {int(level.get(qn("w:ilvl"), "-1")): level for level in abstract.findall(qn("w:lvl"))}


def _value(element: etree._Element, child_name: str) -> str | None:
    child = element.find(qn(child_name))
    return None if child is None else child.get(qn("w:val"))


def _indent(level: etree._Element) -> tuple[int, int]:
    p_pr = level.find(qn("w:pPr"))
    assert p_pr is not None
    indent = p_pr.find(qn("w:ind"))
    assert indent is not None
    return int(indent.get(qn("w:left"), "-1")), int(indent.get(qn("w:hanging"), "-1"))


def _tab_position(level: etree._Element) -> int:
    p_pr = level.find(qn("w:pPr"))
    assert p_pr is not None
    tab = p_pr.find(f"{qn('w:tabs')}/{qn('w:tab')}")
    assert tab is not None
    return int(tab.get(qn("w:pos"), "-1"))
