from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from opencounsel.briefs.cite_check import plan_citation_review


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _citation_brief(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Id. supplies no approved antecedent.")
    document.add_paragraph(
        "See 42 U.S.C. § 1983; Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234; "
        "and R. 12\u201313."
    )
    linked = document.add_paragraph("See ")
    _add_hyperlink(
        linked,
        "550 U.S. 544",
        "https://www.courtlistener.com/opinion/1/example/",
    )
    document.save(path)
    return path


def test_citation_review_is_stable_offline_and_location_exact(tmp_path: Path) -> None:
    source = _citation_brief(tmp_path / "brief.docx")

    first = plan_citation_review(source)
    second = plan_citation_review(source)

    assert first == second
    assert first.record_citation_count == 1
    assert first.authority_citation_count == 4
    assert len(first.findings) == 5
    assert len({item.correction_id for item in first.findings}) == 5
    assert all(item.application_status == "review-only" for item in first.findings)
    assert all(item.stage == "cite-check" for item in first.findings)
    assert all(item.original_text == item.replacement_text for item in first.findings)
    assert all(item.basis_sha256 == first.input_sha256 for item in first.findings)
    assert {item.item_type for item in first.findings} == {
        "record-citation",
        "authority-citation",
    }
    notes = "\n".join(item.note for item in first.findings)
    assert "no ROA package" in notes
    assert "no uniquely approved antecedent" in notes
    assert "official source candidate" in notes
    assert "licensed or manually supplied source" in notes
    assert "Existing navigation link" in notes
    assert "does not establish source identity or substantive verification" in notes


def test_resolved_short_form_names_its_offline_source_candidate(tmp_path: Path) -> None:
    source = tmp_path / "short.docx"
    document = Document()
    document.add_paragraph(
        "See People v Doe, 2024 NY Slip Op 01234. Id. supplies the governing rule."
    )
    document.save(source)

    plan = plan_citation_review(source)

    assert plan.authority_citation_count == 2
    assert len(plan.findings) == 2
    assert all("official source candidate" in item.note for item in plan.findings)
    assert plan.findings[1].original_text == "Id."
