from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from opencounsel.briefs.docx import inspect_brief_docx
from opencounsel.briefs.proof import plan_proof_corrections
from opencounsel.briefs.proof_docx import ProofProjectionError, apply_proof_corrections


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _proof_brief(path: Path) -> Path:
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("Argument ")
    first.bold = True
    second = paragraph.add_run(", controls ; but ellipsis . . . remains.")
    second.italic = True
    linked = document.add_paragraph("See ")
    _add_hyperlink(linked, "linked , authority", "https://example.test/authority")
    document.save(path)
    return path


def test_plans_only_bounded_spacing_corrections_with_stable_ids(tmp_path: Path) -> None:
    source = _proof_brief(tmp_path / "source.docx")

    first = plan_proof_corrections(source)
    second = plan_proof_corrections(source)

    assert first == second
    assert [item.original_text for item in first.corrections] == [" ,", " ;", " ,"]
    assert [item.replacement_text for item in first.corrections] == [",", ";", ","]
    assert all(item.stage == "proof" for item in first.corrections)
    assert all(item.correction_id.startswith("corr-") for item in first.corrections)
    assert ". . ." not in "".join(item.original_text for item in first.corrections)


def test_applies_safe_edits_across_runs_and_routes_hyperlink_to_review(
    tmp_path: Path,
) -> None:
    source = _proof_brief(tmp_path / "source.docx")
    before = source.read_bytes()
    plan = plan_proof_corrections(source)
    output = tmp_path / "corrected.docx"

    result = apply_proof_corrections(source, output, plan)

    assert source.read_bytes() == before
    assert result.input_sha256 == plan.input_sha256
    assert result.output_sha256 != result.input_sha256
    assert result.applied_count == 2
    assert result.review_count == 1
    assert [item.application_status for item in result.corrections] == [
        "applied",
        "applied",
        "review-only",
    ]
    assert "hyperlink" in result.corrections[-1].note
    inspection = inspect_brief_docx(output)
    assert inspection.paragraphs[0].text == "Argument, controls; but ellipsis . . . remains."
    assert inspection.paragraphs[1].text == "See linked , authority"
    assert inspection.paragraphs[1].hyperlinks[0].target == "https://example.test/authority"
    rendered = Document(output)
    assert rendered.paragraphs[0].runs[0].bold is True
    assert rendered.paragraphs[0].runs[1].italic is True


def test_proof_projection_rejects_stale_plan_and_existing_output(tmp_path: Path) -> None:
    source = _proof_brief(tmp_path / "source.docx")
    plan = plan_proof_corrections(source)
    altered = tmp_path / "altered.docx"
    document = Document(source)
    document.add_paragraph("Changed.")
    document.save(altered)

    with pytest.raises(ProofProjectionError, match="does not match"):
        apply_proof_corrections(altered, tmp_path / "out.docx", plan)

    output = tmp_path / "existing.docx"
    output.write_bytes(b"occupied")
    with pytest.raises(ProofProjectionError, match="already exists"):
        apply_proof_corrections(source, output, plan)


def test_signed_package_and_word_field_are_review_only(tmp_path: Path) -> None:
    unsigned = _proof_brief(tmp_path / "unsigned.docx")
    signed = tmp_path / "signed.docx"
    with zipfile.ZipFile(unsigned) as incoming, zipfile.ZipFile(signed, "w") as outgoing:
        for member in incoming.infolist():
            outgoing.writestr(member, incoming.read(member))
        outgoing.writestr("_xmlsignatures/sig1.xml", "<signature/>")

    signed_plan = plan_proof_corrections(signed)
    signed_output = tmp_path / "signed-corrected.docx"
    signed_result = apply_proof_corrections(signed, signed_output, signed_plan)

    assert signed_result.applied_count == 0
    assert signed_result.review_count == 3
    assert signed_output.read_bytes() == signed.read_bytes()
    assert all("digitally signed" in item.note for item in signed_result.corrections)

    field_source = tmp_path / "field.docx"
    document = Document()
    paragraph = document.add_paragraph("Field paragraph , unchanged.")
    run = OxmlElement("w:r")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    run.append(field)
    paragraph._p.append(run)
    document.save(field_source)
    field_output = tmp_path / "field-corrected.docx"

    field_result = apply_proof_corrections(
        field_source,
        field_output,
        plan_proof_corrections(field_source),
    )

    assert field_result.applied_count == 0
    assert field_result.review_count == 1
    assert "Word field" in field_result.corrections[0].note
    assert field_output.read_bytes() == field_source.read_bytes()
