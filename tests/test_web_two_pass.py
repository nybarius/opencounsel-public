from __future__ import annotations

import csv
import io
import json
import time
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from starlette.testclient import TestClient

from opencounsel.web.app import create_app
from opencounsel.web.jobs import JobContext, JobManager, ProgressUpdate

WORD_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def _docx_bytes() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph(
        'Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234, held that "the agreement is unambiguous."'
    )
    document.save(stream)
    return stream.getvalue()


def _text_pdf(text: str) -> bytes:
    stream = io.BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    page[NameObject("/Resources")] = resources
    content = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(stream)
    return stream.getvalue()


def _processor(context: JobContext, progress: ProgressUpdate) -> dict[str, object]:
    progress("reporting", "Building source and review reports")
    export = context.job_dir / "export"
    bundle = export / "reports" / "authority-sources"
    bundle.mkdir(parents=True, mode=0o700)
    authority_id = "auth-1234567890abcdef"
    manifest = {
        "schema_version": 1,
        "brief_sha256": "1" * 64,
        "unresolved_occurrence_count": 0,
        "authorities": [
            {
                "authority_id": authority_id,
                "canonical_citation": "Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234",
                "case_name": "Smith v. Jones",
                "category": "cases",
                "court": "N.Y. Sup. Ct.",
                "year": "2026",
                "docket": None,
                "source_status": "source-needed",
                "source_url": None,
                "source_resolver": None,
                "expected_filename": f"{authority_id}.pdf",
                "assertions": [
                    {
                        "part": "body",
                        "paragraph_index": 1,
                        "citation_text": "2026 N.Y. Misc. LEXIS 1234",
                        "citation_type": "case",
                        "context": 'Smith held that "the agreement is unambiguous."',
                        "resolution_basis": "full-citation",
                    }
                ],
                "verification_targets": [
                    "citation-exists",
                    "quotation-exact",
                    "quotation-signals-complete",
                    "characterization-supported",
                ],
            }
        ],
    }
    (bundle / "authority-manifest.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    with (bundle / "authority-intake.csv").open(
        "w", encoding="utf-8", newline=""
    ) as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=(
                "authority_id",
                "canonical_citation",
                "case_name",
                "expected_filename",
                "source_file",
                "source_url",
            ),
        )
        writer.writeheader()
        writer.writerow(
            {
                "authority_id": authority_id,
                "canonical_citation": manifest["authorities"][0]["canonical_citation"],
                "case_name": "Smith v. Jones",
                "expected_filename": f"{authority_id}.pdf",
                "source_file": "",
                "source_url": "",
            }
        )
    for name, data in {
        "published.docx": b"immutable prepared docx",
        "published.pdf": b"%PDF immutable prepared pdf",
        "opencounsel-filing-package.zip": b"PK first pass",
        "authority-source-report.json": json.dumps(manifest).encode(),
    }.items():
        (export / name).write_bytes(data)
    return {
        "summary": {
            "record_citation_count": 0,
            "resolved_record_citation_count": 0,
            "unresolved_record_citation_count": 0,
            "toa_authority_count": 1,
            "toa_occurrence_count": 1,
            "toc_entry_count": 1,
            "hyperlink_inserted_count": 0,
            "review_item_count": 1,
            "page_count": 1,
            "body_font_pt": 12,
            "footnote_font_pt": 10,
            "source_copy_required_count": 1,
        },
        "sources": [
            {
                "authority_id": authority_id,
                "citation": manifest["authorities"][0]["canonical_citation"],
                "category": "cases",
                "status": "source-copy-required",
                "url": None,
                "occurrence_count": 1,
            }
        ],
        "audit": {
            "input_sha256": "1" * 64,
            "output_docx_sha256": "2" * 64,
            "output_pdf_sha256": "3" * 64,
        },
        "artifacts": {
            "package": {
                "path": "export/opencounsel-filing-package.zip",
                "filename": "opencounsel-filing-package.zip",
                "media_type": "application/zip",
            },
            "document": {
                "path": "export/published.docx",
                "filename": "published.docx",
                "media_type": WORD_MEDIA_TYPE,
            },
            "pdf": {
                "path": "export/published.pdf",
                "filename": "published.pdf",
                "media_type": "application/pdf",
            },
            "sources": {
                "path": "export/authority-source-report.json",
                "filename": "authority-source-report.json",
                "media_type": "application/json",
            },
        },
    }


def _completed(client: TestClient, job_id: str) -> dict[str, object]:
    for _ in range(100):
        response = client.get(f"/api/jobs/{job_id}")
        payload = response.json()
        if payload["status"] in {"completed", "failed"}:
            return payload
        time.sleep(0.01)
    raise AssertionError("job did not finish")


def _prepared_job(client: TestClient, tmp_path: Path) -> tuple[str, str]:
    record = tmp_path / "record.zip"
    record.write_bytes(b"PK synthetic record")
    response = client.post(
        "/api/jobs",
        data={"profile": "fed-sdny-edny-motion-memorandum"},
        files={
            "brief": ("memorandum.docx", _docx_bytes()),
            "roa": ("record.zip", record.read_bytes(), "application/zip"),
        },
    )
    assert response.status_code == 202
    job_id = response.json()["job_id"]
    result = _completed(client, job_id)
    assert result["status"] == "completed"
    return job_id, result["sources"][0]["authority_id"]


def test_declared_source_upload_verifies_and_preserves_first_pass_artifacts(
    tmp_path: Path,
) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_processor)
    with TestClient(create_app(manager=manager)) as client:
        job_id, authority_id = _prepared_job(client, tmp_path)
        before_docx = client.get(f"/api/jobs/{job_id}/artifacts/document").content
        before_pdf = client.get(f"/api/jobs/{job_id}/artifacts/pdf").content
        response = client.post(
            f"/api/jobs/{job_id}/sources",
            data={"authority_id": authority_id},
            files={
                "source": (
                    "opinion.pdf",
                    _text_pdf(
                        "Smith v. Jones. 2026 N.Y. Misc. LEXIS 1234. "
                        "The agreement is unambiguous."
                    ),
                    "application/pdf",
                )
            },
        )
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["source_verification"]["citation_confirmed_count"] == 1
        assert result["source_verification"]["quotation_failure_count"] == 0
        assert result["source_verification"]["characterization_review_count"] == 1
        assert result["sources"][0]["status"] == "verified-source"
        assert client.get(f"/api/jobs/{job_id}/artifacts/document").content == before_docx
        assert client.get(f"/api/jobs/{job_id}/artifacts/pdf").content == before_pdf
        authority_package = client.get(
            f"/api/jobs/{job_id}/artifacts/authority-package"
        )
        assert authority_package.status_code == 200
        verification = client.get(f"/api/jobs/{job_id}/artifacts/verification")
        assert verification.status_code == 200
        final_package = client.get(f"/api/jobs/{job_id}/artifacts/final-package")
        assert final_package.status_code == 200
        assert final_package.content.startswith(b"PK")


def test_source_upload_rejects_invalid_or_undeclared_inputs(tmp_path: Path) -> None:
    manager = JobManager(tmp_path / "jobs", processor=_processor)
    with TestClient(create_app(manager=manager)) as client:
        job_id, authority_id = _prepared_job(client, tmp_path)
        mismatched = client.post(
            f"/api/jobs/{job_id}/sources",
            data={"authority_id": authority_id},
            files={
                "source": (
                    "wrong.pdf",
                    _text_pdf("Unrelated Person v. Entity, 1 F.4th 2"),
                )
            },
        )
        assert mismatched.status_code == 400
        assert "does not match" in mismatched.json()["error"]

        undeclared = client.post(
            f"/api/jobs/{job_id}/sources",
            data={"authority_id": "auth-undeclared"},
            files={"source": ("source.pdf", _text_pdf("Smith v. Jones"))},
        )
        assert undeclared.status_code == 400
        assert "declared authority" in undeclared.json()["error"]

        malformed = client.post(
            f"/api/jobs/{job_id}/sources",
            data={"authority_id": authority_id},
            files={"source": ("source.pdf", b"not a pdf", "application/pdf")},
        )
        assert malformed.status_code == 400
        assert "readable PDF" in malformed.json()["error"]

        duplicate = client.post(
            f"/api/jobs/{job_id}/sources",
            data=[("authority_id", authority_id), ("authority_id", authority_id)],
            files=[
                ("source", ("one.pdf", _text_pdf("Smith v. Jones"), "application/pdf")),
                ("source", ("two.pdf", _text_pdf("Smith v. Jones"), "application/pdf")),
            ],
        )
        assert duplicate.status_code == 400
        assert "duplicate" in duplicate.json()["error"].lower()
