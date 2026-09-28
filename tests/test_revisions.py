from __future__ import annotations

import json
import os
import stat
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

import opencounsel.revisions.store as store_module
from opencounsel.briefs.docx import W_NS, inspect_brief_docx
from opencounsel.revisions import RevisionError, RevisionStore


def _brief(path: Path) -> Path:
    document = Document()
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph("Synthetic proposition in 550 U.S. 544 (2007).")
    document.save(path)
    return path


def _proof_brief(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Synthetic proposition , followed by conclusion ;")
    document.save(path)
    return path


def _front_matter_brief(path: Path) -> Path:
    document = Document()
    document.add_paragraph("[TOC]")
    document.add_paragraph("[TOA]")
    document.add_heading("ARGUMENT", level=1)
    document.add_paragraph(
        "Bell Atlantic Corp. v. Twombly, 550 U.S. 544, 570 (2007), controls. "
        "The claim also arises under 42 U.S.C. § 1983."
    )
    document.save(path)
    return path


def test_revision_creation_is_private_immutable_and_idempotent(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "sensitive-client-name.docx")
    store = RevisionStore(tmp_path / "work", inbox)

    first = store.create(source)
    second = store.create(source)
    manifest = store.load(first.revision_id)
    inspection = store.inspect(first.revision_id)

    assert first.created is True
    assert second.created is False
    assert first.revision_id == second.revision_id
    assert manifest.input_sha256 == first.input_sha256
    assert inspection.paragraph_count == 2
    assert inspection.heading_count == 1
    assert inspection.part_counts == {"document": 2}
    assert "sensitive-client-name" not in json.dumps(manifest.model_dump(mode="json"))
    assert "Synthetic proposition" not in inspection.model_dump_json()
    if os.name == "posix":
        manifest_path = store.manifests / f"{first.revision_id}.json"
        object_path = store.objects.path_for(manifest.input_object_key)
        assert stat.S_IMODE(store.root.stat().st_mode) == 0o700
        assert stat.S_IMODE(manifest_path.stat().st_mode) == 0o600
        assert stat.S_IMODE(object_path.stat().st_mode) == 0o600


def test_revision_store_rejects_unconfined_or_unsafe_sources(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    outside = _brief(tmp_path / "outside.docx")
    store = RevisionStore(tmp_path / "work", inbox)

    with pytest.raises(RevisionError, match="outside"):
        store.create(outside)
    non_docx = inbox / "brief.txt"
    non_docx.write_text("synthetic", encoding="utf-8")
    with pytest.raises(RevisionError, match="regular DOCX"):
        store.create(non_docx)
    unreadable = inbox / "unreadable.docx"
    unreadable.write_text("not OOXML", encoding="utf-8")
    with pytest.raises(RevisionError, match="readable OOXML"):
        store.create(unreadable)
    with pytest.raises(RevisionError, match="does not exist"):
        store.create(inbox / "missing.docx")
    link = inbox / "link.docx"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("symbolic links are unavailable")
    with pytest.raises(RevisionError, match="symbolic link"):
        store.create(link)


def test_revision_store_fails_closed_on_bad_identity_or_manifest(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "brief.docx")
    store = RevisionStore(tmp_path / "work", inbox)
    result = store.create(source)

    with pytest.raises(RevisionError, match="revision id"):
        store.load("../../escape")
    manifest_path = store.manifests / f"{result.revision_id}.json"
    manifest_path.write_text('{"schema_version": 1}', encoding="utf-8")
    with pytest.raises(RevisionError, match="manifest is invalid"):
        store.load(result.revision_id)
    manifest_path.write_text("[]", encoding="utf-8")
    with pytest.raises(RevisionError, match="JSON object"):
        store.load(result.revision_id)


def test_revision_store_detects_manifest_identity_and_object_tampering(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "brief.docx")

    identity_store = RevisionStore(tmp_path / "identity-work", inbox)
    identity = identity_store.create(source)
    manifest_path = identity_store.manifests / f"{identity.revision_id}.json"
    payload = identity_store.load(identity.revision_id).model_dump(mode="json")
    payload["revision_id"] = f"rev-{'1' * 64}"
    manifest_path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(RevisionError, match="identity does not match"):
        identity_store.load(identity.revision_id)

    missing_store = RevisionStore(tmp_path / "missing-work", inbox)
    missing = missing_store.create(source)
    missing_manifest = missing_store.load(missing.revision_id)
    missing_store.objects.path_for(missing_manifest.input_object_key).unlink()
    with pytest.raises(RevisionError, match="object does not exist"):
        missing_store.inspect(missing.revision_id)

    tampered_store = RevisionStore(tmp_path / "tampered-work", inbox)
    tampered = tampered_store.create(source)
    tampered_manifest = tampered_store.load(tampered.revision_id)
    tampered_store.objects.path_for(tampered_manifest.input_object_key).write_bytes(b"changed")
    with pytest.raises(RevisionError, match="readable OOXML"):
        tampered_store.inspect(tampered.revision_id)


def test_revision_store_rejects_missing_roots_and_symlinked_storage(tmp_path: Path) -> None:
    with pytest.raises(RevisionError, match="input root does not exist"):
        RevisionStore(tmp_path / "work", tmp_path / "missing")

    real = tmp_path / "real"
    real.mkdir()
    link = tmp_path / "linked-work"
    try:
        link.symlink_to(real, target_is_directory=True)
    except OSError:
        pytest.skip("symbolic links are unavailable")
    with pytest.raises(RevisionError, match="symbolic link"):
        RevisionStore(link, real)

    input_link = tmp_path / "linked-input"
    input_link.symlink_to(real, target_is_directory=True)
    with pytest.raises(RevisionError, match="input root must not be a symbolic link"):
        RevisionStore(tmp_path / "other-work", input_link)

    storage_file = tmp_path / "storage-file"
    storage_file.write_text("not a directory", encoding="utf-8")
    with pytest.raises(RevisionError, match="directory"):
        RevisionStore(storage_file, real)


def test_process_brief_is_hash_bound_private_and_idempotent(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "sensitive-client-name.docx")
    original = source.read_bytes()
    store = RevisionStore(tmp_path / "work", inbox)
    revision = store.create(source)

    first = store.process(revision.revision_id)
    second = store.process(revision.revision_id)
    manifest = store.load_process(first.process_id)
    ledger = store.load_correction_ledger(first.process_id)
    corrected_path, ledger_path = store.delivery_paths(first.process_id)
    front_matter_path = store.front_matter_path(first.process_id)

    assert first.created is True
    assert second.created is False
    assert first.process_id == second.process_id
    assert first.input_sha256 == first.corrected_sha256 == revision.input_sha256
    assert first.proofed_sha256 == revision.input_sha256
    assert first.correction_count == first.review_item_count == 1
    assert first.applied_correction_count == 0
    assert first.record_citation_count == 0
    assert first.authority_citation_count == 1
    assert first.citation_review_item_count == 1
    assert first.hyperlink_candidate_count == 0
    assert first.hyperlink_inserted_count == 0
    assert first.hyperlink_review_item_count == 0
    assert first.toc_entry_count == 1
    assert first.toa_authority_count == 1
    assert first.toa_occurrence_count == 1
    assert first.front_matter_field_candidate_count == 0
    assert first.hyperlinked_sha256 == first.corrected_sha256
    assert first.front_matter_source_sha256 is not None
    assert first.lawyer_review_status == "pending"
    assert manifest.source_revision_id == revision.revision_id
    assert manifest.corrected_sha256 == first.corrected_sha256
    assert manifest.correction_ledger_sha256 == first.correction_ledger_sha256
    assert ledger.process_id == first.process_id
    assert len(ledger.corrections) == 1
    assert ledger.corrections[0].stage == "cite-check"
    assert ledger.corrections[0].item_type == "authority-citation"
    assert ledger.corrections[0].application_status == "review-only"
    assert ledger.corrections[0].basis_sha256 == first.proofed_sha256
    assert corrected_path.read_bytes() == original
    front_matter = json.loads(front_matter_path.read_text(encoding="utf-8"))
    assert front_matter["brief_sha256"] == first.hyperlinked_sha256
    assert len(front_matter["toc"]) == 1
    assert len(front_matter["toa"]) == 1
    assert json.loads(ledger_path.read_text(encoding="utf-8")) == ledger.model_dump(mode="json")
    assert source.read_bytes() == original
    serialized = repr(
        (
            first.model_dump(mode="json"),
            manifest.model_dump(mode="json"),
            ledger.model_dump(mode="json"),
        )
    )
    assert "sensitive-client-name" not in serialized
    assert "Synthetic proposition" not in serialized
    if os.name == "posix":
        assert stat.S_IMODE(corrected_path.parent.stat().st_mode) == 0o700
        assert stat.S_IMODE(corrected_path.stat().st_mode) == 0o600
        assert stat.S_IMODE(ledger_path.stat().st_mode) == 0o600
        assert stat.S_IMODE(front_matter_path.stat().st_mode) == 0o600


def test_process_brief_compiles_front_matter_and_projects_exact_word_slots(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _front_matter_brief(inbox / "brief.docx")
    store = RevisionStore(tmp_path / "work", inbox)

    result = store.process(store.create(source).revision_id)
    manifest = store.load_process(result.process_id)
    ledger = store.load_correction_ledger(result.process_id)
    corrected, _ = store.delivery_paths(result.process_id)
    front_matter_path = store.front_matter_path(result.process_id)
    front_matter = json.loads(front_matter_path.read_text(encoding="utf-8"))

    assert result.toc_entry_count == 1
    assert result.toa_authority_count == 2
    assert result.toa_occurrence_count == 2
    assert result.front_matter_field_candidate_count == 5
    assert result.front_matter_field_applied_count == 5
    assert result.front_matter_field_review_item_count == 0
    assert result.toc_field_inserted_count == 1
    assert result.toa_field_inserted_count == 1
    assert result.hyperlinked_sha256 is not None
    assert result.hyperlinked_sha256 != result.corrected_sha256
    assert result.front_matter_source_sha256 == manifest.front_matter_source_sha256
    assert front_matter["brief_sha256"] == result.hyperlinked_sha256
    assert len(front_matter["toc"]) == 1
    assert len(front_matter["toa"]) == 2
    assert "\"page\"" not in json.dumps(front_matter)
    assert all(
        item.basis_sha256 == result.hyperlinked_sha256
        for item in ledger.corrections
        if item.stage in {"toc", "toa"}
    )

    with zipfile.ZipFile(corrected) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    instructions = "".join(
        root.xpath(".//w:instrText/text()", namespaces={"w": W_NS})
    )
    assert instructions.count(" TC ") == 1
    assert instructions.count(" TA ") == 2
    assert instructions.count(" TOC ") == 1
    assert instructions.count(" TOA ") == 1


def test_process_brief_rejects_oversized_front_matter_before_storage(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _front_matter_brief(inbox / "brief.docx")
    store = RevisionStore(tmp_path / "work", inbox)
    revision = store.create(source)
    monkeypatch.setattr(store_module, "MAX_FRONT_MATTER_BYTES", 1)

    with pytest.raises(RevisionError, match="front-matter source exceeds size limit"):
        store.process(revision.revision_id)


def test_process_brief_applies_proof_stage_and_records_pending_review(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _proof_brief(inbox / "brief.docx")
    original = source.read_bytes()
    store = RevisionStore(tmp_path / "work", inbox)
    revision = store.create(source)

    first = store.process(revision.revision_id)
    second = store.process(revision.revision_id)
    ledger = store.load_correction_ledger(first.process_id)
    corrected, _ = store.delivery_paths(first.process_id)

    assert first.created is True
    assert second.created is False
    assert first.correction_count == first.applied_correction_count == 2
    assert first.review_item_count == 0
    assert first.proofed_sha256 == first.corrected_sha256
    assert first.record_citation_count == 0
    assert first.authority_citation_count == 0
    assert first.citation_review_item_count == 0
    assert first.hyperlink_candidate_count == 0
    assert first.hyperlink_inserted_count == 0
    assert first.hyperlink_review_item_count == 0
    assert first.lawyer_review_status == "pending"
    assert first.corrected_sha256 != first.input_sha256
    assert [item.original_text for item in ledger.corrections] == [" ,", " ;"]
    assert all(item.application_status == "applied" for item in ledger.corrections)
    assert all(item.review_status == "pending" for item in ledger.corrections)
    assert inspect_brief_docx(corrected).paragraphs[0].text == (
        "Synthetic proposition, followed by conclusion;"
    )
    assert source.read_bytes() == original


def test_process_brief_combines_proof_and_offline_citation_findings(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = inbox / "combined.docx"
    document = Document()
    document.add_paragraph(
        "The claim arises under 42 U.S.C. § 1983 , as the record confirms (R. 7)."
    )
    document.save(source)
    store = RevisionStore(tmp_path / "work", inbox)

    result = store.process(store.create(source).revision_id)
    ledger = store.load_correction_ledger(result.process_id)
    corrected, _ = store.delivery_paths(result.process_id)

    assert result.correction_count == 4
    assert result.applied_correction_count == 2
    assert result.review_item_count == 2
    assert result.record_citation_count == 1
    assert result.authority_citation_count == 1
    assert result.citation_review_item_count == 2
    assert result.hyperlink_candidate_count == 1
    assert result.hyperlink_inserted_count == 1
    assert result.hyperlink_review_item_count == 0
    assert [item.stage for item in ledger.corrections] == [
        "proof",
        "cite-check",
        "cite-check",
        "hyperlink",
    ]
    assert [item.item_type for item in ledger.corrections] == [
        "correction",
        "record-citation",
        "authority-citation",
        "hyperlink",
    ]
    assert inspect_brief_docx(corrected).paragraphs[0].text == (
        "The claim arises under 42 U.S.C. § 1983, as the record confirms (R. 7)."
    )
    assert inspect_brief_docx(corrected).paragraphs[0].hyperlinks[0].target.startswith(
        "https://uscode.house.gov/view.xhtml"
    )
    assert ledger.corrections[0].basis_sha256 == result.input_sha256
    assert all(
        item.basis_sha256 == result.proofed_sha256
        for item in ledger.corrections[1:]
    )


def test_process_brief_projects_official_links_after_proofing(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = inbox / "official.docx"
    document = Document()
    paragraph = document.add_paragraph("The claim arises under 42 U.S.")
    paragraph.add_run("C. § 1983 , and CPLR 3211(a)(7).")
    document.save(source)
    store = RevisionStore(tmp_path / "work", inbox)

    result = store.process(store.create(source).revision_id)
    ledger = store.load_correction_ledger(result.process_id)
    corrected, _ = store.delivery_paths(result.process_id)
    inspection = inspect_brief_docx(corrected)

    assert result.correction_count == 5
    assert result.applied_correction_count == 3
    assert result.review_item_count == 2
    assert result.authority_citation_count == 2
    assert result.citation_review_item_count == 2
    assert result.hyperlink_candidate_count == 2
    assert result.hyperlink_inserted_count == 2
    assert result.hyperlink_review_item_count == 0
    assert result.input_sha256 != result.proofed_sha256
    assert result.proofed_sha256 != result.corrected_sha256
    assert inspection.paragraphs[0].text == (
        "The claim arises under 42 U.S.C. § 1983, and CPLR 3211(a)(7)."
    )
    assert len(inspection.paragraphs[0].hyperlinks) == 2
    assert [item.stage for item in ledger.corrections] == [
        "proof",
        "cite-check",
        "cite-check",
        "hyperlink",
        "hyperlink",
    ]
    assert all(
        item.application_status == "applied"
        for item in ledger.corrections
        if item.stage == "hyperlink"
    )


def test_process_brief_routes_signed_link_candidates_to_review(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    unsigned = inbox / "unsigned.docx"
    document = Document()
    document.add_paragraph("The claim arises under 42 U.S.C. § 1983.")
    document.save(unsigned)
    signed = inbox / "signed.docx"
    with zipfile.ZipFile(unsigned) as incoming, zipfile.ZipFile(signed, "w") as outgoing:
        for member in incoming.infolist():
            outgoing.writestr(member, incoming.read(member))
        outgoing.writestr("_xmlsignatures/sig1.xml", "<signature/>")
    store = RevisionStore(tmp_path / "work", inbox)

    result = store.process(store.create(signed).revision_id)
    ledger = store.load_correction_ledger(result.process_id)
    corrected, _ = store.delivery_paths(result.process_id)

    assert result.input_sha256 == result.proofed_sha256 == result.corrected_sha256
    assert result.correction_count == result.review_item_count == 2
    assert result.applied_correction_count == 0
    assert result.hyperlink_candidate_count == 1
    assert result.hyperlink_inserted_count == 0
    assert result.hyperlink_review_item_count == 1
    hyperlink = ledger.corrections[-1]
    assert hyperlink.stage == "hyperlink"
    assert hyperlink.item_type == "hyperlink"
    assert hyperlink.application_status == "review-only"
    assert "digitally signed" in hyperlink.note
    assert corrected.read_bytes() == signed.read_bytes()


def test_process_brief_leaves_subscription_only_authority_unlinked(
    tmp_path: Path,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = inbox / "vendor.docx"
    document = Document()
    document.add_paragraph("See Smith v. Jones, 2026 N.Y. Misc. LEXIS 1234.")
    document.save(source)
    store = RevisionStore(tmp_path / "work", inbox)

    result = store.process(store.create(source).revision_id)
    ledger = store.load_correction_ledger(result.process_id)
    corrected, _ = store.delivery_paths(result.process_id)

    assert result.authority_citation_count == 1
    assert result.citation_review_item_count == 1
    assert result.hyperlink_candidate_count == 0
    assert result.hyperlink_inserted_count == 0
    assert not inspect_brief_docx(corrected).paragraphs[0].hyperlinks
    assert "licensed or manually supplied source" in ledger.corrections[0].note


def test_process_brief_routes_tracked_link_candidate_to_review(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    plain = inbox / "plain.docx"
    document = Document()
    document.add_paragraph("The claim arises under 42 U.S.C. § 1983.")
    document.save(plain)
    tracked = inbox / "tracked.docx"
    namespace = {"w": W_NS}
    with zipfile.ZipFile(plain) as incoming, zipfile.ZipFile(tracked, "w") as outgoing:
        for member in incoming.infolist():
            data = incoming.read(member)
            if member.filename == "word/document.xml":
                root = etree.fromstring(data)
                run = root.xpath(".//w:body/w:p/w:r", namespaces=namespace)[0]
                parent = run.getparent()
                insertion = etree.Element(f"{{{W_NS}}}ins")
                parent.replace(run, insertion)
                insertion.append(run)
                data = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            outgoing.writestr(member, data)
    store = RevisionStore(tmp_path / "work", inbox)

    result = store.process(store.create(tracked).revision_id)
    ledger = store.load_correction_ledger(result.process_id)
    corrected, _ = store.delivery_paths(result.process_id)

    assert result.hyperlink_candidate_count == 1
    assert result.hyperlink_inserted_count == 0
    assert result.hyperlink_review_item_count == 1
    assert not inspect_brief_docx(corrected).paragraphs[0].hyperlinks
    assert "tracked change" in ledger.corrections[-1].note


def test_process_brief_detects_tampered_delivery_and_manifests(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "brief.docx")
    store = RevisionStore(tmp_path / "work", inbox)
    revision = store.create(source)
    result = store.process(revision.revision_id)

    corrected_path, _ = store.delivery_paths(result.process_id)
    corrected_path.write_bytes(b"tampered")
    with pytest.raises(RevisionError, match="delivery artifact"):
        store.process(revision.revision_id)

    with pytest.raises(RevisionError, match="process id"):
        store.load_process("../../escape")
    process_path = store.processes / f"{result.process_id}.json"
    process_path.write_text('{"schema_version": 1}', encoding="utf-8")
    with pytest.raises(RevisionError, match="process manifest is invalid"):
        store.load_process(result.process_id)


def test_process_loaders_fail_closed_on_missing_or_mismatched_state(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _brief(inbox / "brief.docx")

    missing_store = RevisionStore(tmp_path / "missing-work", inbox)
    with pytest.raises(RevisionError, match="process manifest does not exist"):
        missing_store.load_process(f"proc-{'9' * 64}")

    identity_store = RevisionStore(tmp_path / "identity-work", inbox)
    identity_revision = identity_store.create(source)
    identity_result = identity_store.process(identity_revision.revision_id)
    identity_path = identity_store.processes / f"{identity_result.process_id}.json"
    identity_payload = identity_store.load_process(identity_result.process_id).model_dump(
        mode="json"
    )
    identity_payload["process_id"] = f"proc-{'8' * 64}"
    identity_path.write_text(json.dumps(identity_payload), encoding="utf-8")
    with pytest.raises(RevisionError, match="identity does not match"):
        identity_store.load_process(identity_result.process_id)

    non_object_store = RevisionStore(tmp_path / "non-object-work", inbox)
    non_object_revision = non_object_store.create(source)
    non_object_result = non_object_store.process(non_object_revision.revision_id)
    non_object_path = non_object_store.processes / f"{non_object_result.process_id}.json"
    non_object_path.write_text("[]", encoding="utf-8")
    with pytest.raises(RevisionError, match="must be a JSON object"):
        non_object_store.load_process(non_object_result.process_id)

    mismatch_store = RevisionStore(tmp_path / "mismatch-work", inbox)
    mismatch_revision = mismatch_store.create(source)
    mismatch_result = mismatch_store.process(mismatch_revision.revision_id)
    mismatch_path = mismatch_store.processes / f"{mismatch_result.process_id}.json"
    mismatch_payload = mismatch_store.load_process(mismatch_result.process_id).model_dump(
        mode="json"
    )
    mismatch_payload["input_sha256"] = "0" * 64
    mismatch_path.write_text(json.dumps(mismatch_payload), encoding="utf-8")
    with pytest.raises(RevisionError, match="does not match this transform"):
        mismatch_store.process(mismatch_revision.revision_id)

    missing_delivery_store = RevisionStore(tmp_path / "missing-delivery-work", inbox)
    missing_delivery_revision = missing_delivery_store.create(source)
    missing_delivery_result = missing_delivery_store.process(
        missing_delivery_revision.revision_id
    )
    corrected, _ = missing_delivery_store.delivery_paths(missing_delivery_result.process_id)
    corrected.unlink()
    with pytest.raises(RevisionError, match="missing or unsafe"):
        missing_delivery_store.delivery_paths(missing_delivery_result.process_id)
